"""Migracion m014 - Composicion de los combos: congelar lo ya rebajado (5-oct-2026).

Decision del usuario (Decisiones Pendientes #85, opcion 3; diseno confirmado el 5-oct-2026):
la composicion de un combo se podra ajustar desde el SaaS AUNQUE el combo ya tenga ventas.
Lo ya vendido no se modifica (Regla #40). El dinero ya estaba protegido (cada venta guarda su
precio como numero); el INVENTARIO no: 'Unidades del producto' era Cantidad x TODOS los combos
vendidos, asi que cambiar la cantidad recalculaba tambien lo ya vendido.

Esta migracion solo PREPARA la hoja (no cambia ningun valor calculado):

  A) 'Tabla Combos Dinamicos', TablaComposicionCombos: dos columnas nuevas al final.
       G  'Combos ya congelados'    cuantos combos se habian vendido en el ultimo cambio (hoy 0)
       H  'Unidades ya rebajadas'   unidades del producto rebajadas hasta ese cambio (hoy 0)
     El rango pasa de A15:F{n} a A15:H{n}.
  B) La columna F 'Unidades del producto' deja de ser  D x E  y pasa a ser
       H + D x (E - G)
     es decir: lo ya rebajado (congelado) + la cantidad vigente x los combos vendidos DESPUES
     del ultimo cambio. Con G = 0 y H = 0 da EXACTAMENTE lo mismo que antes.
  C) DiccionarioDatos: 2 filas nuevas al final (las dos columnas nuevas).
  D) Evento MIGRACION en RegistroEventos.

Como se usara (funcion del robot, paquete siguiente; aqui NO se construye):
  * Cambia la CANTIDAD de un componente: el robot anota G = combos vendidos hasta hoy y
    H = unidades rebajadas hasta hoy, y despues escribe la cantidad nueva. F no se mueve.
  * Cambia el PRODUCTO de un componente: el robot agrega una fila de historia (producto
    anterior, cantidad 0, H = lo ya rebajado) para que esas unidades sigan cargadas al
    producto anterior, y la fila vigente empieza en 0 con el producto nuevo.

No se quita ni se mueve nada. 'Control de Inventario'!I (Salidas en Combos) sigue sumando la
columna F por producto, con la misma formula.
Esperado en comparar_libros: tabla A15:F23 -> A15:H23; 'Tabla Combos Dinamicos' 26 celdas
(G15, H15, G16:H23 y las 8 formulas de F16:F23); DiccionarioDatos 28 celdas (2 filas x 14);
RegistroEventos 1 fila; 0 valores calculados cambiados; 0 errores nuevos.

Revisiones ANTES de escribir (si una falla, se detiene y no escribe nada):
  - La tabla es A15:F{n}, tiene 6 columnas y la ultima es 'Unidades del producto'.
  - Cada celda de F es exactamente  =D{fila}*E{fila}  y cada D es un numero.
  - Las columnas G y H estan vacias desde la fila del titulo de la tabla hacia abajo.
  - Las 2 filas nuevas del diccionario estan vacias y sus ID no existen ya.
  - Ningun texto nuevo empieza con '='.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m014_congelar_combos --local   (copia datos_prueba/maestro_prueba.xlsx)
    py -m migraciones.m014_congelar_combos           (Excel MAESTRO de Google Drive)
    py -m migraciones.m014_congelar_combos --demo    (COPIA DE DEMOSTRACION de Google Drive)
Si ya estaba aplicada, responde "Ya estaba aplicada" y no escribe nada.

--demo usa la misma configuracion del SaaS (saas_demo/config.py): toma SAAS_EXCEL_FILE_ID del
.env, no arranca si es el mismo archivo que el maestro, y antes de escribir comprueba que el
archivo de Drive tenga la palabra DEMO en el nombre.
"""
import os
import sys
from copy import copy

if "--local" in sys.argv and "--demo" in sys.argv:
    sys.exit("Use --local o --demo, no los dos.")
# --local ANTES de importar conexion_drive.config (igual que las demas migraciones)
if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

FECHA = "5-oct-2026"
ID_MIGRACION = "m014_congelar_combos"
HOJA, DIC = "Tabla Combos Dinamicos", "DiccionarioDatos"
TABLA = "TablaComposicionCombos"
COL_ULTIMA = "Unidades del producto"
COL_CONGELADOS, COL_REBAJADAS = "Combos ya congelados", "Unidades ya rebajadas"

DIC_FILAS = [
    ["MKT.COMBO.COMP.CONGELADOS", "Marketing", HOJA, TABLA, "G16:G23 (Combos ya congelados)", "Número",
     "Robot (composicion de combos)", "Python", "Excel (columna F de la misma tabla)", "Sí", "No", "—",
     "MKT.COMBO.NOMBRE",
     f"CREADO m014 ({FECHA}, #85 opcion 3): cuantos combos se habian vendido cuando se cambio por ultima vez "
     "la composicion de esa fila. La escribe solo el robot al guardar un cambio de composicion; 0 si nunca "
     "cambio. La columna F cuenta con la cantidad vigente solo los combos vendidos DESPUES de ese numero."],
    ["MKT.COMBO.COMP.UNIDADES_CONGELADAS", "Marketing", HOJA, TABLA, "H16:H23 (Unidades ya rebajadas)", "Número",
     "Robot (composicion de combos)", "Python", "Excel (columna F de la misma tabla)", "Sí", "No", "—",
     "MKT.COMBO.COMP.CONGELADOS",
     f"CREADO m014 ({FECHA}, #85 opcion 3): unidades del producto ya rebajadas del inventario hasta el ultimo "
     "cambio de composicion (Regla #40: lo ya vendido no se modifica). La escribe solo el robot. "
     "'Unidades del producto' (F) = H + Cantidad por combo x (Combos vendidos - Combos ya congelados). "
     "Si cambia el PRODUCTO de un componente, lo ya rebajado queda en una fila de historia con cantidad 0."],
]


class MigracionDetenida(Exception):
    """El Excel no esta como se esperaba: no se escribe nada."""


def _revisar(condicion, mensaje):
    if not condicion:
        raise MigracionDetenida(mensaje + " No se escribio nada.")


def _copiar_estilo(origen, destino):
    if origen.has_style:
        destino.font, destino.border, destino.fill = copy(origen.font), copy(origen.border), copy(origen.fill)
        destino.number_format = origen.number_format
        destino.alignment, destino.protection = copy(origen.alignment), copy(origen.protection)


def _partes_ref(ref):
    """'A15:F23' -> ('A', 15, 'F', 23)"""
    from openpyxl.utils.cell import coordinate_from_string
    inicio, fin = ref.split(":")
    c1, f1 = coordinate_from_string(inicio)
    c2, f2 = coordinate_from_string(fin)
    return c1, f1, c2, f2


def formula_nueva(fila):
    return f"=H{fila}+D{fila}*(E{fila}-G{fila})"


def ya_aplicada(wb):
    tabla = wb[HOJA].tables.get(TABLA)
    return tabla is not None and [c.name for c in tabla.tableColumns][-2:] == [COL_CONGELADOS, COL_REBAJADAS]


def fila_diccionario(ws):
    """Primera fila vacia despues de la ultima con ID en la columna A."""
    ultima = max((c.row for c in ws["A"] if c.value not in (None, "")), default=4)
    return ultima + 1


def revisar_antes(wb):
    ws = wb[HOJA]
    _revisar(TABLA in ws.tables, f"No existe {TABLA} en '{HOJA}'.")
    tabla = ws.tables[TABLA]
    c1, f1, c2, f2 = _partes_ref(tabla.ref)
    _revisar((c1, c2) == ("A", "F") and f2 > f1, f"{TABLA} deberia ser A(n):F(n) y es {tabla.ref}.")
    nombres = [c.name for c in tabla.tableColumns]
    _revisar(len(nombres) == 6 and nombres[-1] == COL_ULTIMA,
             f"{TABLA} deberia tener 6 columnas terminando en '{COL_ULTIMA}': {nombres}.")
    _revisar(ws[f"F{f1}"].value == COL_ULTIMA, f"'{HOJA}'!F{f1} no es '{COL_ULTIMA}'.")
    for r in range(f1 + 1, f2 + 1):
        formula = str(ws[f"F{r}"].value or "").replace(" ", "")
        _revisar(formula == f"=D{r}*E{r}", f"'{HOJA}'!F{r} deberia ser =D{r}*E{r} y es {ws[f'F{r}'].value!r}.")
        cantidad = ws[f"D{r}"].value
        _revisar(isinstance(cantidad, (int, float)) and not isinstance(cantidad, bool),
                 f"'{HOJA}'!D{r} (Cantidad por combo) deberia ser un numero y es {cantidad!r}.")
    ocupadas = [ws[f"{col}{r}"].coordinate for col in ("G", "H") for r in range(f1 - 1, ws.max_row + 1)
                if ws[f"{col}{r}"].value not in (None, "")]
    _revisar(not ocupadas, f"Las columnas G y H deberian estar vacias desde la fila {f1 - 1} y tienen: {ocupadas[:10]}.")

    dic = wb[DIC]
    _revisar([c.value for c in dic[4]][:2] == ["ID variable", "Módulo"], "DiccionarioDatos fila 4 no es el encabezado.")
    ids = {str(c.value).strip() for c in dic["A"] if c.value}
    repetidos = [f[0] for f in DIC_FILAS if f[0] in ids]
    _revisar(not repetidos, f"Ya existen en DiccionarioDatos: {repetidos}.")
    inicio = fila_diccionario(dic)
    ocupadas = [c.coordinate for r in range(inicio, inicio + len(DIC_FILAS)) for c in dic[r] if c.value is not None]
    _revisar(not ocupadas, f"Las filas nuevas del diccionario no estan vacias: {ocupadas[:10]}.")

    textos = [t for fila in DIC_FILAS for t in fila] + [COL_CONGELADOS, COL_REBAJADAS]
    _revisar(not any(isinstance(t, str) and t.startswith("=") for t in textos), "Hay un texto que empieza con '='.")
    return {"fila_titulos": f1, "ultima_fila_tabla": f2, "fila_dic": inicio}


def aplicar(wb, info):
    from openpyxl.worksheet.table import TableColumn

    ws = wb[HOJA]
    tabla = ws.tables[TABLA]
    f1, f2 = info["fila_titulos"], info["ultima_fila_tabla"]
    # A) columnas nuevas: encabezado y ceros, con el estilo de la columna F
    for col, nombre in (("G", COL_CONGELADOS), ("H", COL_REBAJADAS)):
        _copiar_estilo(ws[f"F{f1}"], ws[f"{col}{f1}"])
        ws[f"{col}{f1}"] = nombre
        for r in range(f1 + 1, f2 + 1):
            _copiar_estilo(ws[f"F{r}"], ws[f"{col}{r}"])
            ws[f"{col}{r}"] = 0
    siguiente_id = max(c.id for c in tabla.tableColumns) + 1
    tabla.tableColumns.append(TableColumn(id=siguiente_id, name=COL_CONGELADOS))
    tabla.tableColumns.append(TableColumn(id=siguiente_id + 1, name=COL_REBAJADAS))
    tabla.ref = f"A{f1}:H{f2}"
    if tabla.autoFilter is not None:
        tabla.autoFilter.ref = tabla.ref

    # B) la columna F: lo congelado + cantidad vigente x combos vendidos despues del ultimo cambio
    for r in range(f1 + 1, f2 + 1):
        ws[f"F{r}"] = formula_nueva(r)

    # C) DiccionarioDatos: filas nuevas al final (estilo de la ultima fila con datos)
    dic = wb[DIC]
    modelo = info["fila_dic"] - 1
    for i, valores in enumerate(DIC_FILAS):
        r = info["fila_dic"] + i
        for j, valor in enumerate(valores, start=1):
            _copiar_estilo(dic.cell(modelo, j), dic.cell(r, j))
            dic.cell(r, j).value = valor
    filas = f2 - f1
    return {"A": 2 + 2 * filas, "B": filas, "C": len(DIC_FILAS) * len(DIC_FILAS[0])}


def verificar(wb, info):
    ws = wb[HOJA]
    tabla = ws.tables[TABLA]
    f1, f2 = info["fila_titulos"], info["ultima_fila_tabla"]
    _revisar(tabla.ref == f"A{f1}:H{f2}", f"Verificacion: la tabla quedo {tabla.ref}.")
    _revisar([c.name for c in tabla.tableColumns][-2:] == [COL_CONGELADOS, COL_REBAJADAS], "Verificacion: columnas nuevas.")
    _revisar(ws[f"G{f1}"].value == COL_CONGELADOS and ws[f"H{f1}"].value == COL_REBAJADAS, "Verificacion: encabezados G y H.")
    _revisar(all(ws[f"F{r}"].value == formula_nueva(r) and ws[f"G{r}"].value == 0 and ws[f"H{r}"].value == 0
                 for r in range(f1 + 1, f2 + 1)), "Verificacion: formulas de F y ceros de G y H.")
    _revisar(wb[DIC].cell(info["fila_dic"], 1).value == DIC_FILAS[0][0], "Verificacion: DiccionarioDatos.")


def main():
    demo = "--demo" in sys.argv
    palabra_copia = "DEMO"
    if demo:
        # La misma configuracion del SaaS: valida SAAS_EXCEL_FILE_ID (nunca el maestro) y deja al
        # robot apuntando a la copia. Debe importarse ANTES de conexion_drive.
        try:
            from saas_demo import config as config_saas
        except Exception as error:      # noqa: BLE001 (falta una variable SAAS_ o apunta al maestro)
            sys.exit(f"DETENIDA: la configuracion de la copia de demostracion no es valida: {error} "
                     "No se escribio nada.")
        if config_saas.ENSAYO:
            sys.exit("--demo no se usa en modo ensayo. No se escribio nada.")
        palabra_copia = config_saas.PALABRA_COPIA

    from conexion_drive import config
    from conexion_drive.drive_excel import editar_excel, obtener_metadatos
    from conexion_drive.tablas import registrar_evento

    if config.ORIGEN_EXCEL == "local":
        destino = "COPIA LOCAL " + str(config.EXCEL_LOCAL_PATH)
    else:
        nombre = str(obtener_metadatos().get("name") or "")
        es_copia = palabra_copia.lower() in nombre.lower()
        if demo and not es_copia:
            sys.exit(f"DETENIDA: el archivo de Drive se llama '{nombre}' y no contiene la palabra "
                     f"'{palabra_copia}': no parece la copia de demostracion. No se escribio nada.")
        if not demo and es_copia:
            sys.exit(f"DETENIDA: el archivo de Drive se llama '{nombre}': parece la copia de demostracion. "
                     "Para ella se usa --demo. No se escribio nada.")
        destino = f"GOOGLE DRIVE, {'COPIA DE DEMOSTRACION' if demo else 'EXCEL MAESTRO'}: {nombre}"
    print(f"=== {ID_MIGRACION} -> {destino} ===")

    def editar(wb, wb_valores):
        if ya_aplicada(wb):
            return {"hubo_cambios": False, "estado": "YA_APLICADA"}
        info = revisar_antes(wb)
        conteo = aplicar(wb, info)
        verificar(wb, info)
        f1, f2 = info["fila_titulos"], info["ultima_fila_tabla"]
        registrar_evento(
            wb, modulo="Migraciones", tipo_evento="MIGRACION", entidad=ID_MIGRACION, id_entidad="m014",
            accion="Combos: columnas para congelar lo ya rebajado al cambiar la composicion",
            estado="OK",
            detalle=(f"A) {TABLA} A{f1}:F{f2} -> A{f1}:H{f2} (G {COL_CONGELADOS}, H {COL_REBAJADAS}, en 0) | "
                     f"B) F{f1 + 1}:F{f2} de D x E a H + D x (E - G), mismo resultado | C) DiccionarioDatos filas "
                     f"{info['fila_dic']}-{info['fila_dic'] + len(DIC_FILAS) - 1} | #85 opcion 3, Regla #40"),
            origen="Robot Python",
        )
        return {"hubo_cambios": True, "estado": "APLICADA", "conteo": conteo, "info": info}

    try:
        resultado = editar_excel(editar, descripcion=f"Migracion {ID_MIGRACION}")
    except MigracionDetenida as error:
        print(f"DETENIDA: {error}")
        sys.exit(1)
    if resultado["estado"] == "YA_APLICADA":
        print("Ya estaba aplicada. No se escribio nada.")
        return
    info = resultado["info"]
    f1, f2 = info["fila_titulos"], info["ultima_fila_tabla"]
    print(f"  Bloque A: {TABLA} A{f1}:F{f2} -> A{f1}:H{f2} (G {COL_CONGELADOS}, H {COL_REBAJADAS})")
    print(f"  Bloque B: F{f1 + 1}:F{f2} ahora es H + D x (E - G)")
    print(f"  Bloque C: DiccionarioDatos filas {info['fila_dic']} a {info['fila_dic'] + len(DIC_FILAS) - 1}")
    for bloque, n in resultado["conteo"].items():
        print(f"  Bloque {bloque}: {n} celdas")
    print("LISTO.")


if __name__ == "__main__":
    main()
