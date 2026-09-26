"""Migracion m009 - Inventario de prueba cuadrado (Decisiones Pendientes #18).

Decisiones del usuario (26-sep-2026): es un estudio academico; se trata de CUADRAR un sistema
completo con datos supuestos. Las cantidades de 'Ingreso de Inventario' eran relleno de cuando se
creo el Excel. Las rebajas de inventario las hace SIEMPRE el Excel con sus formulas
('Control de Inventario' = compras + produccion - ventas - combos - consumo en produccion);
lo digitado son valores, segun el libro de mapas.

La semana de prueba es del 2 al 8 de marzo de 2026 (ventas de despacho, pedidos digitales y
hornadas del 'Historico Diario'). Que hace:

  A) Fechas: 'Ingreso de Inventario' A4:A9 (compras externas) pasan a la semana de prueba
     (2, 2, 4, 5, 6 y 7 de marzo) y A10:A20 (materia prima) al 2 de marzo, antes de hornear.
  B) Cantidades de compra (columna D) = lo justo para las hornadas del Historico, redondeado
     hacia arriba a empaques enteros (harina 16, levadura 19, mejorante 7, bolsas 15, burbuja 6,
     margarina 2, azucar 1, sal 1, manteca 1, cafe 1, mermelada 1) y leche 33 L (20 + 13).
  C) Produccion vieja H4:H9 (550 Baguette / 200 CupCake de relleno, no corresponden a ninguna
     hornada) -> 0; la produccion real de prueba entra en D). G4:G9 y el resto de la fila no cambian.
  D) Una fila de produccion en 'Ingreso de Inventario' y sus filas de consumo en TablaConsumoMP
     por CADA hornada del 'Historico Diario' (16: 1,500 Baguette y 1,430 CupCake), escritas con
     las MISMAS funciones del robot (hornadas.py / ingreso_inventario.py). Lote H-AAAAMMDD-HHMM-cod.
  E) Compra de prueba de Natilla (15 unid., 2 de marzo, factura F004): se vendian 12 y nunca
     se habia comprado (existencia -12).

Resultado esperado (lo calcula el Excel): ninguna existencia negativa en 'Control de Inventario'
(Baguette 246, CupCake 124); materia prima de prueba = 490,800.

Seguridad: revisa TODO antes de escribir (celdas con el valor esperado, 16 hornadas con esos
totales); se detiene sin cambiar nada si algo no calza; se puede correr 2 veces (no repite lotes
ni la factura F004); relee cada celda; evento MIGRACION con los cambios.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m009_inventario_prueba --local
    py -m migraciones.m009_inventario_prueba
    py -m migraciones.m009_inventario_prueba --archivo "ruta.xlsx"
"""
import os
import sys
from datetime import datetime

if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

ID_MIGRACION = "m009_inventario_prueba"
FECHA = "26-sep-2026"
ING = "Ingreso de Inventario"
HIS = "Historico Diario"


def _d(mes, dia):
    return datetime(2026, mes, dia)


# A) fechas: celda -> (antes, despues)
FECHAS = {"A4": (_d(5, 4), _d(3, 2)), "A5": (_d(5, 4), _d(3, 2)), "A6": (_d(5, 6), _d(3, 4)),
          "A7": (_d(5, 7), _d(3, 5)), "A8": (_d(5, 8), _d(3, 6)), "A9": (_d(5, 9), _d(3, 7))}
for _f in range(10, 21):
    FECHAS[f"A{_f}"] = (_d(5, _f), _d(3, 2))

# B) cantidades de compra: celda -> (producto esperado en C, antes, despues)
CANTIDADES = {
    "D7": ("Leche Semidescremada Dos Pinos -1 L", 10, 13),
    "D10": ("MEJORANTE PANNIVITA PRO-PAN 1 KG HOMO-PLUS", 2, 7),
    "D11": ("HARINA FLORES 25 KG FUERTE", 10, 16),
    "D12": ("AZUCAR DOÑA MARIA 2KG (24 KG)", 15, 1),
    "D13": ("SAL SARU COMUN ECONOMICA 25 KG SACO BRIN", 20, 1),
    "D14": ("MANTECA MAX PAN BULK 25 KGS", 12, 1),
    "D15": ("LEVADURA FRESCA GOLONDRINA 400 GR CIP", 20, 19),
    "D16": ("MERMELADA GUAYABA HORNEO 5 KG CCR CAJA C2 (FRUTA TRITURADA)", 30, 1),
    "D17": ("BOLSAS PAPEL KRAFT 2 BAGUETTE 100 UNDS BIODEGRADABLE", 100, 15),
    "D18": ("BURBUJA FOAM #6 125 UNID POSTRE BL", 100, 6),
    "D19": ("Café Naranjo puro 250gr x 10 Unidades", 20, 1),
    "D20": ("MARGARINA TUMBA NUMAR 7.5 KILOS 60 BARRITAS X125GR", 20, 2),
}

# C) produccion vieja: celda -> (antes, despues)
PRODUCCION_VIEJA = {"H4": (500, 0), "H5": (50, 0), "H6": (50, 0), "H7": (50, 0), "H8": (50, 0), "H9": (50, 0)}

# D) hornadas esperadas en el Historico
HORNADAS_ESPERADAS = 16
TOTALES_ESPERADOS = {"P01": 1500, "Pa01": 1430}

# E) compra de natilla
NATILLA = ("Natilla Dos Pinos La Ganja Con Sal - 300g", 15, "F004", _d(3, 2))


class MigracionDetenida(Exception):
    """Una celda no tenia el valor esperado: no se escribio nada."""


def _hornadas(wb):
    """Hornadas del Historico: lista de (fecha, hora, codigo, cantidad, lote)."""
    from conexion_drive.tablas import leer_tabla, obtener_tabla
    from microservicios.grupo2_inventario.hornadas import _codigo, _como_hora, _lote
    ws, tabla = obtener_tabla(wb, HIS, "HistoricoHorneadas")
    lista = []
    for fila in leer_tabla(ws, tabla):
        fecha, cantidad = fila.get("Fecha"), fila.get("Cantidad producida")
        if not isinstance(fecha, datetime) or not isinstance(cantidad, (int, float)) or cantidad <= 0:
            continue
        hora = _como_hora(fila.get("Hora"))
        codigo = _codigo(fila.get("Producto"))
        fecha = datetime(fecha.year, fecha.month, fecha.day)
        lista.append((fecha, hora, codigo, int(cantidad), _lote(fecha, hora, codigo)))
    return lista


def _valores_columna(ws, col, desde=4):
    return {ws[f"{col}{r}"].value for r in range(desde, ws.max_row + 1)}


def migrar(wb, wb_valores=None):
    from conexion_drive.tablas import obtener_tabla, registrar_evento
    from microservicios.grupo1_captura_operativa.ingreso_inventario import escribir_fila_ingreso
    from microservicios.grupo2_inventario.hornadas import _escribir_consumo, _nombre_maestro

    if wb_valores is None:
        raise MigracionDetenida("Se necesitan los valores calculados del libro (receta de consumo).")
    ws_i = wb[ING]

    # 1) Revisar TODO antes de escribir
    problemas, celdas = [], []
    for celda, (antes, despues) in FECHAS.items():
        v = ws_i[celda].value
        if v != despues:
            if v != antes:
                problemas.append(f"'{ING}'!{celda}: se esperaba {antes:%d-%m-%Y} y tiene {v!r}")
            else:
                celdas.append((celda, antes, despues, "A"))
    for celda, (producto, antes, despues) in CANTIDADES.items():
        fila = celda[1:]
        if ws_i[f"C{fila}"].value != producto:
            problemas.append(f"'{ING}'!C{fila}: se esperaba {producto!r} y tiene {ws_i[f'C{fila}'].value!r}")
        v = ws_i[celda].value
        if v != despues:
            if v != antes:
                problemas.append(f"'{ING}'!{celda}: se esperaba {antes!r} y tiene {v!r}")
            else:
                celdas.append((celda, antes, despues, "B"))
    for celda, (antes, despues) in PRODUCCION_VIEJA.items():
        v = ws_i[celda].value
        if v != despues:
            if v != antes:
                problemas.append(f"'{ING}'!{celda}: se esperaba {antes!r} y tiene {v!r}")
            else:
                celdas.append((celda, antes, despues, "C"))

    hornadas = _hornadas(wb)
    totales = {}
    for _, _, codigo, cantidad, _ in hornadas:
        totales[codigo] = totales.get(codigo, 0) + cantidad
    if len(hornadas) != HORNADAS_ESPERADAS or totales != TOTALES_ESPERADOS:
        problemas.append(f"'{HIS}': se esperaban {HORNADAS_ESPERADAS} hornadas con {TOTALES_ESPERADOS} y hay "
                         f"{len(hornadas)} con {totales}")
    lotes_existentes = _valores_columna(ws_i, "B")
    pendientes_h = [h for h in hornadas if h[4] not in lotes_existentes]
    natilla_pendiente = NATILLA[2] not in lotes_existentes

    if problemas:
        raise MigracionDetenida("NO SE CAMBIO NADA. Celdas distintas a lo esperado:\n  " + "\n  ".join(problemas))
    if not celdas and not pendientes_h and not natilla_pendiente:
        return {"hubo_cambios": False, "cambios": ["Ya estaba aplicada: 0 celdas pendientes."], "detalle": []}

    # 2) Escribir A, B, C y releer
    for celda, antes, despues, bloque in celdas:
        ws_i[celda] = despues
    for celda, antes, despues, bloque in celdas:
        if ws_i[celda].value != despues:
            raise MigracionDetenida(f"Verificacion fallo en '{ING}'!{celda}")

    # 3) D) hornadas: produccion + consumo, con las funciones del robot
    ws_t, tabla_i = obtener_tabla(wb, ING, "TablaIngresos")
    filas_prod, filas_cons = [], []
    for fecha, hora, codigo, cantidad, lote in pendientes_h:
        nombre_inv = _nombre_maestro(wb_valores, codigo)
        filas_prod.append(escribir_fila_ingreso(ws_t, tabla_i, fecha, lote, nombre_inv, cantidad, es_produccion=True))
        filas_cons += _escribir_consumo(wb, wb_valores, fecha, lote, codigo, nombre_inv, cantidad)

    # 4) E) compra de natilla
    fila_natilla = None
    if natilla_pendiente:
        nombre, cantidad, factura, fecha = NATILLA
        fila_natilla = escribir_fila_ingreso(ws_t, tabla_i, fecha, factura, nombre, cantidad,
                                             es_produccion=False, proveedor="Proveedor A")

    por_bloque = {}
    for c in celdas:
        por_bloque[c[3]] = por_bloque.get(c[3], 0) + 1
    resumen = [f"Bloque {b}: {n} celdas" for b, n in sorted(por_bloque.items())]
    resumen.append(f"Bloque D: {len(pendientes_h)} hornadas -> {len(filas_prod)} filas de produccion y "
                   f"{len(filas_cons)} filas de consumo")
    resumen.append(f"Bloque E: compra de natilla {'fila ' + str(fila_natilla) if natilla_pendiente else 'ya estaba'}")
    antes_despues = " ; ".join(
        f"{c}: {a:%d-%m-%Y} -> {d:%d-%m-%Y}" if isinstance(a, datetime) else f"{c}: {a} -> {d}"
        for c, a, d, b in celdas)
    registrar_evento(
        wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
        id_entidad=ID_MIGRACION, accion="Inventario de prueba cuadrado (semana 2-8 mar 2026)", estado="OK",
        detalle=(f"A) fechas B) compras justas C) produccion vieja a 0 | D) {len(pendientes_h)} hornadas del "
                 f"Historico: produccion filas {filas_prod}, consumo {len(filas_cons)} filas | E) natilla fila "
                 f"{fila_natilla} | {antes_despues}"),
    )
    return {"hubo_cambios": True, "cambios": resumen, "detalle": celdas}


def _aplicar_a_archivo(ruta):
    import io
    from pathlib import Path
    import openpyxl
    from conexion_drive.recalculo import reparar_prefijos_xlfn, recalcular_bytes

    ruta = Path(ruta)
    contenido, _ = reparar_prefijos_xlfn(ruta.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    wb_valores = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    resultado = migrar(wb, wb_valores)
    if resultado["hubo_cambios"]:
        salida = io.BytesIO()
        wb.save(salida)
        ruta.write_bytes(recalcular_bytes(salida.getvalue()))
    return resultado


if __name__ == "__main__":
    try:
        if "--archivo" in sys.argv:
            ruta = sys.argv[sys.argv.index("--archivo") + 1]
            print(f"Excel de destino: ARCHIVO {ruta}")
            resultado = _aplicar_a_archivo(ruta)
        else:
            from conexion_drive import config
            from conexion_drive.drive_excel import editar_excel
            print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
            resultado = editar_excel(migrar, descripcion="Migracion m009_inventario_prueba")
    except MigracionDetenida as error:
        print(str(error))
        sys.exit(1)
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
