"""Migracion 001 — preparar el Excel maestro para el Grupo 2 (se corre UNA sola vez).

Que cambia (decisiones confirmadas por el usuario, 23-sep-2026):

  1. 'Tabla Combos Dinamicos': agrega la tabla TablaComposicionCombos (A15:F23)
     con lo que trae cada combo, sacado de sus formulas de precio:
        Combo 1 = 3 Baguette + 1 Leche      Combo 2 = 2 CupCake + 1 Coca Cola
        Combo 3 = 2 Baguette + 1 Natilla    Combo 4 = 3 Baguette + 1 Queso
     Cada fila cuenta cuantos combos se vendieron (despacho ORDEN CORRECTA +
     digital CONFIRMADO) y cuantas unidades del producto salieron por ellos.
     Todo son formulas: los nombres se leen de TABLA CRUCE MAESTRA y de
     Tabla Combos Dinamicos, asi que si un nombre cambia, la tabla lo sigue.

  2. 'Control de Inventario': columna nueva I 'Salidas en Combos' y
     Existencia Actual = Entradas - Ventas Despacho - Ventas Digitales - Salidas en Combos.

  3. 'Historico Diario':
     a) Filas 5 y 6 dejan de ser formulas vivas enlazadas a las hojas de costo
        y quedan como VALORES FIJOS (paso 3 del 'Diseño registro' del mapa).
        Sin esto, cada hornada nueva sobrescribiria la anterior.
     b) Las filas de prueba que dicen 'Baguette' pasan a 'pan baguette', el
        nombre que busca 'flujo caja anual' (sin esto, no se cuentan).

  4. Precio de referencia (idea del usuario, 23-sep-2026): en cada hoja de costo
     se agrega UNA celda que digita la administracion con el precio de venta:
        'costo pan baguette'!G28 (arranca en 842.69)   'costo cupcakes'!G27 (arranca en 343.16)
     - TABLA CRUCE MAESTRA (Precio de Baguette y CupCake) pasa a leer esa celda,
       asi el precio al publico ya no cambia con el tamaño de cada hornada.
     - El precio calculado por costo (G29 / G28 = costo x margen) SIGUE IGUAL,
       como referencia de como se mueve el costo.
     - Columna H: 'Utilidad real' y '% Utilidad real' con el precio de referencia,
       para ver la ganancia extra al producir mas.

Es SEGURA de repetir: si ya se aplico, no cambia nada.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m001_grupo2 --local    <- primero, sobre datos_prueba/maestro_prueba.xlsx
    py -m migraciones.m001_grupo2            <- despues, sobre el Excel real en Google Drive
"""
import os
import sys

if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

import logging                                                    # noqa: E402
from copy import copy                                             # noqa: E402

from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo  # noqa: E402

from conexion_drive import editar_excel                           # noqa: E402
from conexion_drive import config                                 # noqa: E402
from conexion_drive.tablas import registrar_evento                # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")

HOJA_COMBOS = "Tabla Combos Dinamicos"
TABLA_COMPOSICION = "TablaComposicionCombos"
FILA_TITULO, FILA_ENCABEZADO = 14, 15

# (fila del combo en Tabla Combos Dinamicos, fila del producto en TABLA CRUCE MAESTRA, cantidad)
# Sacado de las formulas de 'Precio de Oferta' (E5:E8):
#   E5 = (3*G8)*0.9 + G10   E6 = (2*G9)*0.9 + G13   E7 = (2*G8)*0.9 + G14   E8 = (3*G8)*0.9 + G11
COMPOSICION = [
    (5, 8, 3), (5, 10, 1),     # Combo 1: 3 Baguette + 1 Leche Suli
    (6, 9, 2), (6, 13, 1),     # Combo 2: 2 CupCake  + 1 Coca Cola
    (7, 8, 2), (7, 14, 1),     # Combo 3: 2 Baguette + 1 Natilla
    (8, 8, 3), (8, 11, 1),     # Combo 4: 3 Baguette + 1 Queso
]
ENCABEZADOS = ["Codigo combo", "Nombre del combo", "Producto", "Cantidad por combo",
               "Combos vendidos", "Unidades del producto"]

FORMULA_VENDIDOS = (
    "=SUMIFS('Ventas despacho'!$F$5:$F$50000,'Ventas despacho'!$E$5:$E$50000,B{r},"
    "'Ventas despacho'!$I$5:$I$50000,\"ORDEN CORRECTA\")"
    "+SUMIFS('Detalle comanda digital'!$D$4:$D$50010,'Detalle comanda digital'!$C$4:$C$50010,B{r},"
    "'Detalle comanda digital'!$E$4:$E$50010,\"CONFIRMADO\")"
)


def _copiar_estilo(origen, destino):
    destino.font = copy(origen.font)
    destino.border = copy(origen.border)
    destino.alignment = copy(origen.alignment)
    destino.number_format = origen.number_format


def paso_1_composicion(wb, cambios):
    ws = wb[HOJA_COMBOS]
    if TABLA_COMPOSICION in ws.tables:
        cambios.append("1) Composicion de combos: ya existia, sin cambios")
        return False
    ultima = FILA_ENCABEZADO + len(COMPOSICION)
    ocupadas = [c.coordinate for fila in ws.iter_rows(min_row=FILA_TITULO, max_row=ultima, max_col=6)
                for c in fila if c.value is not None]
    if ocupadas:
        raise RuntimeError(f"El espacio A{FILA_TITULO}:F{ultima} de '{HOJA_COMBOS}' no esta vacio: {ocupadas[:5]}")

    ws[f"A{FILA_TITULO}"] = "COMPOSICIÓN DE COMBOS (para descontar inventario — no borrar)"
    _copiar_estilo(ws["C3"], ws[f"A{FILA_TITULO}"])
    for i, texto in enumerate(ENCABEZADOS):
        celda = ws.cell(FILA_ENCABEZADO, 1 + i, texto)
        _copiar_estilo(ws.cell(4, 1 + i), celda)

    for n, (fila_combo, fila_producto, cantidad) in enumerate(COMPOSICION):
        r = FILA_ENCABEZADO + 1 + n
        ws[f"A{r}"] = f"=$F${fila_combo}"
        ws[f"B{r}"] = f"=$C${fila_combo}"
        ws[f"C{r}"] = f"='TABLA CRUCE MAESTRA'!$C${fila_producto}"
        ws[f"D{r}"] = cantidad
        ws[f"E{r}"] = FORMULA_VENDIDOS.format(r=r)
        ws[f"F{r}"] = f"=D{r}*E{r}"

    tabla = Table(displayName=TABLA_COMPOSICION, ref=f"A{FILA_ENCABEZADO}:F{ultima}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tabla)
    cambios.append(f"1) Tabla {TABLA_COMPOSICION} creada en '{HOJA_COMBOS}'!A{FILA_ENCABEZADO}:F{ultima}")
    return True


def paso_2_control_inventario(wb, cambios):
    ws = wb["Control de Inventario"]
    tabla = ws.tables["TablaControlInventario"]
    if any(col.name == "Salidas en Combos" for col in tabla.tableColumns):
        cambios.append("2) Control de Inventario: columna 'Salidas en Combos' ya existia, sin cambios")
        return False
    primera, ultima = 5, int(tabla.ref.split(":")[1][1:])
    ws["I4"] = "Salidas en Combos"
    _copiar_estilo(ws["H4"], ws["I4"])
    ws.column_dimensions["I"].width = max(ws.column_dimensions["E"].width or 14, 18)
    for r in range(primera, ultima + 1):
        if ws[f"A{r}"].value in (None, ""):
            continue
        ws[f"I{r}"] = f"=SUMIF({TABLA_COMPOSICION}[Producto],A{r},{TABLA_COMPOSICION}[Unidades del producto])"
        _copiar_estilo(ws[f"E{r}"], ws[f"I{r}"])
        ws[f"F{r}"] = f"=B{r}+C{r}-D{r}-E{r}-I{r}"
    tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name="Salidas en Combos"))
    tabla.ref = f"A4:I{ultima}"
    if tabla.autoFilter is not None:
        tabla.autoFilter.ref = tabla.ref
    cambios.append(f"2) Control de Inventario: columna I 'Salidas en Combos' + Existencia actualizada (filas {primera}-{ultima})")
    return True


def paso_3_historico(wb, wb_valores, cambios):
    ws, wv = wb["Historico Diario"], wb_valores["Historico Diario"]
    congeladas = 0
    for r in (5, 6):
        for col in "ABCDEFGHI":
            valor = ws[f"{col}{r}"].value
            if isinstance(valor, str) and valor.startswith("="):
                calculado = wv[f"{col}{r}"].value
                if calculado is None and col not in "I":
                    raise RuntimeError(f"Historico Diario!{col}{r} no tiene valor calculado; no se congela.")
                ws[f"{col}{r}"] = round(calculado, 2) if isinstance(calculado, float) else calculado
                congeladas += 1
    renombradas = 0
    for r in range(5, ws.max_row + 1):
        if ws[f"C{r}"].value == "Baguette":
            ws[f"C{r}"] = "pan baguette"
            renombradas += 1
    cambios.append(f"3) Historico Diario: {congeladas} celdas congeladas como valor, "
                   f"{renombradas} filas 'Baguette' -> 'pan baguette'")
    return bool(congeladas or renombradas)


# hoja: (celda precio referencia, celda etiqueta, celda costo unitario, fila utilidad, fila %,
#        celda precio por costo, celda TABLA CRUCE MAESTRA)
PRECIO_REFERENCIA = {
    "costo pan baguette": ("G28", "F28", "G30", 31, 32, "G29", "G8"),
    "costo cupcakes": ("G27", "F27", "G29", 30, 31, "G28", "G9"),
}


def paso_4_precio_referencia(wb, wb_valores, cambios):
    hechos = []
    for hoja, (celda_ref, celda_etq, celda_costo, fila_util, fila_pct, celda_calc, celda_tcm) in PRECIO_REFERENCIA.items():
        ws, tcm = wb[hoja], wb["TABLA CRUCE MAESTRA"]
        destino = f"='{hoja}'!${celda_ref[0]}${celda_ref[1:]}"
        if tcm[celda_tcm].value == destino:
            continue
        ocupada = [c for c in (celda_ref, celda_etq, f"H{fila_util - 2}", f"H{fila_util}", f"H{fila_pct}")
                   if ws[c].value not in (None, "")]
        if ocupada:
            raise RuntimeError(f"En '{hoja}' ya hay datos en {ocupada}; no se agrega el precio de referencia.")
        precio_actual = wb_valores["TABLA CRUCE MAESTRA"][celda_tcm].value
        ws[celda_etq] = "Precio de referencia (lo digita la administración)"
        _copiar_estilo(ws[f"F{fila_util - 2}"] if ws[f"F{fila_util - 2}"].value else ws[celda_calc.replace("G", "F")], ws[celda_etq])
        ws[celda_ref] = round(float(precio_actual), 2)
        _copiar_estilo(ws[celda_calc], ws[celda_ref])
        ws[celda_ref].fill = copy(ws["G4"].fill)                # mismo color que las celdas que se digitan
        ws[f"H{fila_util - 2}"] = "Con precio de referencia"
        ws[f"H{fila_util}"] = f"={celda_ref}-{celda_costo}"
        ws[f"H{fila_pct}"] = f"=H{fila_util}/{celda_costo}"
        _copiar_estilo(ws[f"G{fila_util}"], ws[f"H{fila_util}"])
        _copiar_estilo(ws[f"G{fila_pct}"], ws[f"H{fila_pct}"])
        tcm[celda_tcm] = destino
        hechos.append(f"{hoja}!{celda_ref}={round(float(precio_actual), 2)}")
    if not hechos:
        cambios.append("4) Precio de referencia: ya existia, sin cambios")
        return False
    cambios.append("4) Precio de referencia agregado (" + ", ".join(hechos) + "); TABLA CRUCE MAESTRA G8/G9 lo leen")
    return True


def migrar(wb, wb_valores):
    cambios = []
    hubo = [paso_1_composicion(wb, cambios),
            paso_2_control_inventario(wb, cambios),
            paso_3_historico(wb, wb_valores, cambios),
            paso_4_precio_referencia(wb, wb_valores, cambios)]
    if not any(hubo):
        return {"hubo_cambios": False, "cambios": cambios}
    registrar_evento(wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
                     id_entidad="m001_grupo2", accion="Preparar Grupo 2", estado="OK",
                     detalle=" | ".join(cambios))
    return {"hubo_cambios": True, "cambios": cambios}


if __name__ == "__main__":
    print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
    resultado = editar_excel(migrar, descripcion="Migracion m001_grupo2")
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
