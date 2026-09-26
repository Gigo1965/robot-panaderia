"""Migracion m008 - Nombres unicos de los productos propios (Baguette / CupCake).

Decision del usuario (26-sep-2026): un solo nombre por producto en TODO el documento, propios y
externos, alineado a la TABLA CRUCE MAESTRA tal como sale ahi: "Baguette" (C8) y "CupCake" (C9).
REEMPLAZA la decision del Grupo 2 de usar 'pan baguette' / 'cupcakes' en 'Historico Diario'
(libro de mapas: 'Diseño registro mapa de celdas' A25). Los productos externos y la materia
prima ya usan el nombre tecnico de la TCM / TCM MP: no cambian.

Que cambia (solo TEXTO; la logica de ninguna formula cambia):
  A) Datos: 'Ventas despacho' E, 'Detalle comanda digital' C, 'Historico Diario' C:
     celdas cuyo valor es exactamente 'cupcakes' / 'pan baguette' -> 'CupCake' / 'Baguette'.
     'Interface digital' F (resumen del pedido, ej. '14 Baguette, 6 cupcakes'): la palabra.
  B) Hojas de costo A6 (el robot hornadas.py lee aqui el nombre que escribe en el Historico):
     'costo pan baguette'!A6 'pan baguette' -> 'Baguette'; 'costo cupcakes'!A6 'cupcakes' -> 'CupCake'.
  C) 'flujo caja anual': las formulas que buscan en el Historico cambian SOLO el texto entre
     comillas: "pan baguette" -> "Baguette" y "cupcakes" -> "CupCake" (deben ser 120).
  D) Etiquetas y titulos: 'flujo caja anual' A, 'Punto de Equilibrio' A, 'flujo caja proyecto',
     'proyeccion ventas', 'proyeccion anual ventas', 'costo cupcakes' G3.

NO cambia: los NOMBRES DE LAS HOJAS ('costo cupcakes', 'costo pan baguette'), 'DiccionarioDatos'
y 'RegistroEventos' (son historial: no se reescriben), ni las bolsas kraft ('BAGUETTE' en mayuscula
es parte del nombre del proveedor).

Efecto esperado: 'flujo caja anual' da EXACTAMENTE los mismos numeros (Historico y formulas cambian
juntos). 'Control de Inventario' D6/E6 (CupCake) empiezan a encontrar las ventas de prueba que
decian 'cupcakes' (hoy D6 = 0): se cuadra en la m009 (inventario de prueba).

Seguridad (igual que m005/m007): revisa TODO antes de escribir; se detiene sin cambiar nada si
las celdas fijas no tienen el valor esperado o si las formulas de C no son 120; se puede correr
2 veces; relee cada celda; evento MIGRACION con cada cambio anterior -> nuevo.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m008_nombres_unicos --local
    py -m migraciones.m008_nombres_unicos
    py -m migraciones.m008_nombres_unicos --archivo "ruta.xlsx"
"""
import os
import re
import sys

if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

ID_MIGRACION = "m008_nombres_unicos"
FECHA = "26-sep-2026"
BAG, CUP = "Baguette", "CupCake"      # TABLA CRUCE MAESTRA C8 / C9

# A) columnas de datos donde la celda completa es el nombre viejo
COLUMNAS_EXACTAS = [("Ventas despacho", "E"), ("Detalle comanda digital", "C"), ("Historico Diario", "C")]
EXACTOS = {"cupcakes": CUP, "pan baguette": BAG}
# A) resumen de pedido en texto libre
COLUMNAS_TEXTO = [("Interface digital", "F")]

# B) celdas fijas: (hoja, celda, antes, despues)
FIJAS = [("costo pan baguette", "A6", "pan baguette", BAG),
         ("costo cupcakes", "A6", "cupcakes", CUP)]

# C) formulas
HOJA_FORMULAS = "flujo caja anual"
FORMULAS_ESPERADAS = 120
LITERALES = [('"pan baguette"', f'"{BAG}"'), ('"cupcakes"', f'"{CUP}"')]

# D) etiquetas: hoja -> columnas (None = toda la hoja)
ETIQUETAS = {"flujo caja anual": "A", "Punto de Equilibrio": "A", "flujo caja proyecto": None,
             "proyeccion ventas": None, "proyeccion anual ventas": None}
ETIQUETAS_FIJAS = [("costo cupcakes", "G3", "            Cup cakes", f"            {CUP}")]

# Reemplazos de palabras en texto (A de Interface digital y D). Sensibles a mayusculas:
# 'BAGUETTE' (bolsas kraft) y los nombres de hojas quedan fuera.
PALABRAS = [(re.compile(r"\bpan bague(?:t)?te\b"), BAG),      # incluye la errata 'pan baguete'
            (re.compile(r"\b[Cc]upcakes?\b"), CUP),
            (re.compile(r"\bbaguette\b"), BAG)]


class MigracionDetenida(Exception):
    """Una celda no tenia el valor esperado: no se escribio nada."""


def _cambiar_palabras(texto):
    nuevo = texto
    for patron, reemplazo in PALABRAS:
        nuevo = patron.sub(reemplazo, nuevo)
    return nuevo


def _es_formula(v):
    return isinstance(v, str) and v.startswith("=")


def _lista_cambios(wb):
    """Devuelve (cambios, problemas). cambios = (hoja, celda, antes, despues, bloque)."""
    cambios, problemas = [], []

    # A) datos exactos
    for hoja, col in COLUMNAS_EXACTAS:
        ws = wb[hoja]
        for celda in ws[col]:
            v = celda.value
            if isinstance(v, str) and v in EXACTOS:
                cambios.append((hoja, celda.coordinate, v, EXACTOS[v], "A"))
    # A) resumen de pedidos
    for hoja, col in COLUMNAS_TEXTO:
        for celda in wb[hoja][col]:
            v = celda.value
            if isinstance(v, str) and not _es_formula(v):
                nuevo = _cambiar_palabras(v)
                if nuevo != v:
                    cambios.append((hoja, celda.coordinate, v, nuevo, "A"))

    # B) celdas fijas (se verifican: antes o ya aplicada)
    for hoja, celda, antes, despues in FIJAS:
        v = wb[hoja][celda].value
        if v == antes:
            cambios.append((hoja, celda, antes, despues, "B"))
        elif v != despues:
            problemas.append(f"'{hoja}'!{celda}: se esperaba {antes!r} y tiene {v!r}")

    # C) formulas del flujo caja anual
    ws = wb[HOJA_FORMULAS]
    formulas = []
    ya_aplicadas = 0
    for fila in ws.iter_rows():
        for celda in fila:
            v = celda.value
            if not _es_formula(v):
                continue
            nuevo = v
            for viejo, bueno in LITERALES:
                nuevo = nuevo.replace(viejo, bueno)
            if nuevo != v:
                formulas.append((HOJA_FORMULAS, celda.coordinate, v, nuevo, "C"))
            elif any(bueno in v for _, bueno in LITERALES) and "Historico Diario" in v:
                ya_aplicadas += 1
    if formulas and len(formulas) + ya_aplicadas != FORMULAS_ESPERADAS:
        problemas.append(f"'{HOJA_FORMULAS}': se esperaban {FORMULAS_ESPERADAS} formulas y hay "
                         f"{len(formulas)} por cambiar + {ya_aplicadas} ya cambiadas")
    cambios += formulas

    # D) etiquetas
    for hoja, col in ETIQUETAS.items():
        ws = wb[hoja]
        celdas = ws[col] if col else [c for fila in ws.iter_rows() for c in fila]
        for celda in celdas:
            v = celda.value
            if isinstance(v, str) and not _es_formula(v):
                nuevo = _cambiar_palabras(v)
                if nuevo != v:
                    cambios.append((hoja, celda.coordinate, v, nuevo, "D"))
    for hoja, celda, antes, despues in ETIQUETAS_FIJAS:
        v = wb[hoja][celda].value
        if v == antes:
            cambios.append((hoja, celda, antes, despues, "D"))
        elif v != despues:
            problemas.append(f"'{hoja}'!{celda}: se esperaba {antes!r} y tiene {v!r}")

    return cambios, problemas


def migrar(wb, wb_valores=None):
    # 1) Revisar TODO antes de escribir
    cambios, problemas = _lista_cambios(wb)
    if problemas:
        raise MigracionDetenida("NO SE CAMBIO NADA. Celdas distintas a lo esperado:\n  "
                                + "\n  ".join(problemas))

    # 2) Ningun TEXTO empieza con "="; las formulas (bloque C) si, y siguen siendo formulas
    for hoja, celda, antes, despues, bloque in cambios:
        if bloque != "C" and _es_formula(despues):
            raise MigracionDetenida(f"Texto que empieza con '=' en '{hoja}'!{celda}. No se escribio nada.")
        if bloque == "C" and not (_es_formula(antes) and _es_formula(despues)):
            raise MigracionDetenida(f"'{hoja}'!{celda} dejaria de ser formula. No se escribio nada.")

    if not cambios:
        return {"hubo_cambios": False, "cambios": ["Ya estaba aplicada: 0 celdas pendientes."],
                "detalle": []}

    # 3) Escribir
    for hoja, celda, antes, despues, bloque in cambios:
        wb[hoja][celda] = despues

    # 4) Volver a leer y confirmar
    for hoja, celda, antes, despues, bloque in cambios:
        if wb[hoja][celda].value != despues:
            raise MigracionDetenida(f"Verificacion fallo en '{hoja}'!{celda}")

    por_bloque = {}
    for c in cambios:
        por_bloque[c[4]] = por_bloque.get(c[4], 0) + 1
    resumen = [f"Bloque {b}: {n} celdas" for b, n in sorted(por_bloque.items())]
    # En el evento: B y D completos; A y C resumidos (son muchos y todos iguales)
    antes_despues = " ; ".join(f"'{h}'!{c}: {a!r} -> {d!r}" for h, c, a, d, b in cambios if b in "BD")

    from conexion_drive.tablas import registrar_evento
    registrar_evento(
        wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
        id_entidad=ID_MIGRACION, accion="Nombres unicos Baguette / CupCake (TABLA CRUCE MAESTRA)", estado="OK",
        detalle=(f"A) datos 'cupcakes'/'pan baguette' -> CupCake/Baguette: {por_bloque.get('A', 0)} celdas | "
                 f"B) hojas de costo A6 | C) flujo caja anual: {por_bloque.get('C', 0)} formulas, solo el texto "
                 f"buscado | D) etiquetas | Total: {len(cambios)} celdas | {antes_despues}"),
    )
    return {"hubo_cambios": True, "cambios": resumen, "detalle": cambios}


def _aplicar_a_archivo(ruta):
    import io
    from pathlib import Path
    import openpyxl
    from conexion_drive.recalculo import reparar_prefijos_xlfn, recalcular_bytes

    ruta = Path(ruta)
    contenido, _ = reparar_prefijos_xlfn(ruta.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    resultado = migrar(wb, None)
    if resultado["hubo_cambios"]:
        salida = io.BytesIO()
        wb.save(salida)
        ruta.write_bytes(recalcular_bytes(salida.getvalue()))
    return resultado


if __name__ == "__main__":
    try:
        if "--archivo" in sys.argv:
            ruta = sys.argv[sys.argv.index("--archivo") + 1]
            print(f"Excel de destino: ARCHIVO {ruta}")
            resultado = _aplicar_a_archivo(ruta)
        else:
            from conexion_drive import config
            from conexion_drive.drive_excel import editar_excel
            print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
            resultado = editar_excel(migrar, descripcion="Migracion m008_nombres_unicos")
    except MigracionDetenida as error:
        print(str(error))
        sys.exit(1)
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
