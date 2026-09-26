"""Migracion m010 - Descuento de los combos editable por la administracion (Pendiente #19).

Decision del usuario (26-sep-2026):
  - Si sube el precio de un producto externo (queso, natilla, leche, Coca Cola), el Precio de Oferta
    del combo SUBE aunque tenga descuento (refleja la realidad; no se compensa). Asi ya funcionaba:
    la formula lee el precio en vivo de la TABLA CRUCE MAESTRA. Resuelve el #19.
  - El numero del descuento lo DIGITA A MANO la administracion en 'Tabla Combos Dinamicos'
    columna L ("Descuento %"), por ejemplo 5 o 10.

Hallazgo: la columna L existia pero NINGUNA formula la leia; el 10% estaba escrito fijo (x 0.9)
dentro de cada formula de la columna E. Si la administracion cambiaba L, el precio no cambiaba.

Que cambia: SOLO el 0.9 fijo de E5:E8 por el descuento de su propia fila:
    0.9  ->  (1-IF($Lf>1,$Lf/100,$Lf))
Acepta que la administracion escriba 10 o 10% (Excel guarda 10% como 0.1). L vacia = sin descuento.
Todo lo demas de la formula queda IGUAL: las mismas celdas de la TABLA CRUCE MAESTRA, el mismo
producto externo a precio completo y el descuento solo sobre el pan / cupcake (diseño de la tabla).

No cambia la ruta de dependencias: E sigue en el mismo lugar y con el mismo significado; ninguna
otra hoja lee la columna E ('Ventas despacho' L9:L12 lee C5:C8 y 'Control de Inventario' I lee
TablaComposicionCombos). Con L = 10 los 4 precios quedan EXACTAMENTE iguales.

Seguridad (igual que m007): revisa TODO antes de escribir; se detiene sin cambiar nada si una
formula no es la esperada; se puede correr 2 veces; relee cada celda; evento MIGRACION.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m010_descuento_combos --local
    py -m migraciones.m010_descuento_combos
    py -m migraciones.m010_descuento_combos --archivo "ruta.xlsx"
"""
import os
import sys

if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

ID_MIGRACION = "m010_descuento_combos"
HOJA = "Tabla Combos Dinamicos"
TCM = "'TABLA CRUCE MAESTRA'!"


def _desc(f):
    return f"(1-IF($L{f}>1,$L{f}/100,$L{f}))"


CAMBIOS = [
    (HOJA, "E5", f"=(3*{TCM}G8)*0.9 +{TCM}G10", f"=(3*{TCM}G8)*{_desc(5)} +{TCM}G10"),   # Combo 1
    (HOJA, "E6", f"=(2*{TCM}G9)*0.9 +{TCM}G13", f"=(2*{TCM}G9)*{_desc(6)} +{TCM}G13"),   # Combo 2
    (HOJA, "E7", f"=(2*{TCM}G8)*0.9 +{TCM}G14", f"=(2*{TCM}G8)*{_desc(7)} +{TCM}G14"),   # Combo 3
    (HOJA, "E8", f"=(3*{TCM}G8)*0.9+{TCM}G11", f"=(3*{TCM}G8)*{_desc(8)}+{TCM}G11"),     # Combo 4
]
# La administracion debe tener un numero en L (hoy 10 en los 4 combos)
DESCUENTO_ACTUAL = {"L5": 10, "L6": 10, "L7": 10, "L8": 10}


class MigracionDetenida(Exception):
    """Una celda no tenia el valor esperado: no se escribio nada."""


def migrar(wb, wb_valores=None):
    ws = wb[HOJA]

    # 1) Revisar TODO antes de escribir
    problemas, pendientes = [], []
    for celda, esperado in DESCUENTO_ACTUAL.items():
        if ws[celda].value != esperado:
            problemas.append(f"'{HOJA}'!{celda}: se esperaba {esperado!r} y tiene {ws[celda].value!r}")
    for hoja, celda, antes, despues in CAMBIOS:
        valor = ws[celda].value
        if valor == despues:
            continue
        if valor != antes:
            problemas.append(f"'{hoja}'!{celda}: se esperaba {antes!r} y tiene {valor!r}")
        else:
            pendientes.append((hoja, celda, antes, despues))
    if problemas:
        raise MigracionDetenida("NO SE CAMBIO NADA. Celdas distintas a lo esperado:\n  " + "\n  ".join(problemas))
    if not pendientes:
        return {"hubo_cambios": False, "cambios": ["Ya estaba aplicada: 0 celdas pendientes."], "detalle": []}

    # 2) Escribir y 3) releer
    for hoja, celda, antes, despues in pendientes:
        ws[celda] = despues
    for hoja, celda, antes, despues in pendientes:
        if ws[celda].value != despues:
            raise MigracionDetenida(f"Verificacion fallo en '{hoja}'!{celda}")

    from conexion_drive.tablas import registrar_evento
    registrar_evento(
        wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
        id_entidad=ID_MIGRACION, accion="Descuento de combos editable en la columna L", estado="OK",
        detalle=(f"{HOJA} E5:E8: el 0.9 fijo pasa a leer el Descuento % de la columna L (acepta 10 o 10%). "
                 f"Precios iguales con L = 10. Total: {len(pendientes)} celdas | "
                 + " ; ".join(f"{c}: {a} -> {d}" for h, c, a, d in pendientes)),
    )
    return {"hubo_cambios": True, "cambios": [f"Combos: {len(pendientes)} formulas (E5:E8)"], "detalle": pendientes}


def _aplicar_a_archivo(ruta):
    import io
    from pathlib import Path
    import openpyxl
    from conexion_drive.recalculo import reparar_prefijos_xlfn, recalcular_bytes

    ruta = Path(ruta)
    contenido, _ = reparar_prefijos_xlfn(ruta.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    resultado = migrar(wb, None)
    if resultado["hubo_cambios"]:
        salida = io.BytesIO()
        wb.save(salida)
        ruta.write_bytes(recalcular_bytes(salida.getvalue()))
    return resultado


if __name__ == "__main__":
    try:
        if "--archivo" in sys.argv:
            ruta = sys.argv[sys.argv.index("--archivo") + 1]
            print(f"Excel de destino: ARCHIVO {ruta}")
            resultado = _aplicar_a_archivo(ruta)
        else:
            from conexion_drive import config
            from conexion_drive.drive_excel import editar_excel
            print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
            resultado = editar_excel(migrar, descripcion="Migracion m010_descuento_combos")
    except MigracionDetenida as error:
        print(str(error))
        sys.exit(1)
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
