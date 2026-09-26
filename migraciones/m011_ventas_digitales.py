"""Migracion m011 - Ventas digitales y envios separados en 'Flujo caja diario' (Pendiente #15).

Decisiones del usuario (26-sep-2026):
  - "Una cosa es venta digital y otra el cobro del envio." En 'Interface digital' TOTAL_A_COBRAR
    = comanda + delivery es CORRECTO (lo que paga el cliente) y NO cambia. En 'Flujo caja diario'
    se separan: "Ventas digitales" = solo la comanda (subtotal - descuento) y una columna NUEVA
    "Envios cobrados" = solo el delivery. El total de ingresos del dia no cambia de significado.
  - 'Detalle comanda digital' NO lleva envio (es por producto; el envio es del pedido completo).
  - Completar los pedidos de prueba con metodos de pago mezclados (Efectivo / SINPE Movil),
    mas cerca de la realidad. Todo es estudio academico: se cuadra un sistema con datos supuestos.

Que hace:
  A) 'Interface digital' WA-102..WA-118 (filas 9-25): les faltaban TOTAL_A_COBRAR, METODO_PAGO y
     ESTADO, por eso el flujo no los veia. J = subtotal - descuento + envio (regla de
     pedidos_digitales.py), K = Efectivo / SINPE Movil alternados, M = CONFIRMADO. Valores fijos.
  B) WA-101 (2-abr-2026, Jueves Santo = feriado): se alinea con su linea de 'Detalle comanda
     digital' (10 Baguette): DETALLE_COMANDA '10 Baguette', subtotal 10 x 842.69 = 8,426.90,
     descuento de feriado 5% ('Interface digital' B5) = 421.35, envio 500, total 8,505.55.
     En 'Detalle comanda digital' fila 4: subtotal de linea 8,005.55, detonador FERIADO, 0.05.
  C) 'Flujo caja diario' (tblFlujoCajaDiario):
     - C "Ventas digitales" = SUMIFS(SUBTOTAL_BRUTO) - SUMIFS(DESCUENTO_TOTAL) (solo CONFIRMADO).
     - K NUEVA "Envios cobrados" = SUMIFS(ENVIO_MAPS) (solo CONFIRMADO). La tabla pasa a A4:K369.
     - D "Total ingresos dia" = B + C + K.
     - El bloque RESUMEN / TOTALES se corre 2 columnas (K3:L13 -> M3:N13), sin cambiar sus
       formulas (usan referencias estructuradas), y se agrega "Envios cobrados" en M14:N14.
  D) DiccionarioDatos N28: nota '|| m011' (se AGREGA texto, no se borra).

Seguridad (igual que m007): revisa TODO antes de escribir; se detiene sin cambiar nada si algo no
tiene el valor esperado; se puede correr 2 veces; relee; evento MIGRACION.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m011_ventas_digitales --local
    py -m migraciones.m011_ventas_digitales
    py -m migraciones.m011_ventas_digitales --archivo "ruta.xlsx"
"""
import os
import sys
from copy import copy

if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

ID_MIGRACION = "m011_ventas_digitales"
FECHA = "26-sep-2026"
INT, DET, FCD, DIC = "Interface digital", "Detalle comanda digital", "Flujo caja diario", "DiccionarioDatos"
TABLA_FCD = "tblFlujoCajaDiario"
PRIMERA, ULTIMA = 5, 368                      # filas con formulas en el flujo
ENVIOS = "Envíos cobrados"
PAGOS = ("Efectivo", "SINPE Móvil")

_ID = "'Interface digital'!"
_CRIT = f"{_ID}$B$8:$B$50000,$A{{r}},{_ID}$M$8:$M$50000,\"CONFIRMADO\""
C_ANTES = f"=IF($A{{r}}=\"\",0,SUMIFS({_ID}$J$8:$J$50000,{_CRIT}))"
C_DESPUES = f"=IF($A{{r}}=\"\",0,SUMIFS({_ID}$G$8:$G$50000,{_CRIT})-SUMIFS({_ID}$H$8:$H$50000,{_CRIT}))"
K_DESPUES = f"=IF($A{{r}}=\"\",0,SUMIFS({_ID}$I$8:$I$50000,{_CRIT}))"
D_ANTES = "=IF($A{r}=\"\",\"\",B{r}+C{r})"
D_DESPUES = "=IF($A{r}=\"\",\"\",B{r}+C{r}+K{r})"

# B) WA-101: celda -> (antes, despues)
WA101_INT = {"F8": ("10 Bollo Pan, 1 Natilla, 1 Queso", "10 Baguette"), "G8": (5200, 8426.9),
             "H8": (200, 421.35), "I8": (500, 500), "J8": (5500, 8505.55)}
WA101_DET = {"F4": (None, 8005.55), "H4": (None, "FERIADO"), "I4": (None, 0.05)}

NOTA_N28 = (f" || m011 ({FECHA}): 'Flujo caja diario' ya NO suma TOTAL_A_COBRAR: 'Ventas digitales' = "
            "SUBTOTAL_BRUTO - DESCUENTO_TOTAL (solo la comanda) y la columna nueva 'Envios cobrados' = "
            "ENVIO_MAPS (solo el delivery). TOTAL_A_COBRAR sigue siendo lo que paga el cliente (comanda + "
            "delivery). Decision del usuario: una cosa es venta digital y otra el cobro del envio.")


class MigracionDetenida(Exception):
    """Una celda no tenia el valor esperado: no se escribio nada."""


def _copiar_estilo(origen, destino):
    if origen.has_style:
        destino.font, destino.border, destino.fill = copy(origen.font), copy(origen.border), copy(origen.fill)
        destino.number_format = origen.number_format
        destino.alignment, destino.protection = copy(origen.alignment), copy(origen.protection)


def _revisar(wb):
    """Devuelve (pendientes, problemas, ya_aplicada_flujo)."""
    problemas, cambios = [], []
    wi, wd, wf = wb[INT], wb[DET], wb[FCD]

    # A) pedidos de prueba WA-102..WA-118
    for i, fila in enumerate(range(9, 26)):
        esperado = f"WA-{102 + i}"
        if wi[f"A{fila}"].value != esperado:
            problemas.append(f"'{INT}'!A{fila}: se esperaba {esperado} y tiene {wi[f'A{fila}'].value!r}")
            continue
        g, h, e = wi[f"G{fila}"].value, wi[f"H{fila}"].value, wi[f"I{fila}"].value
        if not all(isinstance(x, (int, float)) for x in (g, h, e)):
            problemas.append(f"'{INT}' fila {fila}: subtotal/descuento/envio no son numeros ({g!r}, {h!r}, {e!r})")
            continue
        for col, nuevo in (("J", round(g - h + e, 2)), ("K", PAGOS[i % 2]), ("M", "CONFIRMADO")):
            v = wi[f"{col}{fila}"].value
            if v == nuevo:
                continue
            if v is not None:
                problemas.append(f"'{INT}'!{col}{fila}: se esperaba vacia y tiene {v!r}")
            else:
                cambios.append((INT, f"{col}{fila}", None, nuevo, "A"))

    # B) WA-101
    if wi["A8"].value != "WA-101" or wd["A4"].value != "WA-101" or wd["C4"].value != "Baguette" or wd["D4"].value != 10:
        problemas.append("WA-101 no esta en 'Interface digital' fila 8 / 'Detalle comanda digital' fila 4 (10 Baguette)")
    for hoja, mapa in ((INT, WA101_INT), (DET, WA101_DET)):
        for celda, (antes, despues) in mapa.items():
            v = wb[hoja][celda].value
            if v == despues:
                continue
            if v != antes:
                problemas.append(f"'{hoja}'!{celda}: se esperaba {antes!r} y tiene {v!r}")
            else:
                cambios.append((hoja, celda, antes, despues, "B"))

    # C) flujo caja diario
    tabla = wf.tables[TABLA_FCD]
    ya = tabla.ref == "A4:K369"
    if not ya:
        if tabla.ref != "A4:J369":
            problemas.append(f"'{FCD}': la tabla deberia ser A4:J369 y es {tabla.ref}")
        for celda, esperado in (("K3", "RESUMEN / TOTALES (Tabla Microsoft - referencias estructuradas)"),
                                ("K4", "Concepto"), ("L4", "Total"), ("K13", "Saldo de caja actual")):
            if wf[celda].value != esperado:
                problemas.append(f"'{FCD}'!{celda}: se esperaba {esperado!r} y tiene {wf[celda].value!r}")
        for celda in ("M3", "N3", "M4", "N14", "M14", "K369"):
            if wf[celda].value is not None:
                problemas.append(f"'{FCD}'!{celda} deberia estar vacia y tiene {wf[celda].value!r}")
    for r in range(PRIMERA, ULTIMA + 1):
        for col, antes, despues in (("C", C_ANTES, C_DESPUES), ("D", D_ANTES, D_DESPUES)):
            v = wf[f"{col}{r}"].value
            a, d = antes.format(r=r), despues.format(r=r)
            if v == d:
                continue
            if v != a:
                problemas.append(f"'{FCD}'!{col}{r}: formula distinta a la esperada ({v!r})")
            else:
                cambios.append((FCD, f"{col}{r}", a, d, "C"))

    # D) nota
    n28 = wb[DIC]["N28"].value or ""
    if not n28.endswith(NOTA_N28):
        cambios.append((DIC, "N28", n28, n28 + NOTA_N28, "D"))
    return cambios, problemas, ya


def _mover_resumen_y_agregar_columna(wf):
    # 1) Correr el RESUMEN K3:L13 -> M3:N13 (valores + estilos), y su celda combinada K3:L3 -> M3:N3
    wf.unmerge_cells("K3:L3")
    for r in range(3, 14):
        for c_origen, c_dest in (("K", "M"), ("L", "N")):
            o, d = wf[f"{c_origen}{r}"], wf[f"{c_dest}{r}"]
            d.value = o.value
            _copiar_estilo(o, d)
            o.value = None
    wf.merge_cells("M3:N3")
    wf.column_dimensions["M"].width = wf.column_dimensions["K"].width or 30
    wf.column_dimensions["N"].width = wf.column_dimensions["L"].width or 18
    wf.column_dimensions["K"].width = wf.column_dimensions["J"].width or 18
    # 2) Fila nueva del resumen
    wf["M14"] = "Envíos cobrados (incluidos en Total ingresos)"
    wf["N14"] = f"=SUM({TABLA_FCD}[{ENVIOS}])"
    _copiar_estilo(wf["M13"], wf["M14"])
    _copiar_estilo(wf["N12"], wf["N14"])
    # 3) Columna K de la tabla: encabezado, formulas y estilo (copiado de J)
    for r in range(4, 370):
        _copiar_estilo(wf[f"J{r}"], wf[f"K{r}"])
        wf[f"K{r}"].value = None
    wf["K4"] = ENVIOS
    for r in range(PRIMERA, ULTIMA + 1):
        wf[f"K{r}"] = K_DESPUES.format(r=r)
    from openpyxl.worksheet.table import TableColumn
    tabla = wf.tables[TABLA_FCD]
    tabla.tableColumns.append(TableColumn(id=11, name=ENVIOS))
    tabla.ref = "A4:K369"
    if tabla.autoFilter is not None:
        tabla.autoFilter.ref = "A4:K369"


def migrar(wb, wb_valores=None):
    cambios, problemas, ya_flujo = _revisar(wb)
    if problemas:
        raise MigracionDetenida("NO SE CAMBIO NADA. Celdas distintas a lo esperado:\n  " + "\n  ".join(problemas[:30]))
    for hoja, celda, antes, despues, bloque in cambios:
        if bloque in "ABD" and isinstance(despues, str) and despues.startswith("="):
            raise MigracionDetenida(f"Texto que empieza con '=' en '{hoja}'!{celda}. No se escribio nada.")
    if not cambios and ya_flujo:
        return {"hubo_cambios": False, "cambios": ["Ya estaba aplicada: 0 celdas pendientes."], "detalle": []}

    wi, wd, wf = wb[INT], wb[DET], wb[FCD]
    # estilos de las celdas nuevas de Interface (copiados de la fila 8)
    for hoja, celda, antes, despues, bloque in cambios:
        if hoja == INT and bloque == "A":
            _copiar_estilo(wi[f"{celda[0]}8"], wi[celda])
        wb[hoja][celda] = despues
    if not ya_flujo:
        _mover_resumen_y_agregar_columna(wf)

    for hoja, celda, antes, despues, bloque in cambios:
        if wb[hoja][celda].value != despues:
            raise MigracionDetenida(f"Verificacion fallo en '{hoja}'!{celda}")
    if wf.tables[TABLA_FCD].ref != "A4:K369" or wf["K4"].value != ENVIOS or wf["M5"].value != "Ventas despacho":
        raise MigracionDetenida("Verificacion fallo en la columna nueva o en el resumen del flujo")

    por_bloque = {}
    for c in cambios:
        por_bloque[c[4]] = por_bloque.get(c[4], 0) + 1
    resumen = [f"Bloque {b}: {n} celdas" for b, n in sorted(por_bloque.items())]
    if not ya_flujo:
        resumen.append(f"Flujo caja diario: columna nueva K '{ENVIOS}' ({ULTIMA - PRIMERA + 1} formulas), "
                       "resumen movido a M3:N14")
    from conexion_drive.tablas import registrar_evento
    registrar_evento(
        wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
        id_entidad=ID_MIGRACION, accion="Ventas digitales y envios separados en Flujo caja diario", estado="OK",
        detalle=(f"A) WA-102..118: TOTAL_A_COBRAR, METODO_PAGO (Efectivo/SINPE Movil), ESTADO | B) WA-101 alineado "
                 f"(10 Baguette, feriado 5%) | C) Flujo: C = subtotal - descuento, K nueva '{ENVIOS}', D = B+C+K, "
                 f"resumen K3:L13 -> M3:N13 + M14:N14 | D) nota N28 | {resumen}"),
    )
    return {"hubo_cambios": True, "cambios": resumen, "detalle": cambios}


def _aplicar_a_archivo(ruta):
    import io
    from pathlib import Path
    import openpyxl
    from conexion_drive.recalculo import reparar_prefijos_xlfn, recalcular_bytes

    ruta = Path(ruta)
    contenido, _ = reparar_prefijos_xlfn(ruta.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    resultado = migrar(wb, None)
    if resultado["hubo_cambios"]:
        salida = io.BytesIO()
        wb.save(salida)
        ruta.write_bytes(recalcular_bytes(salida.getvalue()))
    return resultado


if __name__ == "__main__":
    try:
        if "--archivo" in sys.argv:
            ruta = sys.argv[sys.argv.index("--archivo") + 1]
            print(f"Excel de destino: ARCHIVO {ruta}")
            resultado = _aplicar_a_archivo(ruta)
        else:
            from conexion_drive import config
            from conexion_drive.drive_excel import editar_excel
            print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
            resultado = editar_excel(migrar, descripcion="Migracion m011_ventas_digitales")
    except MigracionDetenida as error:
        print(str(error))
        sys.exit(1)
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
