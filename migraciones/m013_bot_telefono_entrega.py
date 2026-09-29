"""Migracion m013 - Datos nuevos para el bot de WhatsApp real (Sesion A del bot, 29-sep-2026).

Decisiones del usuario (29-sep-2026):
  - #33 opcion B: se GUARDA el telefono del cliente. Columna nueva TELEFONO al final de
    TablaInterfaceDigital. CLIENTE se llena con el nombre de perfil de WhatsApp.
  - Recibo al entregar, opcion C: funcion compartida marcar_entregado(id_pedido) (bot hoy,
    SaaS despues). La hora de entrega va en una columna nueva ENTREGADO. La columna ESTADO
    NO se toca: 'Flujo caja diario' y 'Control de Inventario' suman solo "CONFIRMADO".
  - Datos nuevos opcion C: el numero de SINPE Movil de la panaderia vive en NODO INTEGRAL
    fila 22 (numero SUPUESTO, estudio academico, Regla #27). Los numeros autorizados para el
    comando ENTREGADO NO van al Excel (van en el .env).
  - El bot es solo en espanol (el #16 queda sin efecto).

Cambios (bloques):
  A) 'Interface digital' TablaInterfaceDigital: columnas nuevas P 'TELEFONO' (texto) y
     Q 'ENTREGADO' (fecha y hora). El rango pasa de A7:O{n} a A7:Q{n}. Los pedidos existentes
     quedan con P y Q VACIAS (no se inventa nada). Estilo copiado de la columna O.
  B) 'NODO INTEGRAL' fila 22 (estaba vacia): Bot WhatsApp | Numero SINPE Movil | 8888-8888 |
     Manual (Dato del negocio) | explicacion. Estilo copiado de la fila 21.
  C) DiccionarioDatos: 3 filas nuevas al final (OPS.WA.TELEFONO, OPS.WA.ENTREGADO,
     NODO.SINPE_MOVIL), igual que m004 y m006 (filas agregadas despues de la ultima).
  D) Evento MIGRACION en RegistroEventos.

No se quita ni se mueve nada. Ninguna formula existente cambia.
Esperado en comparar_libros: tabla A7:O25 -> A7:Q25; Interface digital 2 celdas (P7, Q7);
NODO INTEGRAL 5 celdas (A22:E22); DiccionarioDatos 42 celdas (3 filas x 14); RegistroEventos
1 fila; 0 valores calculados cambiados; 0 errores nuevos.

Revisiones ANTES de escribir (si una falla, se detiene y no escribe nada):
  - La tabla es A7:O{n}, su ultima columna es CODIGO_OFERTA y tiene 15 columnas.
  - Las columnas P y Q de 'Interface digital' estan vacias en TODA la hoja.
  - NODO INTEGRAL A22:E22 esta vacia y no existe ya la etiqueta 'Numero SINPE Movil'.
  - Las 3 filas nuevas del diccionario estan vacias y sus ID no existen ya.
  - Ningun texto nuevo empieza con '='.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m013_bot_telefono_entrega --local   (copia datos_prueba/maestro_prueba.xlsx)
    py -m migraciones.m013_bot_telefono_entrega           (Excel real de Google Drive)
Si ya estaba aplicada, responde "Ya estaba aplicada" y no escribe nada.
"""
import os
import sys
from copy import copy

# --local ANTES de importar conexion_drive.config (igual que las demas migraciones)
if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

FECHA = "29-sep-2026"
ID_MIGRACION = "m013_bot_telefono_entrega"
INT, NODO, DIC = "Interface digital", "NODO INTEGRAL", "DiccionarioDatos"
TABLA_INT = "TablaInterfaceDigital"
COL_TEL, COL_ENT = "TELEFONO", "ENTREGADO"
FORMATO_TEL, FORMATO_ENT = "@", "d-mmm hh:mm"
ETIQUETA_SINPE = "Número SINPE Móvil"
FILA_NODO = 22

NODO_FILA = {
    "A": "Bot WhatsApp",
    "B": ETIQUETA_SINPE,
    "C": "8888-8888",
    "D": "Manual (Dato del negocio)",
    "E": ("Lectura del bot: se muestra al cliente que elige pagar por SINPE Móvil. "
          f"Número SUPUESTO (estudio académico, Regla #27). CREADO m013 ({FECHA})."),
}

DIC_FILAS = [
    ["OPS.WA.TELEFONO", "Operaciones", INT, TABLA_INT, "P8 (TELEFONO)", "Texto (+506...)",
     "WhatsApp Bot (Twilio)", "Python", "Bot (recibo) / SaaS", "No", "No", "WhatsApp (Twilio)",
     "OPS.WA.ENTREGADO",
     f"CREADO m013 ({FECHA}, #33 opcion B): numero de WhatsApp del cliente, tal como lo envia "
     "Twilio. Se guarda solo para avisarle sobre su pedido (recibo al entregar); el cliente lo acepta "
     "en el mensaje de confirmacion. Dato personal: no se usa para publicidad. Los pedidos anteriores "
     "a la m013 quedan vacios."],
    ["OPS.WA.ENTREGADO", "Operaciones", INT, TABLA_INT, "Q8 (ENTREGADO)", "Fecha y hora",
     "Robot (marcar_entregado)", "Python", "SaaS / BI", "No", "No", "WhatsApp (Twilio)",
     "OPS.WA.TELEFONO",
     f"CREADO m013 ({FECHA}, recibo al entregar opcion C): fecha y hora en que el pedido se "
     "entrego. La escribe la funcion compartida marcar_entregado(id_pedido): hoy por el comando "
     "'ENTREGADO WA-xxx' de un numero autorizado (lista en el .env, no en el Excel) y despues por el "
     "boton de administracion del SaaS. Al marcarlo se envia el recibo al cliente. ESTADO no cambia "
     "(sigue CONFIRMADO) porque el flujo de caja y el inventario suman solo CONFIRMADO."],
    ["NODO.SINPE_MOVIL", "Bot WhatsApp", NODO, "(sin Tabla formal)", f"C{FILA_NODO} ({ETIQUETA_SINPE})",
     "Texto", "Manual (administracion)", "Usuario", "Bot WhatsApp", "No", "No", "—",
     "OPS.WA.METODO_PAGO",
     f"CREADO m013 ({FECHA}): numero de SINPE Movil de la panaderia que el bot muestra al cliente "
     "que elige pagar por SINPE Movil. En el estudio es un numero SUPUESTO (Regla #27). Se busca por "
     "su etiqueta en la columna B, no por numero de fila."],
]


class MigracionDetenida(Exception):
    """El Excel no esta como se esperaba: no se escribe nada."""


def _revisar(condicion, mensaje):
    if not condicion:
        raise MigracionDetenida(mensaje + " No se escribio nada.")


def _copiar_estilo(origen, destino):
    if origen.has_style:
        destino.font, destino.border, destino.fill = copy(origen.font), copy(origen.border), copy(origen.fill)
        destino.number_format = origen.number_format
        destino.alignment, destino.protection = copy(origen.alignment), copy(origen.protection)


def _partes_ref(ref):
    """'A7:O25' -> ('A', 7, 'O', 25)"""
    from openpyxl.utils.cell import coordinate_from_string
    inicio, fin = ref.split(":")
    c1, f1 = coordinate_from_string(inicio)
    c2, f2 = coordinate_from_string(fin)
    return c1, f1, c2, f2


def ya_aplicada(wb):
    ws = wb[INT]
    tabla = ws.tables[TABLA_INT]
    nombres = [c.name for c in tabla.tableColumns]
    return nombres[-2:] == [COL_TEL, COL_ENT] and wb[NODO][f"B{FILA_NODO}"].value == ETIQUETA_SINPE


def fila_diccionario(ws):
    """Primera fila vacia despues de la ultima con ID en la columna A."""
    ultima = max((c.row for c in ws["A"] if c.value not in (None, "")), default=4)
    return ultima + 1


def revisar_antes(wb):
    ws = wb[INT]
    _revisar(TABLA_INT in ws.tables, f"No existe {TABLA_INT} en '{INT}'.")
    tabla = ws.tables[TABLA_INT]
    c1, f1, c2, f2 = _partes_ref(tabla.ref)
    _revisar((c1, f1, c2) == ("A", 7, "O"), f"{TABLA_INT} deberia ser A7:O(n) y es {tabla.ref}.")
    nombres = [c.name for c in tabla.tableColumns]
    _revisar(len(nombres) == 15 and nombres[-1] == "CODIGO_OFERTA",
             f"{TABLA_INT} deberia tener 15 columnas terminando en CODIGO_OFERTA: {nombres}.")
    _revisar(ws["O7"].value == "CODIGO_OFERTA", f"'{INT}'!O7 no es CODIGO_OFERTA.")
    ocupadas = [c.coordinate for col in ("P", "Q") for c in ws[col] if c.value not in (None, "")]
    _revisar(not ocupadas, f"Las columnas P y Q de '{INT}' deberian estar vacias y tienen: {ocupadas[:10]}.")

    nodo = wb[NODO]
    ocupadas = [f"{c}{FILA_NODO}" for c in "ABCDE" if nodo[f"{c}{FILA_NODO}"].value not in (None, "")]
    _revisar(not ocupadas, f"NODO INTEGRAL fila {FILA_NODO} deberia estar vacia y tiene: {ocupadas}.")
    _revisar(nodo["B21"].value == "DATA OFERTA MUSI", "NODO INTEGRAL B21 no es 'DATA OFERTA MUSI'.")
    etiquetas = [str(c.value).strip() for c in nodo["B"] if c.value]
    _revisar(ETIQUETA_SINPE not in etiquetas, f"Ya existe la etiqueta '{ETIQUETA_SINPE}' en NODO INTEGRAL.")

    dic = wb[DIC]
    _revisar([c.value for c in dic[4]][:2] == ["ID variable", "Módulo"], "DiccionarioDatos fila 4 no es el encabezado.")
    ids = {str(c.value).strip() for c in dic["A"] if c.value}
    repetidos = [f[0] for f in DIC_FILAS if f[0] in ids]
    _revisar(not repetidos, f"Ya existen en DiccionarioDatos: {repetidos}.")
    inicio = fila_diccionario(dic)
    ocupadas = [c.coordinate for r in range(inicio, inicio + len(DIC_FILAS)) for c in dic[r] if c.value is not None]
    _revisar(not ocupadas, f"Las filas nuevas del diccionario no estan vacias: {ocupadas[:10]}.")

    textos = list(NODO_FILA.values()) + [t for fila in DIC_FILAS for t in fila] + [COL_TEL, COL_ENT]
    _revisar(not any(isinstance(t, str) and t.startswith("=") for t in textos), "Hay un texto que empieza con '='.")
    return {"ultima_fila_tabla": f2, "fila_dic": inicio}


def aplicar(wb, info):
    from openpyxl.worksheet.table import TableColumn

    ws = wb[INT]
    tabla = ws.tables[TABLA_INT]
    f2 = info["ultima_fila_tabla"]
    # A) columnas nuevas: encabezado + estilo de las filas de datos (valores vacios)
    for col, nombre, formato in (("P", COL_TEL, FORMATO_TEL), ("Q", COL_ENT, FORMATO_ENT)):
        _copiar_estilo(ws["O7"], ws[f"{col}7"])
        ws[f"{col}7"] = nombre
        for r in range(8, f2 + 1):
            _copiar_estilo(ws[f"O{r}"], ws[f"{col}{r}"])
            ws[f"{col}{r}"].number_format = formato
        ws.column_dimensions[col].width = 16 if col == "P" else 15
    siguiente_id = max(c.id for c in tabla.tableColumns) + 1
    tabla.tableColumns.append(TableColumn(id=siguiente_id, name=COL_TEL))
    tabla.tableColumns.append(TableColumn(id=siguiente_id + 1, name=COL_ENT))
    tabla.ref = f"A7:Q{f2}"
    if tabla.autoFilter is not None:
        tabla.autoFilter.ref = tabla.ref

    # B) NODO INTEGRAL fila 22
    nodo = wb[NODO]
    for col, valor in NODO_FILA.items():
        _copiar_estilo(nodo[f"{col}21"], nodo[f"{col}{FILA_NODO}"])
        nodo[f"{col}{FILA_NODO}"] = valor
    nodo[f"C{FILA_NODO}"].number_format = "@"

    # C) DiccionarioDatos: filas nuevas al final (estilo de la ultima fila con datos)
    dic = wb[DIC]
    modelo = info["fila_dic"] - 1
    for i, valores in enumerate(DIC_FILAS):
        r = info["fila_dic"] + i
        for j, valor in enumerate(valores, start=1):
            _copiar_estilo(dic.cell(modelo, j), dic.cell(r, j))
            dic.cell(r, j).value = valor
    return {"A": 2, "B": len(NODO_FILA), "C": len(DIC_FILAS) * len(DIC_FILAS[0])}


def verificar(wb, info):
    ws = wb[INT]
    tabla = ws.tables[TABLA_INT]
    _revisar(tabla.ref == f"A7:Q{info['ultima_fila_tabla']}", f"Verificacion: la tabla quedo {tabla.ref}.")
    _revisar([c.name for c in tabla.tableColumns][-2:] == [COL_TEL, COL_ENT], "Verificacion: columnas nuevas.")
    _revisar(ws["P7"].value == COL_TEL and ws["Q7"].value == COL_ENT, "Verificacion: encabezados P7/Q7.")
    _revisar(wb[NODO][f"B{FILA_NODO}"].value == ETIQUETA_SINPE, "Verificacion: NODO INTEGRAL fila 22.")
    _revisar(wb[DIC].cell(info["fila_dic"], 1).value == DIC_FILAS[0][0], "Verificacion: DiccionarioDatos.")


def main():
    from conexion_drive import config
    from conexion_drive.drive_excel import editar_excel
    from conexion_drive.tablas import registrar_evento

    destino = ("COPIA LOCAL " + str(config.EXCEL_LOCAL_PATH)) if config.ORIGEN_EXCEL == "local" \
        else "GOOGLE DRIVE (real)"
    print(f"=== {ID_MIGRACION} -> {destino} ===")

    def editar(wb, wb_valores):
        if ya_aplicada(wb):
            return {"hubo_cambios": False, "estado": "YA_APLICADA"}
        info = revisar_antes(wb)
        conteo = aplicar(wb, info)
        verificar(wb, info)
        registrar_evento(
            wb, modulo="Migraciones", tipo_evento="MIGRACION", entidad=ID_MIGRACION, id_entidad="m013",
            accion="Bot WhatsApp: columnas TELEFONO y ENTREGADO, SINPE en NODO fila 22, diccionario",
            estado="OK",
            detalle=(f"A) {TABLA_INT} A7:O{info['ultima_fila_tabla']} -> A7:Q{info['ultima_fila_tabla']} "
                     f"(P {COL_TEL}, Q {COL_ENT}, vacias en los pedidos existentes) | B) NODO INTEGRAL "
                     f"A22:E22 SINPE supuesto | C) DiccionarioDatos filas {info['fila_dic']}-"
                     f"{info['fila_dic'] + len(DIC_FILAS) - 1} | #33 opcion B, recibo opcion C"),
            origen="Robot Python",
        )
        return {"hubo_cambios": True, "estado": "APLICADA", "conteo": conteo, "info": info}

    try:
        resultado = editar_excel(editar, descripcion=f"Migracion {ID_MIGRACION}")
    except MigracionDetenida as error:
        print(f"DETENIDA: {error}")
        sys.exit(1)
    if resultado["estado"] == "YA_APLICADA":
        print("Ya estaba aplicada. No se escribio nada.")
        return
    info = resultado["info"]
    print(f"  Bloque A: {TABLA_INT} A7:O{info['ultima_fila_tabla']} -> A7:Q{info['ultima_fila_tabla']} "
          f"(P {COL_TEL}, Q {COL_ENT})")
    print(f"  Bloque B: NODO INTEGRAL A22:E22 ({ETIQUETA_SINPE} = {NODO_FILA['C']})")
    print(f"  Bloque C: DiccionarioDatos filas {info['fila_dic']} a {info['fila_dic'] + len(DIC_FILAS) - 1}")
    for bloque, n in resultado["conteo"].items():
        print(f"  Bloque {bloque}: {n} celdas")
    print("LISTO.")


if __name__ == "__main__":
    main()
