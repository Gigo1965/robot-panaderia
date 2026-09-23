"""Migracion 002 — consumo de materia prima al hornear (se corre UNA sola vez, despues de m001).

Que cambia (decisiones confirmadas por el usuario, 23-sep-2026):

  1. 'TABLA CRUCE MAESTRA MP': se corrigen 2 conversiones de costo por gramo
        E10 Mejorante 3 KILOS: D10/1000 -> D10/3000
        E19 Cafe 250 GR x 3:   D19/2500 -> D19/750
     (el costo del mejorante estaba 3 veces mas alto; baja el costo de la Baguette)

  2. Hoja NUEVA 'Consumo de Materia Prima':
     a) TablaRecetaConsumo (A4:F16): que gasta cada unidad horneada. Todo son
        FORMULAS enlazadas a las recetas de 'costo pan baguette' / 'costo cupcakes'
        y a los nombres de las tablas maestras: si la receta cambia, esto la sigue.
     b) TablaConsumoMP (A19:I...): la lista de SALIDAS por produccion (espejo de
        'Ingreso de Inventario'). La llena el robot en cada hornada con valores fijos;
        al anular una hornada escribe las mismas filas en negativo.

  3. 'Control de Inventario': columna J 'Consumo en Produccion' y
     Existencia = Entradas - Ventas - Combos - Consumo en Produccion.
     'Ingreso de Inventario' y 'Flujo caja diario' NO se tocan (la compra ya salio
     de caja el dia que se compro; hornear no mueve caja).

  4. 'costo cupcakes'!A12: la etiqueta decia '41gr x unid', la receta usa 25 g
     (confirmado por el usuario): se corrige solo el texto.

  5. 'flujo caja proyecto' B9 / C9 (precio de venta de la proyeccion) pasan a leer el
     PRECIO DE REFERENCIA ('costo pan baguette'!G28 / 'costo cupcakes'!G27) en vez del
     precio calculado por costo x margen (G29 / G28). Confirmado por el usuario: los
     INGRESOS usan el precio de referencia; los COSTOS (B11 / C11) siguen siendo el costo
     real dinamico. De B9 / C9 dependen Punto de Equilibrio, TIR VAR y las proyecciones.

Es SEGURA de repetir: si ya se aplico, no cambia nada (cada paso revisa si ya se hizo,
asi que si se agrega un paso nuevo, al volver a correrla solo se aplica ese paso).

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m002_consumo_mp --local    <- primero, sobre la copia de prueba
    py -m migraciones.m002_consumo_mp            <- despues, sobre el Excel real en Drive
"""
import os
import sys

if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

import logging                                                    # noqa: E402
from copy import copy                                             # noqa: E402

from openpyxl.styles import Alignment, Font, PatternFill          # noqa: E402
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo  # noqa: E402

from conexion_drive import config, editar_excel                   # noqa: E402
from conexion_drive.tablas import registrar_evento                # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")

HOJA = "Consumo de Materia Prima"
TABLA_RECETA, TABLA_CONSUMO = "TablaRecetaConsumo", "TablaConsumoMP"
MP = "'TABLA CRUCE MAESTRA MP'"
TCM = "'TABLA CRUCE MAESTRA'"

CONVERSIONES = {"E10": ("=D10/1000", "=D10/3000"), "E19": ("=D19/2500", "=D19/750")}

# (producto horneado, codigo, formula del nombre de la materia prima, formula de la cantidad por unidad,
#  unidad, formula del contenido por paquete)
RECETA = [
    ("Baguette", "P01", f"={MP}!$C$11", "='costo pan baguette'!$I$14", "g", f"={MP}!$D$11/{MP}!$E$11"),
    ("Baguette", "P01", f"={MP}!$C$14", "='costo pan baguette'!$I$17", "g", f"={MP}!$D$14/{MP}!$E$14"),
    ("Baguette", "P01", f"={MP}!$C$10", "='costo pan baguette'!$I$21", "g", f"={MP}!$D$10/{MP}!$E$10"),
    ("Baguette", "P01", f"={MP}!$C$13", "='costo pan baguette'!$I$24", "g", f"={MP}!$D$13/{MP}!$E$13"),
    ("Baguette", "P01", f"={MP}!$C$15", "='costo pan baguette'!$I$12", "g", f"={MP}!$D$15/{MP}!$E$15"),
    ("Baguette", "P01", f"={MP}!$C$17", 1, "und", f"={MP}!$D$17/{MP}!$E$17"),
    ("CupCake", "Pa01", f"={MP}!$C$11", "='costo cupcakes'!$I$14", "g", f"={MP}!$D$11/{MP}!$E$11"),
    ("CupCake", "Pa01", f"={MP}!$C$20", "='costo cupcakes'!$I$17", "g", f"={MP}!$D$20/{MP}!$E$20"),
    ("CupCake", "Pa01", f"={MP}!$C$12", "='costo cupcakes'!$I$19", "g", f"={MP}!$D$12/{MP}!$E$12"),
    ("CupCake", "Pa01", f"={TCM}!$C$10", "='costo cupcakes'!$I$21", "ml", 1000),
    ("CupCake", "Pa01", f"={MP}!$C$15", "='costo cupcakes'!$I$15", "g", f"={MP}!$D$15/{MP}!$E$15"),
    ("CupCake", "Pa01", f"={MP}!$C$18", 0.5, "und", f"={MP}!$D$18/{MP}!$E$18"),
]
ENC_RECETA = ["Producto horneado", "Codigo", "Materia prima", "Cantidad por unidad", "Unidad",
              "Contenido por paquete"]
ENC_CONSUMO = ["Fecha", "Lote", "Producto horneado", "Cantidad horneada", "Materia prima",
               "Consumo por unidad", "Consumo total", "Contenido por paquete", "Paquetes consumidos"]
FILA_RECETA, FILA_CONSUMO = 4, 19

NEGRITA = Font(name="Arial", size=10, bold=True)
NORMAL = Font(name="Arial", size=10)
ENCABEZADO = PatternFill("solid", fgColor="FFD9E8D2")


def paso_1_conversiones(wb, cambios):
    ws = wb["TABLA CRUCE MAESTRA MP"]
    hechas = []
    for celda, (antes, despues) in CONVERSIONES.items():
        actual = ws[celda].value
        if actual == despues:
            continue
        if actual != antes:
            raise RuntimeError(f"{MP}!{celda} tiene '{actual}', se esperaba '{antes}'. No se cambia.")
        ws[celda] = despues
        hechas.append(f"{celda}: {antes} -> {despues}")
    cambios.append("1) Conversiones MP: " + (", ".join(hechas) if hechas else "ya corregidas, sin cambios"))
    return bool(hechas)


def _encabezado(ws, fila, textos):
    for j, texto in enumerate(textos, start=1):
        c = ws.cell(fila, j, texto)
        c.font, c.fill = NEGRITA, ENCABEZADO
        c.alignment = Alignment(wrap_text=True, vertical="top")


def paso_2_hoja_consumo(wb, cambios):
    if HOJA in wb.sheetnames:
        cambios.append(f"2) Hoja '{HOJA}': ya existia, sin cambios")
        return False
    posicion = wb.sheetnames.index("Control de Inventario") + 1
    ws = wb.create_sheet(HOJA, index=posicion)
    ws["A1"] = "CONSUMO DE MATERIA PRIMA (salidas de inventario por produccion)"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = ("Arriba: la receta por unidad (formulas enlazadas a las hojas de costo — no digitar). "
                "Abajo: cada hornada registrada por el robot. No se digita a mano.")
    ws["A2"].font = NORMAL

    _encabezado(ws, FILA_RECETA, ENC_RECETA)
    for n, (producto, codigo, nombre, cantidad, unidad, contenido) in enumerate(RECETA, start=1):
        r = FILA_RECETA + n
        for j, valor in enumerate((producto, codigo, nombre, cantidad, unidad, contenido), start=1):
            ws.cell(r, j, valor).font = NORMAL
        ws.cell(r, 6).number_format = "#,##0"
    fin_receta = FILA_RECETA + len(RECETA)
    t1 = Table(displayName=TABLA_RECETA, ref=f"A{FILA_RECETA}:F{fin_receta}")
    t1.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t1)

    ws.cell(FILA_CONSUMO - 1, 1, "Consumo por hornada (lo escribe el robot)").font = NEGRITA
    _encabezado(ws, FILA_CONSUMO, ENC_CONSUMO)
    for j in range(1, len(ENC_CONSUMO) + 1):            # fila de datos vacia con formato (la tabla necesita 1)
        c = ws.cell(FILA_CONSUMO + 1, j)
        c.font = NORMAL
    ws.cell(FILA_CONSUMO + 1, 1).number_format = "dd/mm/yyyy"
    ws.cell(FILA_CONSUMO + 1, 9).number_format = "0.0000"
    t2 = Table(displayName=TABLA_CONSUMO, ref=f"A{FILA_CONSUMO}:I{FILA_CONSUMO + 1}")
    t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t2)

    for letra, ancho in zip("ABCDEFGHI", (18, 22, 46, 16, 46, 16, 16, 18, 18)):
        ws.column_dimensions[letra].width = ancho
    cambios.append(f"2) Hoja '{HOJA}' creada: {TABLA_RECETA} (A{FILA_RECETA}:F{fin_receta}) "
                   f"y {TABLA_CONSUMO} (A{FILA_CONSUMO}:I{FILA_CONSUMO + 1})")
    return True


def paso_3_control_inventario(wb, cambios):
    ws = wb["Control de Inventario"]
    tabla = ws.tables["TablaControlInventario"]
    nombres = [c.name for c in tabla.tableColumns]
    if "Consumo en Producción" in nombres:
        cambios.append("3) Control de Inventario: columna 'Consumo en Producción' ya existia, sin cambios")
        return False
    if "Salidas en Combos" not in nombres:
        raise RuntimeError("Primero hay que aplicar la migracion m001_grupo2.")
    ultima = int(tabla.ref.split(":")[1][1:])
    ws["J4"] = "Consumo en Producción"
    for attr in ("font", "border", "alignment", "fill"):
        setattr(ws["J4"], attr, copy(getattr(ws["I4"], attr)))
    ws.column_dimensions["J"].width = max(ws.column_dimensions["I"].width or 18, 20)
    for r in range(5, ultima + 1):
        if ws[f"A{r}"].value in (None, ""):
            continue
        ws[f"J{r}"] = f"=SUMIF({TABLA_CONSUMO}[Materia prima],A{r},{TABLA_CONSUMO}[Paquetes consumidos])"
        ws[f"J{r}"].font = copy(ws[f"I{r}"].font)
        ws[f"J{r}"].border = copy(ws[f"I{r}"].border)
        ws[f"J{r}"].number_format = "0.00"
        ws[f"F{r}"] = f"=B{r}+C{r}-D{r}-E{r}-I{r}-J{r}"
    tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name="Consumo en Producción"))
    tabla.ref = f"A4:J{ultima}"
    if tabla.autoFilter is not None:
        tabla.autoFilter.ref = tabla.ref
    cambios.append(f"3) Control de Inventario: columna J 'Consumo en Producción' + Existencia actualizada")
    return True


def paso_4_etiqueta(wb, cambios):
    ws = wb["costo cupcakes"]
    actual = str(ws["A12"].value or "")
    if "41gr" not in actual:
        cambios.append("4) Etiqueta harina CupCake: ya corregida, sin cambios")
        return False
    ws["A12"] = actual.replace("41gr", "25gr")
    cambios.append("4) 'costo cupcakes'!A12: etiqueta '41gr x unid' -> '25gr x unid'")
    return True


PRECIO_PROYECCION = {
    "B9": ("=ROUND('costo pan baguette'!$G$29,2)", "=ROUND('costo pan baguette'!$G$28,2)"),
    "C9": ("=ROUND('costo cupcakes'!$G$28,2)", "=ROUND('costo cupcakes'!$G$27,2)"),
}


def paso_5_precio_proyeccion(wb, cambios):
    ws = wb["flujo caja proyecto"]
    hechas = []
    for celda, (antes, despues) in PRECIO_PROYECCION.items():
        actual = ws[celda].value
        if actual == despues:
            continue
        if actual != antes:
            raise RuntimeError(f"'flujo caja proyecto'!{celda} tiene '{actual}', se esperaba '{antes}'. No se cambia.")
        ws[celda] = despues
        hechas.append(celda)
    cambios.append("5) flujo caja proyecto: " + (f"{', '.join(hechas)} ahora leen el precio de referencia"
                                                 if hechas else "precio de referencia ya aplicado, sin cambios"))
    return bool(hechas)


def migrar(wb, wb_valores):
    cambios = []
    hubo = [paso_1_conversiones(wb, cambios), paso_2_hoja_consumo(wb, cambios),
            paso_3_control_inventario(wb, cambios), paso_4_etiqueta(wb, cambios),
            paso_5_precio_proyeccion(wb, cambios)]
    if not any(hubo):
        return {"hubo_cambios": False, "cambios": cambios}
    registrar_evento(wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
                     id_entidad="m002_consumo_mp", accion="Consumo de materia prima", estado="OK",
                     detalle=" | ".join(cambios))
    return {"hubo_cambios": True, "cambios": cambios}


if __name__ == "__main__":
    print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
    resultado = editar_excel(migrar, descripcion="Migracion m002_consumo_mp")
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
