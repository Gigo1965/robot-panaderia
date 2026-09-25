"""Comparacion INDEPENDIENTE de dos copias del Excel maestro (segunda revision del protocolo).

No usa el codigo de ninguna migracion: abre los dos archivos y compara TODAS las hojas,
celda por celda, en formulas y en valores calculados. Sirve para confirmar que una migracion
cambio SOLO lo que anuncio.

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.comparar_libros "datos_prueba/maestro_prueba_ANTES_m005.xlsx" "datos_prueba/maestro_prueba.xlsx"

Que muestra:
  * Hojas agregadas o quitadas, y el rango de cada tabla si cambio.
  * Por hoja: celdas con formula/contenido distinto (antes -> despues).
  * Celdas con resultado de ERROR (#REF!, #DIV/0!, #NAME?...) en el libro nuevo que no estaban antes.
  * Resumen: cuantas formulas quedaron iguales y cuantos cambios hubo por hoja.
Los cambios de valor calculado se cuentan aparte (son la cascada normal de una formula corregida).
"""
import sys
from pathlib import Path

from openpyxl import load_workbook

ERRORES = ("#REF!", "#DIV/0!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
MAX_LINEAS = 60   # por hoja, para no llenar la pantalla


def _texto(valor):
    """Las formulas matriciales (ej. NODO C19/C20) se comparan por su texto y su rango."""
    if hasattr(valor, "text") and hasattr(valor, "ref"):
        ref = str(valor.ref)
        if ":" in ref and ref.split(":")[0] == ref.split(":")[1]:
            ref = ref.split(":")[0]          # "C19:C19" es lo mismo que "C19"
        return f"{{matricial {ref}}} {valor.text}"
    return valor


def _celdas(ws):
    return {c.coordinate: _texto(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None}


def _corto(valor, largo=70):
    texto = repr(valor)
    return texto if len(texto) <= largo else texto[: largo - 3] + "..."


def comparar(ruta_antes, ruta_despues):
    f_antes, f_desp = load_workbook(ruta_antes), load_workbook(ruta_despues)
    v_antes = load_workbook(ruta_antes, data_only=True)
    v_desp = load_workbook(ruta_despues, data_only=True)

    hojas_a, hojas_d = f_antes.sheetnames, f_desp.sheetnames
    print(f"Hojas: antes {len(hojas_a)}, despues {len(hojas_d)}")
    for h in set(hojas_a) ^ set(hojas_d):
        print(f"  HOJA {'QUITADA' if h in hojas_a else 'NUEVA'}: {h}")

    total_iguales, total_cambios, total_valores, errores_nuevos = 0, 0, 0, []
    for hoja in hojas_d:
        if hoja not in hojas_a:
            continue
        wa, wd = f_antes[hoja], f_desp[hoja]
        for nombre in wd.tables:
            ref_nueva = wd.tables[nombre].ref
            ref_antes = wa.tables[nombre].ref if nombre in wa.tables else None
            if ref_antes != ref_nueva:
                print(f"  [{hoja}] tabla {nombre}: {ref_antes} -> {ref_nueva}")

        ca, cd = _celdas(wa), _celdas(wd)
        cambios, cambiadas = [], set()
        for celda in sorted(set(ca) | set(cd), key=lambda c: (len(c), c)):
            a, d = ca.get(celda), cd.get(celda)
            if a == d:
                if isinstance(d, str) and d.startswith("="):
                    total_iguales += 1
            else:
                cambios.append(f"    {celda}: {_corto(a)} -> {_corto(d)}")
                cambiadas.add(celda)

        va_hoja, vd_hoja = _celdas(v_antes[hoja]), _celdas(v_desp[hoja])
        n_valores = 0
        for celda, valor in vd_hoja.items():
            if valor != va_hoja.get(celda) and celda not in cambiadas:
                n_valores += 1
            if isinstance(valor, str) and valor in ERRORES and va_hoja.get(celda) not in ERRORES:
                errores_nuevos.append(f"'{hoja}'!{celda} = {valor}")
        total_valores += n_valores

        if cambios:
            total_cambios += len(cambios)
            print(f"\n[{hoja}] {len(cambios)} celda(s) con contenido distinto:")
            for linea in cambios[:MAX_LINEAS]:
                print(linea)
            if len(cambios) > MAX_LINEAS:
                print(f"    ... y {len(cambios) - MAX_LINEAS} mas")

    print("\n================ RESUMEN ================")
    print(f"Formulas iguales:                    {total_iguales:,}")
    print(f"Celdas con contenido distinto:       {total_cambios:,}")
    print(f"Valores calculados que cambiaron:    {total_valores:,}  (cascada de las formulas corregidas)")
    print(f"Errores NUEVOS de formula:           {len(errores_nuevos)}")
    for e in errores_nuevos[:20]:
        print("    " + e)
    print("TODO BIEN: 0 errores nuevos." if not errores_nuevos else "REVISAR: hay errores nuevos.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    antes, despues = Path(sys.argv[1]), Path(sys.argv[2])
    for ruta in (antes, despues):
        if not ruta.exists():
            print(f"No existe: {ruta}")
            sys.exit(1)
    comparar(antes, despues)
