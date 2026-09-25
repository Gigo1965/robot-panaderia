"""Migracion m004 - Grupo 3C: celdas de la competencia (Musi; Musmanni desactivado).

Actualizacion 25-sep-2026: solo se monitorea a Musi (antes Musmanni Mini Super; ambas
marcas de Heineken Costa Rica desde ene-2026). Musmanni queda desactivado en competencia.py;
sus celdas (NODO C11, Tabla Combos D10/E10) no se tocan. Las notas C lo documentan.

Que hace (diseno aprobado por el usuario el 25-sep-2026):
  A) NODO INTEGRAL: B11 pasa a "DATA OFERTA MUSMANNI" y se crea la fila 21 para Musi
     (A21:E21, mismo formato que la fila 11). Se usa la fila 21 (primera libre) porque
     insertar una fila correria C12:C20, que usan formulas y nodo_integral.py.
  B) Tabla Combos Dinamicos: fila 11 para Musi (C11:E11, mismo formato que la fila 10).
     Las filas 10 y 11 NO son combos propios: las formulas de NODO C19/C20 solo leen
     las filas 5 a 8. Sirven para ver la oferta de la competencia junto a los combos.
  C) DiccionarioDatos: notas "|| m004" en L252, N252 y N89 (se AGREGA texto, no se borra).
  D) DiccionarioDatos: filas nuevas 303 (NODO.PRECIOS_COMPETENCIA_MUSI) y
     304 (MKT.COMPETENCIA.MUSI), con el formato de la fila 252.

Seguridad (protocolo del proyecto, igual que la m003):
  * Antes de escribir revisa TODAS las celdas. Si alguna no tiene el valor
    "antes" esperado (ni el "despues"), se detiene SIN cambiar nada.
  * Si una celda ya tiene el valor "despues", la salta (se puede correr 2 veces).
  * Despues de escribir, vuelve a leer cada celda y confirma el valor.
  * Ningun texto escrito empieza con "=".

Uso (desde la carpeta ROBOT PANADERO):
    py -m migraciones.m004_competencia --local            (copia datos_prueba)
    py -m migraciones.m004_competencia                    (Excel real en Drive)
    py -m migraciones.m004_competencia --archivo "ruta.xlsx"   (un archivo suelto)
"""
import os
import sys
from copy import copy

# --local: igual que m001/m002/m003. Debe ir ANTES de importar conexion_drive.config.
if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

ID_MIGRACION = "m004_competencia"

NODO = "NODO INTEGRAL"
COMBOS = "Tabla Combos Dinamicos"
DIC = "DiccionarioDatos"

URL_MUSI = "https://tiendasmusi.com/promociones-para-panaderia-y-reposteria/"
API_MUSI = ("https://tiendasmusi.com/wp-json/wp/v2/pages"
            "?slug=promociones-para-panaderia-y-reposteria")

# ---------------------------------------------------------------- Bloque A (NODO INTEGRAL)
BLOQUE_A = [
    (NODO, "B11", "DATA OFERTA ", "DATA OFERTA MUSMANNI"),
    (NODO, "A21", None, "Precios de la competencia"),
    (NODO, "B21", None, "DATA OFERTA MUSI"),
    (NODO, "C21", None, "valor /scraping web"),
    (NODO, "D21", None, "Bloqueado (Lectura)"),
    (NODO, "E21", None, "Input alfanumérico (API WordPress + visión Gemini)"),
]

# ---------------------------------------------------------------- Bloque B (Tabla Combos)
BLOQUE_B = [
    (COMBOS, "C11", None, "Promocion Musi"),
    (COMBOS, "D11", None, "imput detalle de oferta "),
    (COMBOS, "E11", None, "valor /scraping web + IA google identifica los datos en las imágenes "),
]

# ---------------------------------------------------------------- Bloque C (notas agregadas)
NOTA_L252 = (
    " || m004 (25-sep-2026): URL VIGENTE https://musmanni.com/promociones/ (la de arriba esta "
    "desactualizada desde mar-2024). API: https://musmanni.com/wp-json/wp/v2/pages?slug=promociones "
    "(+ subpaginas promociones-reposteria, -almuerzos, -pasteles) y "
    "https://musmanni.com/wp-json/wp/v2/promos"
)
NOTA_N252 = (
    " || m004 (25-sep-2026): Desde 25-sep-2026 solo se monitorea a Musi. Musi nacio en 2017 "
    "como el nuevo nombre de los Musmanni Mini Super; las panaderias Musmanni siguen como marca "
    "aparte, sin promociones web nuevas desde jul-2025. Ambas marcas son de Heineken Costa Rica "
    "desde ene-2026 (antes FIFCO). Musmanni queda desactivado en competencia.py (se reactiva "
    "cambiando una marca); sus celdas (NODO C11, Tabla Combos D10/E10) no se tocan. Musi (la "
    "mitad del local es panaderia, confirmado por el usuario) esta en NODO INTEGRAL C21 "
    "(NODO.PRECIOS_COMPETENCIA_MUSI). Motor de vision = Google Gemini, modelo 'flash' mas "
    "nuevo disponible para la clave (25-sep-2026: gemini-3.8-flash; opcional GEMINI_MODELO "
    "en el .env; clave GEMINI_API_KEY en el .env). Frecuencia: 1 vez al dia, inicio 7:00 a. m. "
    "hora CR (resultado listo antes de las 8:00), en la misma tarea programada del NODO "
    "INTEGRAL (GitHub Actions, 13:00 UTC). Robot: "
    "microservicios/grupo3_integraciones/competencia.py."
)
NOTA_N89 = (
    " || m004 (25-sep-2026): Musmanni (fila 10) DESACTIVADO desde 25-sep-2026: solo se "
    "monitorea a Musi (fila 11, MKT.COMPETENCIA.MUSI); D10/E10 no se tocan. Musi = antes "
    "Musmanni Mini Super (2017); ambas marcas de Heineken Costa Rica desde ene-2026. "
    "IMPLEMENTADO en competencia.py: capa 1 = API de WordPress (sin IA), solo imagenes "
    "VISIBLES (las ocultas con elementor-hidden-* no cuentan); capa 2 = Gemini lee producto, "
    "precio y vigencia de cada imagen (una sola vez, con cache). D11 = detalle; E11 = precio "
    "mas bajo vigente en CRC. Sin promo vigente: 'Sin promocion vigente (fecha)' y 0. Si falla "
    "la conexion o la lectura queda incompleta (alguna imagen visible sin leer) NO se toca. "
    "No es un combo propio: sirve para comparar en el SaaS y ajustar los combos propios si "
    "hace falta."
)
NOTAS_C = [(DIC, "L252", NOTA_L252), (DIC, "N252", NOTA_N252), (DIC, "N89", NOTA_N89)]

# ---------------------------------------------------------------- Bloque D (filas nuevas)
FILA_303 = [
    "NODO.PRECIOS_COMPETENCIA_MUSI", "Precios de la competencia", NODO, "(sin Tabla formal)",
    "C21 (DATA OFERTA MUSI)", "Alfanumérico (resumen del día)",
    "API WordPress + visión Gemini (Musi)", "Robot (competencia.py)",
    "MKT.COMPETENCIA.MUSI (Tabla Combos Dinamicos)", "No", "No", URL_MUSI,
    "MKT.COMPETENCIA.MUSI",
    "CREADO m004 (25-sep-2026): resumen del dia de la competencia Musi, p. ej. "
    "'Sin promocion vigente (25-09-2026)' o 'N promocion(es) vigente(s) (fecha)'. "
    "Si falla la conexion o la lectura queda incompleta NO se toca. Frecuencia: 1 vez al dia, "
    "inicio 7:00 a. m. hora CR, junto con el NODO INTEGRAL.",
]
FILA_304 = [
    "MKT.COMPETENCIA.MUSI", "Marketing", COMBOS, "(sin Tabla formal)",
    "C11:E11 ('Promocion Musi')", "Texto (D11) y número CRC (E11)",
    "API WordPress + visión Gemini", "Robot (competencia.py)", "SaaS / BI", "No", "No",
    API_MUSI, "NODO.PRECIOS_COMPETENCIA_MUSI",
    "CREADO m004 (25-sep-2026): D11 = detalle de las promociones vigentes de Musi "
    "(producto, precio, vigencia); E11 = precio mas bajo vigente en CRC, 0 si no hay. "
    "No es un combo propio: se muestra junto a los combos para comparar. Mismas reglas que "
    "MKT.COMPETENCIA.MUSMANI (fila 10).",
]
COLUMNAS = "ABCDEFGHIJKLMN"
BLOQUE_D = ([(DIC, f"{c}303", None, v) for c, v in zip(COLUMNAS, FILA_303)]
            + [(DIC, f"{c}304", None, v) for c, v in zip(COLUMNAS, FILA_304)])

# ---------------------------------------------------------------- Formato (copiar estilo)
# (hoja, fila origen, fila destino, columnas)
FORMATOS = [
    (NODO, 11, 21, "ABCDE"),
    (COMBOS, 10, 11, "CDE"),
    (DIC, 252, 303, COLUMNAS),
    (DIC, 252, 304, COLUMNAS),
]


class MigracionDetenida(Exception):
    """Una celda no tenia el valor esperado: no se escribio nada."""


def _lista_cambios(wb):
    """Arma la lista final (hoja, celda, antes, despues, bloque)."""
    cambios = [(h, c, a, d, "A") for h, c, a, d in BLOQUE_A]
    cambios += [(h, c, a, d, "B") for h, c, a, d in BLOQUE_B]
    for h, c, nota in NOTAS_C:
        actual = wb[h][c].value
        if isinstance(actual, str) and actual.endswith(nota):
            cambios.append((h, c, actual[: -len(nota)], actual, "C"))
        else:
            cambios.append((h, c, actual, (actual or "") + nota, "C"))
    cambios += [(h, c, a, d, "D") for h, c, a, d in BLOQUE_D]
    return cambios


def _copiar_formato(ws, fila_origen, fila_destino, columnas):
    for col in columnas:
        origen, destino = ws[f"{col}{fila_origen}"], ws[f"{col}{fila_destino}"]
        if origen.has_style:
            destino.font = copy(origen.font)
            destino.fill = copy(origen.fill)
            destino.border = copy(origen.border)
            destino.alignment = copy(origen.alignment)
            destino.number_format = origen.number_format
            destino.protection = copy(origen.protection)
    altura = ws.row_dimensions[fila_origen].height
    if altura:
        ws.row_dimensions[fila_destino].height = altura


def migrar(wb, wb_valores=None):
    cambios = _lista_cambios(wb)

    # 1) Revisar TODO antes de escribir
    problemas, pendientes = [], []
    for hoja, celda, antes, despues, bloque in cambios:
        valor = wb[hoja][celda].value
        if valor == despues:
            continue
        if valor != antes:
            problemas.append(f"'{hoja}'!{celda}: se esperaba {antes!r} y tiene {valor!r}")
        else:
            pendientes.append((hoja, celda, antes, despues, bloque))
    if problemas:
        raise MigracionDetenida("NO SE CAMBIO NADA. Celdas distintas a lo esperado:\n  "
                                + "\n  ".join(problemas))

    # 2) Regla del proyecto: ningun texto nuevo empieza con "="
    for hoja, celda, antes, despues, bloque in pendientes:
        if isinstance(despues, str) and despues.startswith("="):
            raise MigracionDetenida(f"Texto que empieza con '=' en '{hoja}'!{celda}. No se escribio nada.")

    if not pendientes:
        return {"hubo_cambios": False, "cambios": ["Ya estaba aplicada: 0 celdas pendientes."],
                "detalle": []}

    # 3) Escribir valores y copiar formatos
    for hoja, celda, antes, despues, bloque in pendientes:
        wb[hoja][celda] = despues
    for hoja, origen, destino, columnas in FORMATOS:
        _copiar_formato(wb[hoja], origen, destino, columnas)

    # 4) Volver a leer y confirmar
    for hoja, celda, antes, despues, bloque in pendientes:
        if wb[hoja][celda].value != despues:
            raise MigracionDetenida(f"Verificacion fallo en '{hoja}'!{celda}")

    por_bloque = {}
    for p in pendientes:
        por_bloque[p[4]] = por_bloque.get(p[4], 0) + 1
    resumen = [f"Bloque {b}: {n} celdas" for b, n in sorted(por_bloque.items())]

    from conexion_drive.tablas import registrar_evento
    registrar_evento(
        wb, modulo="Sistema", tipo_evento="MIGRACION", entidad="Libro Maestro",
        id_entidad=ID_MIGRACION, accion="Competencia 3C (Musmanni + Musi)", estado="OK",
        detalle=("A) NODO INTEGRAL B11 = DATA OFERTA MUSMANNI + fila 21 para Musi | "
                 "B) Tabla Combos Dinamicos fila 11 para Musi | C) notas '|| m004' en "
                 "DiccionarioDatos L252, N252, N89 | D) DiccionarioDatos filas 303 y 304 | "
                 f"Total: {len(pendientes)} celdas"),
    )
    return {"hubo_cambios": True, "cambios": resumen, "detalle": pendientes}


def _aplicar_a_archivo(ruta):
    """Aplica la migracion a un .xlsx suelto y lo recalcula con LibreOffice."""
    import io
    from pathlib import Path
    import openpyxl
    from conexion_drive.recalculo import reparar_prefijos_xlfn, recalcular_bytes

    ruta = Path(ruta)
    contenido, _ = reparar_prefijos_xlfn(ruta.read_bytes())
    wb = openpyxl.load_workbook(io.BytesIO(contenido))
    resultado = migrar(wb, None)
    if resultado["hubo_cambios"]:
        salida = io.BytesIO()
        wb.save(salida)
        ruta.write_bytes(recalcular_bytes(salida.getvalue()))
    return resultado


if __name__ == "__main__":
    try:
        if "--archivo" in sys.argv:
            ruta = sys.argv[sys.argv.index("--archivo") + 1]
            print(f"Excel de destino: ARCHIVO {ruta}")
            resultado = _aplicar_a_archivo(ruta)
        else:
            from conexion_drive import config
            from conexion_drive.drive_excel import editar_excel
            print(f"Excel de destino: {'COPIA LOCAL ' + str(config.EXCEL_LOCAL_PATH) if config.ORIGEN_EXCEL == 'local' else 'GOOGLE DRIVE (real)'}")
            resultado = editar_excel(migrar, descripcion="Migracion m004_competencia")
    except MigracionDetenida as error:
        print(str(error))
        sys.exit(1)
    for linea in resultado["cambios"]:
        print("  " + linea)
    print("LISTO." if resultado["hubo_cambios"] else "Ya estaba aplicada: no se cambio nada.")
