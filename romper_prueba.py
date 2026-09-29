"""Rompe A PROPOSITO las formulas de NODO INTEGRAL C19/C20 en la copia de
prueba, quitando el prefijo _xlfn. (igual que lo hizo Google Sheets).
Antes de romper, guarda un respaldo de la copia de prueba.
Nunca toca el Excel real de Drive.

|| V15 (28-sep-2026, Decisiones Pendientes #27, decision del usuario = opcion B):
   Para que sirve: es la EVIDENCIA del escudo _xlfn (#11, commit 67d3361). Imita el
   dano que hizo Google Sheets (Regla #12) para demostrar que el robot repara solo
   esas formulas la proxima vez que abre el Excel.
   Dos seguros nuevos:
     1) PREGUNTA antes de romper: hay que escribir SI. Si alguien lo corre por
        accidente (boton Run de VS Code), no hace nada.
     2) El respaldo NUNCA se pisa: si maestro_prueba_RESPALDO.xlsx ya existe, el
        respaldo nuevo lleva fecha y hora en el nombre.
   Uso (desde la carpeta ROBOT PANADERO):  py romper_prueba.py
"""
import os
import shutil
from datetime import datetime

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

ruta = 'datos_prueba/maestro_prueba.xlsx'
respaldo = 'datos_prueba/maestro_prueba_RESPALDO.xlsx'


def respaldo_sin_pisar():
    """Devuelve un nombre de respaldo que todavia no existe (nunca sobrescribe)."""
    if not os.path.exists(respaldo):
        return respaldo
    marca = datetime.now().strftime('%Y%m%d_%H%M%S')
    return respaldo.replace('.xlsx', f'_{marca}.xlsx')


def main():
    print('ATENCION: esto ROMPE a proposito NODO INTEGRAL C19/C20 en la COPIA LOCAL')
    print(f'  ({ruta}). Nunca toca Google Drive. Sirve para probar el escudo _xlfn (#11).')
    if input('Escriba SI para continuar (cualquier otra cosa cancela): ').strip().upper() != 'SI':
        print('Cancelado. No se toco nada.')
        return

    destino = respaldo_sin_pisar()
    shutil.copy2(ruta, destino)
    print('Respaldo guardado en', destino)

    wb = openpyxl.load_workbook(ruta)
    ws = wb['NODO INTEGRAL']

    rotas = 0
    for celda in ('C19', 'C20'):
        valor = ws[celda].value
        if isinstance(valor, ArrayFormula):          # formula matricial
            antes = valor.text
            valor.text = antes.replace('_xlfn.', '')
            ahora = valor.text
        elif isinstance(valor, str) and valor.startswith('='):   # formula normal
            antes = valor
            ws[celda] = antes.replace('_xlfn.', '')
            ahora = ws[celda].value
        else:
            print(celda, 'no tiene formula, se salta')
            continue
        print(celda, 'antes:', antes[:60])
        print(celda, 'ahora:', ahora[:60])
        rotas += 1

    wb.save(ruta)
    print(f'{rotas} formula(s) rotas a proposito (igual que lo hizo Google Sheets)')


if __name__ == '__main__':
    main()
