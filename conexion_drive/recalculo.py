"""Recalcula TODAS las formulas del Excel con LibreOffice (sin ventanas).

Por que hace falta: openpyxl guarda las formulas pero BORRA sus resultados
guardados. Si se subiera asi a Google Drive, la siguiente lectura del robot
(precios, combos activos, etc.) veria celdas vacias. Por eso, antes de subir,
LibreOffice abre el archivo, calcula todo y lo guarda con los resultados.

Probado con el Excel maestro v68: 4443 formulas, 0 errores, tablas y
formulas intactas despues del recalculo.
"""
import os
import platform
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from openpyxl import load_workbook
from io import BytesIO

from . import config

_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalcularYGuardar()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""

_ERRORES_EXCEL = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


class ErrorRecalculo(Exception):
    pass


def encontrar_soffice():
    """Busca el programa de LibreOffice en las rutas normales de Windows, Linux y Mac."""
    candidatos = []
    if config.SOFFICE_PATH:
        candidatos.append(config.SOFFICE_PATH)
    for nombre in ("soffice", "libreoffice"):
        encontrado = shutil.which(nombre)
        if encontrado:
            candidatos.append(encontrado)
    if platform.system() == "Windows":
        for base in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                     os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")):
            # soffice.com espera a que termine; soffice.exe a veces no
            candidatos.append(str(Path(base) / "LibreOffice" / "program" / "soffice.com"))
            candidatos.append(str(Path(base) / "LibreOffice" / "program" / "soffice.exe"))
    candidatos.append("/Applications/LibreOffice.app/Contents/MacOS/soffice")

    for ruta in candidatos:
        if ruta and Path(ruta).exists():
            return ruta
    raise ErrorRecalculo(
        "No encontre LibreOffice. Instalalo desde https://www.libreoffice.org "
        "o escribe su ruta en SOFFICE_PATH dentro del archivo .env"
    )


def _entorno():
    entorno = os.environ.copy()
    if platform.system() == "Linux":
        entorno.setdefault("SAL_USE_VCLPLUGIN", "svp")
    return entorno


def recalcular_bytes(contenido: bytes, timeout: int = 180) -> bytes:
    """Recibe el .xlsx en bytes y devuelve el .xlsx recalculado en bytes."""
    soffice = encontrar_soffice()
    with tempfile.TemporaryDirectory(prefix="robot_recalculo_") as carpeta:
        carpeta = Path(carpeta)
        archivo = carpeta / "maestro.xlsx"
        archivo.write_bytes(contenido)
        perfil = carpeta / "perfil_libreoffice"
        perfil_url = perfil.as_uri()

        # 1) Crear un perfil temporal de LibreOffice (no toca tu LibreOffice normal)
        subprocess.run(
            [soffice, "--headless", "--norestore", "--terminate_after_init",
             f"-env:UserInstallation={perfil_url}"],
            capture_output=True, timeout=timeout, env=_entorno(),
        )
        carpeta_macro = perfil / "user" / "basic" / "Standard"
        if not carpeta_macro.exists():
            raise ErrorRecalculo("LibreOffice no creo su perfil temporal; no se recalculo.")
        (carpeta_macro / "Module1.xba").write_text(_MACRO, encoding="utf-8")

        # 2) Abrir el archivo, recalcular todo y guardar
        antes = archivo.stat().st_mtime_ns
        resultado = subprocess.run(
            [soffice, "--headless", "--norestore", f"-env:UserInstallation={perfil_url}",
             "vnd.sun.star.script:Standard.Module1.RecalcularYGuardar?language=Basic&location=application",
             str(archivo)],
            capture_output=True, text=True, timeout=timeout, env=_entorno(),
        )
        if resultado.returncode != 0:
            raise ErrorRecalculo(f"LibreOffice fallo: {resultado.stderr.strip()}")

        # En Windows a veces el programa devuelve el control antes de terminar de guardar
        limite = time.monotonic() + 30
        while archivo.stat().st_mtime_ns == antes and time.monotonic() < limite:
            time.sleep(1)
        if archivo.stat().st_mtime_ns == antes:
            raise ErrorRecalculo(
                "LibreOffice no guardo el archivo. Cierra cualquier ventana de LibreOffice "
                "abierta y vuelve a intentar."
            )
        return archivo.read_bytes()


def buscar_errores(contenido: bytes, maximo: int = 20):
    """Devuelve una lista de celdas con errores (#REF!, #N/A, etc.) en el archivo."""
    wb = load_workbook(BytesIO(contenido), data_only=True)
    errores = []
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and celda.value.strip() in _ERRORES_EXCEL:
                    errores.append(f"{ws.title}!{celda.coordinate} = {celda.value}")
                    if len(errores) >= maximo:
                        return errores
    return errores


def valores_recalculados(wb):
    """Recalcula una COPIA del libro que se esta editando y la devuelve con valores.

    Sirve cuando una edicion necesita leer un resultado que depende de lo que
    se acaba de escribir (ej. escribir la cantidad de la hornada en G4 y luego
    leer el costo unitario que la receta calcula con esa cantidad). El libro
    original 'wb' no se modifica.
    """
    memoria = BytesIO()
    wb.save(memoria)
    return load_workbook(BytesIO(recalcular_bytes(memoria.getvalue())), data_only=True)
