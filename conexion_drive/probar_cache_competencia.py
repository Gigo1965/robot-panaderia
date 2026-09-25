"""Prueba de SOLO LECTURA: el robot puede leer CACHE_COMPETENCIA.xlsx en Google Drive.

Que revisa (no escribe NADA en Drive):
  1) Credenciales de la cuenta de servicio (carpeta credenciales/).
  2) Datos del archivo en Drive: nombre, tipo, tamano y permisos del robot.
     - El tipo debe ser Excel real (.xlsx), NO una hoja de Google Sheets.
     - El robot debe poder EDITAR y DESCARGAR.
  3) Descarga el archivo y lo abre con openpyxl:
     hojas Cache / Control / Instrucciones, tabla TablaCache y sus 14 columnas.

El ID del archivo se toma de:
  - el argumento:   py -m conexion_drive.probar_cache_competencia <ID>
  - o del .env:     CACHE_COMPETENCIA_FILE_ID=<ID>
"""
import io
import os
import sys
from pathlib import Path

from conexion_drive import config  # carga el .env

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_SHEETS = "application/vnd.google-apps.spreadsheet"
COLUMNAS = ["ID_LECTURA", "FECHA_LECTURA", "COMPETIDOR", "URL_IMAGEN", "FECHA_SUBIDA_IMAGEN",
            "ES_PROMOCION", "PRODUCTO", "PRECIO_CRC", "VIGENCIA_DESDE", "VIGENCIA_HASTA",
            "TEXTO_VIGENCIA", "ESTADO_AL_LEER", "MODELO_GEMINI", "ORIGEN"]
SCOPES = ["https://www.googleapis.com/auth/drive"]


def _credenciales():
    ruta = os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
    if ruta and Path(ruta).exists():
        return Path(ruta)
    candidatos = sorted((Path(config.RAIZ_PROYECTO) / "credenciales").glob("*.json"))
    if not candidatos:
        raise SystemExit("FALLO: no hay ningun .json en la carpeta credenciales/")
    return candidatos[0]


def main():
    ok = True
    file_id = (sys.argv[1] if len(sys.argv) > 1 else os.getenv("CACHE_COMPETENCIA_FILE_ID") or "").strip()
    if not file_id:
        raise SystemExit("FALLO: falta el ID. Agrega CACHE_COMPETENCIA_FILE_ID=<ID> al .env "
                         "o pasalo como argumento.")

    from google.oauth2 import service_account
    from google.auth.transport.requests import AuthorizedSession

    print("=== 1) Credenciales de la cuenta de servicio ===")
    ruta = _credenciales()
    cred = service_account.Credentials.from_service_account_file(str(ruta), scopes=SCOPES)
    print(f"OK  {ruta.name}  ({cred.service_account_email})")
    sesion = AuthorizedSession(cred)

    print("\n=== 2) Datos del archivo en Drive ===")
    r = sesion.get(f"https://www.googleapis.com/drive/v3/files/{file_id}",
                   params={"fields": "name,mimeType,size,modifiedTime,capabilities(canEdit,canDownload)",
                           "supportsAllDrives": "true"}, timeout=30)
    if r.status_code != 200:
        print(f"FALLO  HTTP {r.status_code}: {r.text[:300]}")
        print("  404 = el ID esta mal o el archivo NO esta compartido con la cuenta de servicio.")
        raise SystemExit(1)
    meta = r.json()
    cap = meta.get("capabilities", {})
    print(f"    Nombre:     {meta.get('name')}")
    print(f"    Modificado: {meta.get('modifiedTime')}   Tamano: {meta.get('size')} bytes")
    if meta.get("mimeType") == MIME_XLSX:
        print("OK  Tipo: Excel real (.xlsx)")
    elif meta.get("mimeType") == MIME_SHEETS:
        print("FALLO  Tipo: Google Sheets. Se convirtio al subirlo: hay que borrarlo de Drive, "
              "desmarcar 'Convertir archivos subidos' y subirlo otra vez.")
        raise SystemExit(1)
    else:
        print(f"FALLO  Tipo inesperado: {meta.get('mimeType')}")
        ok = False
    print(("OK  " if cap.get("canEdit") else "FALLO  ") + f"El robot puede EDITAR: {cap.get('canEdit')}")
    print(("OK  " if cap.get("canDownload") else "FALLO  ") + f"El robot puede DESCARGAR: {cap.get('canDownload')}")
    ok = ok and bool(cap.get("canEdit")) and bool(cap.get("canDownload"))

    print("\n=== 3) Descargando y abriendo con openpyxl ===")
    r = sesion.get(f"https://www.googleapis.com/drive/v3/files/{file_id}",
                   params={"alt": "media", "supportsAllDrives": "true"}, timeout=60)
    if r.status_code != 200:
        print(f"FALLO  descarga HTTP {r.status_code}: {r.text[:300]}")
        raise SystemExit(1)
    print(f"OK  Descargado: {len(r.content)} bytes")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    print(f"    Hojas: {wb.sheetnames}")
    faltan = [h for h in ("Cache", "Control", "Instrucciones") if h not in wb.sheetnames]
    if faltan:
        print(f"FALLO  faltan hojas: {faltan}")
        raise SystemExit(1)
    ws = wb["Cache"]
    if "TablaCache" not in ws.tables:
        print("FALLO  la hoja Cache no tiene la tabla TablaCache")
        raise SystemExit(1)
    ref = ws.tables["TablaCache"].ref
    encabezados = [c.value for c in ws[ref.split(":")[0][0] + "1":ref.split(":")[1][0] + "1"][0]]
    print(f"OK  Tabla TablaCache en {ref}")
    if encabezados == COLUMNAS:
        print(f"OK  Las 14 columnas estan en orden")
    else:
        print(f"FALLO  columnas distintas: {encabezados}")
        ok = False
    lecturas = sum(1 for fila in ws.iter_rows(min_row=2, max_col=1, values_only=True) if fila[0])
    ctl = wb["Control"]
    print(f"    Lecturas guardadas: {lecturas}")
    print(f"    Limite del ciclo (Control B4): {ctl['B4'].value:,}   Ciclo actual (B8): {ctl['B8'].value}")

    print("\n" + ("TODO BIEN: el robot puede leer (y tiene permiso de editar) el cache en Drive. "
                  "No se escribio nada." if ok else "HAY FALLOS: revisa las lineas marcadas con FALLO."))


if __name__ == "__main__":
    main()
