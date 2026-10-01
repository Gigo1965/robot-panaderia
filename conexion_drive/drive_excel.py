"""Descargar, editar y subir el Excel maestro en Google Drive (Cuenta de Servicio).

Funciones publicas:
    leer_excel(solo_valores=True)   -> libro openpyxl para SOLO LEER
    editar_excel(funcion, ...)      -> descarga, aplica tu funcion, recalcula y sube
    obtener_metadatos()             -> nombre, version, fecha de modificacion del archivo

Proteccion contra choques (dos procesos escribiendo a la vez, ej. el bot de
WhatsApp y un microservicio): Google Drive numera cada version del archivo.
Justo antes de subir, el robot revisa que la version siga siendo la misma que
descargo. Si alguien subio otra version en medio, NO sobrescribe: vuelve a
descargar y repite la edicion (hasta 3 intentos).

|| V15 (1-oct-2026, Pendiente #37, opcion B1 aprobada por el usuario): Google
Drive sube el NUMERO de version por su cuenta segundos despues de cada subida,
aunque nadie cambie el contenido. Eso producia "conflictos" falsos (la compra
F006 del 28-sep entro recien en el 3er intento). Ahora el choque se detecta con
la HUELLA del contenido (md5Checksum de Drive), que solo cambia si alguien de
verdad subio un archivo distinto. Se mantienen los 3 intentos. El numero de
version se sigue leyendo solo como informacion (probar_conexion, registro).
El modo local (copia de prueba) no cambia: sigue usando la hora del archivo.
"""
import hashlib
import json
import logging
import time
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from . import config
from .recalculo import ErrorRecalculo, buscar_errores, recalcular_bytes, reparar_prefijos_xlfn

log = logging.getLogger("conexion_drive")

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PERMISOS = ["https://www.googleapis.com/auth/drive"]


class ErrorConexionExcel(Exception):
    """Cualquier problema al conectar, descargar o subir el Excel."""


class ConflictoDeVersion(ErrorConexionExcel):
    """Otro proceso modifico el Excel mientras este lo editaba."""


# ----------------------------------------------------------------------------
# Conexion con Google Drive
# ----------------------------------------------------------------------------
@lru_cache(maxsize=1)
def _servicio_drive():
    # Se importan aqui para que el modo "local" funcione aunque no esten instaladas
    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
    except ImportError as error:
        raise ErrorConexionExcel(
            "Faltan las librerias de Google. En la Terminal ejecuta:  pip install -r requirements.txt"
        ) from error

    if config.GOOGLE_SERVICE_ACCOUNT_JSON:
        info = json.loads(config.GOOGLE_SERVICE_ACCOUNT_JSON)
        credenciales = service_account.Credentials.from_service_account_info(info, scopes=PERMISOS)
    elif config.GOOGLE_SERVICE_ACCOUNT_FILE and Path(config.GOOGLE_SERVICE_ACCOUNT_FILE).exists():
        credenciales = service_account.Credentials.from_service_account_file(
            str(config.GOOGLE_SERVICE_ACCOUNT_FILE), scopes=PERMISOS
        )
    else:
        raise ErrorConexionExcel(
            "No encontre la clave de la Cuenta de Servicio. Revisa GOOGLE_SERVICE_ACCOUNT_FILE "
            f"en el archivo .env (ruta buscada: {config.GOOGLE_SERVICE_ACCOUNT_FILE})"
        )
    return build("drive", "v3", credentials=credenciales, cache_discovery=False)


def _id_archivo():
    if not config.EXCEL_MAESTRO_FILE_ID or "pega_aqui" in config.EXCEL_MAESTRO_FILE_ID:
        raise ErrorConexionExcel("Falta EXCEL_MAESTRO_FILE_ID en el archivo .env")
    return config.EXCEL_MAESTRO_FILE_ID


def _modo_local():
    return config.ORIGEN_EXCEL == "local"


def obtener_metadatos():
    """Datos basicos del archivo (sirve para probar la conexion)."""
    if _modo_local():
        ruta = Path(config.EXCEL_LOCAL_PATH)
        return {"name": ruta.name, "version": str(ruta.stat().st_mtime_ns),
                "mimeType": MIME_XLSX, "modo": "local"}
    try:
        return _servicio_drive().files().get(
            fileId=_id_archivo(),
            fields="id,name,mimeType,version,modifiedTime,size,md5Checksum",
            supportsAllDrives=True,
        ).execute()
    except ErrorConexionExcel:
        raise
    except Exception as error:  # errores de red, permisos, ID mal escrito...
        raise ErrorConexionExcel(_explicar_error_google(error)) from error


def _explicar_error_google(error):
    texto = str(error)
    if "404" in texto or "notFound" in texto:
        return ("Google Drive dice que el archivo no existe o no esta compartido con la "
                "Cuenta de Servicio. Revisa el EXCEL_MAESTRO_FILE_ID y que el archivo este "
                "compartido como Editor con robot-panaderia@robot-panaderia-canas.iam.gserviceaccount.com")
    if "403" in texto:
        return ("Google Drive nego el permiso (403). Revisa que la Google Drive API este "
                "habilitada en el proyecto robot-panaderia-canas y que la cuenta sea Editor.")
    return f"Error de Google Drive: {texto}"


# ----------------------------------------------------------------------------
# Descargar y subir
# ----------------------------------------------------------------------------
def _aplicar_escudo_xlfn(contenido: bytes) -> bytes:
    """Repara en memoria las formulas que perdieron el prefijo _xlfn. y avisa.

    V2 (24-sep-2026): al reparar, openpyxl guarda las formulas pero BORRA los
    resultados calculados (precios, totales...). Sin recalcular, el robot se
    quedaria "ciego" (precios en None -> PrecioSinCalcular) hasta el siguiente
    guardado. Por eso, SOLO cuando hubo reparaciones, se recalcula con
    LibreOffice aqui mismo (~10 s extra, solo en ese caso raro).
    """
    contenido, reparaciones = reparar_prefijos_xlfn(contenido)
    if reparaciones:
        log.warning(
            "Escudo _xlfn: se repararon %s formula(s) danadas (ej. por Google Sheets):",
            len(reparaciones),
        )
        for detalle in reparaciones[:10]:
            log.warning("    %s", detalle)
        if config.RECALCULAR:
            try:
                contenido = recalcular_bytes(contenido)
            except ErrorRecalculo as error:
                raise ErrorConexionExcel(
                    "El escudo _xlfn reparo formulas pero no pudo recalcular con LibreOffice: "
                    f"{error}"
                ) from error
            log.warning("Escudo _xlfn: libro recalculado despues de la reparacion.")
    return contenido


def _descargar():
    """Devuelve (contenido_en_bytes, marca), ya con el escudo _xlfn aplicado.

    marca = lo que _subir usa para saber si alguien mas escribio en medio:
      - modo local: la hora del archivo (texto), igual que siempre.
      - Drive (|| V15, #37): {"md5": huella de los bytes DESCARGADOS (antes del
        escudo), "version": numero de version (solo informativo)}.
    """
    if _modo_local():
        ruta = Path(config.EXCEL_LOCAL_PATH)
        if not ruta.exists():
            raise ErrorConexionExcel(f"No existe la copia local: {ruta}")
        return _aplicar_escudo_xlfn(ruta.read_bytes()), str(ruta.stat().st_mtime_ns)

    from googleapiclient.http import MediaIoBaseDownload

    metadatos = obtener_metadatos()
    if metadatos.get("mimeType") != MIME_XLSX:
        raise ErrorConexionExcel(
            f"El archivo en Drive no es un .xlsx (es {metadatos.get('mimeType')}). "
            "No debe convertirse a Google Sheets."
        )
    try:
        solicitud = _servicio_drive().files().get_media(fileId=_id_archivo(), supportsAllDrives=True)
        memoria = BytesIO()
        descargador = MediaIoBaseDownload(memoria, solicitud)
        terminado = False
        while not terminado:
            _, terminado = descargador.next_chunk()
    except Exception as error:
        raise ErrorConexionExcel(_explicar_error_google(error)) from error
    crudo = memoria.getvalue()
    # La huella se calcula sobre lo que REALMENTE se bajo (no sobre los metadatos
    # pedidos un instante antes): asi no hay hueco entre leer metadatos y descargar.
    marca = {"md5": hashlib.md5(crudo).hexdigest(), "version": metadatos.get("version")}
    return _aplicar_escudo_xlfn(crudo), marca


def _subir(contenido: bytes, version_descargada: str):
    if _modo_local():
        ruta = Path(config.EXCEL_LOCAL_PATH)
        if str(ruta.stat().st_mtime_ns) != version_descargada:
            raise ConflictoDeVersion("La copia local cambio mientras se editaba.")
        temporal = ruta.with_suffix(".tmp")
        temporal.write_bytes(contenido)
        temporal.replace(ruta)
        return {"version": str(ruta.stat().st_mtime_ns)}

    from googleapiclient.http import MediaIoBaseUpload

    # || V15 (#37): antes se comparaba el numero de version (ver docstring); ahora la huella.
    actuales = obtener_metadatos()
    md5_actual = actuales.get("md5Checksum")
    md5_descargado = version_descargada["md5"]
    if not md5_actual:
        raise ErrorConexionExcel(
            "Google Drive no devolvio la huella (md5Checksum) del Excel; no se subio nada por seguridad."
        )
    if md5_actual != md5_descargado:
        raise ConflictoDeVersion(
            "El Excel cambio en Drive mientras se editaba (contenido distinto; version "
            f"{version_descargada.get('version')} -> {actuales.get('version')})."
        )
    if actuales.get("version") != version_descargada.get("version"):
        log.info(
            "Drive cambio solo el numero de version (%s -> %s) con el mismo contenido: no es un choque.",
            version_descargada.get("version"), actuales.get("version"),
        )
    try:
        media = MediaIoBaseUpload(BytesIO(contenido), mimetype=MIME_XLSX, resumable=True)
        return _servicio_drive().files().update(
            fileId=_id_archivo(),
            media_body=media,
            fields="id,version,modifiedTime",
            supportsAllDrives=True,
        ).execute()
    except Exception as error:
        raise ErrorConexionExcel(_explicar_error_google(error)) from error


# ----------------------------------------------------------------------------
# Funciones publicas
# ----------------------------------------------------------------------------
def leer_excel(solo_valores: bool = True):
    """Descarga el Excel y lo abre para LEER.

    solo_valores=True  -> las celdas con formula traen su resultado (ej. el precio)
    solo_valores=False -> las celdas con formula traen el texto de la formula
    """
    contenido, _ = _descargar()
    return load_workbook(BytesIO(contenido), data_only=solo_valores)


def editar_excel(funcion_edicion, descripcion: str = "", intentos: int = 3):
    """Descarga el Excel, ejecuta funcion_edicion(wb, wb_valores), recalcula y sube.

    - wb:         libro con FORMULAS. Todo cambio se hace aqui.
    - wb_valores: el mismo libro con los VALORES calculados. Solo para leer
                  (precios, combos activos...). Nunca se sube.

    Si funcion_edicion devuelve un diccionario con "hubo_cambios": False
    (por ejemplo, una venta duplicada), no se sube nada.
    """
    for intento in range(1, intentos + 1):
        contenido, version = _descargar()
        wb = load_workbook(BytesIO(contenido))
        wb_valores = load_workbook(BytesIO(contenido), data_only=True)

        resultado = funcion_edicion(wb, wb_valores)
        if isinstance(resultado, dict) and resultado.get("hubo_cambios") is False:
            log.info("%s: sin cambios, no se sube nada.", descripcion)
            return resultado

        memoria = BytesIO()
        wb.save(memoria)
        nuevo = memoria.getvalue()

        if config.RECALCULAR:
            nuevo = recalcular_bytes(nuevo)
            errores_antes = set(buscar_errores(contenido, maximo=500))
            errores_nuevos = [e for e in buscar_errores(nuevo, maximo=500) if e not in errores_antes]
            if errores_nuevos:
                raise ErrorConexionExcel(
                    "El cambio produjo errores de formula; NO se subio nada. "
                    f"Primeras celdas: {errores_nuevos[:10]}"
                )
        else:
            log.warning("RECALCULAR=no: el archivo se sube sin resultados de formulas.")

        try:
            _subir(nuevo, version)
            log.info("%s: guardado (intento %s).", descripcion, intento)
            return resultado
        except ConflictoDeVersion as conflicto:
            log.warning("%s: %s Reintentando...", descripcion, conflicto)
            time.sleep(2 * intento)

    raise ConflictoDeVersion(f"{descripcion}: no se pudo guardar tras {intentos} intentos.")
