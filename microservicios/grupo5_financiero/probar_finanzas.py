"""Prueba completa de finanzas.py sobre una COPIA TEMPORAL del Excel de prueba.

No toca datos_prueba/maestro_prueba.xlsx ni Google Drive: copia el Excel de prueba a
datos_prueba/_prueba_finanzas.xlsx (se reemplaza en cada corrida) y ahi ejecuta 26 casos:
9 cambios validos, 2 repetidos (SIN_CAMBIO) y 15 que DEBEN ser rechazados. Despues revisa lo que
calculo el Excel (egresos del dia, saldo de caja), los eventos, que la fila 26 de 'flujo caja anual'
siga vacia y que TODAS las formulas del libro quedaron identicas al original.

|| 6-oct-2026 (paquete 2): casos 27 a 43, los cuatro supuestos financieros ('flujo caja proyecto' C28 y
B31; 'Punto de Equilibrio' B11 y B22): 8 cambios validos, 1 repetido y 8 rechazados. Son 43 casos:
17 validos, 3 repetidos y 23 rechazados. Se revisa como guarda cada celda el numero (12.51 en C28;
la fraccion en B11 y B22), el aviso del plazo distinto de 60 meses, que el VAN, la TIR y el punto de
equilibrio vuelven a su valor al devolver los datos, y que los textos fijos A29 y A23 no cambiaron.

La prueba no da por hecho ningun valor del Excel: compara contra lo que lee al empezar.

Uso (desde la carpeta ROBOT PANADERO):
    py -m microservicios.grupo5_financiero.probar_finanzas
Tarda cerca de 4 minutos: cada cambio valido se recalcula con LibreOffice, igual que en produccion.
"""
import os
import shutil
from datetime import timedelta
from pathlib import Path

ORIGINAL = Path("datos_prueba/maestro_prueba.xlsx")
TEMPORAL = Path("datos_prueba/_prueba_finanzas.xlsx")

# Antes de importar conexion_drive: trabajar SOLO sobre la copia temporal
os.environ["ORIGEN_EXCEL"] = "local"
os.environ["EXCEL_LOCAL_PATH"] = str(TEMPORAL)

from openpyxl import load_workbook                                       # noqa: E402

from conexion_drive.recalculo import buscar_errores                     # noqa: E402
from conexion_drive.tablas import leer_tabla, obtener_tabla              # noqa: E402
from conexion_drive.tiempo import ahora_local                            # noqa: E402
from microservicios.grupo5_financiero import finanzas                    # noqa: E402
from microservicios.grupo5_financiero.finanzas import FinanzasInvalido   # noqa: E402

resultados = []
# Donde dice el mapa que esta cada supuesto (la funcion los ubica por su etiqueta; aqui se comprueba que coinciden).
SUPUESTO_CELDA = {campo: definicion["celda_mapa"] for campo, definicion in finanzas.SUPUESTOS.items()}


def caso(numero, descripcion, funcion, *args, debe_fallar=False, estado=None, **kwargs):
    try:
        r = funcion(*args, **kwargs)
    except FinanzasInvalido as error:
        ok = debe_fallar
        detalle = f"rechazado: {str(error)[:100]}"
    else:
        ok = (not debe_fallar) and (estado is None or r["estado"] == estado)
        detalle = f"{r['estado']} {r['celda']} {r['antes']!r} -> {r['despues']!r}" + (" | AVISO" if r.get("avisos") else "")
        caso.ultimo = r
    resultados.append(ok)
    print(f"{'OK   ' if ok else 'FALLO'} {numero:>2}. {descripcion}\n         {detalle}")


def revisar(condicion, texto):
    resultados.append(bool(condicion))
    print(f"{'OK   ' if condicion else 'FALLO'}     {texto}")


def formulas(ruta):
    wb = load_workbook(ruta)
    salida = {}
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                v = c.value
                if hasattr(v, "text"):
                    v = v.text
                if isinstance(v, str) and v.startswith("="):
                    salida[(ws.title, c.coordinate)] = v
    return salida


def valores():
    return load_workbook(TEMPORAL, data_only=True)


def eventos(wb):
    ws, tabla = obtener_tabla(wb, "RegistroEventos", "tblRegistroEventos")
    return leer_tabla(ws, tabla)


def dia_de(leido, fecha):
    return next(d for d in leido["dias"] if d["fecha"] == fecha)


def main():
    if not ORIGINAL.exists():
        print(f"No existe {ORIGINAL}")
        return 1
    shutil.copyfile(ORIGINAL, TEMPORAL)
    print(f"Copia temporal: {TEMPORAL} (el original no se toca)\n")

    hoy = ahora_local().date()
    ayer = hoy - timedelta(days=1)
    wb0 = valores()
    inicio = finanzas.leer_flujo_diario(wb0, hoy=hoy, dias=400)
    primera = inicio["primera"]
    eventos_antes = len(eventos(wb0))
    caja_hoy, caja_ayer, caja_primera = (dia_de(inicio, f)["caja_chica"] or 0 for f in (hoy, ayer, primera))
    total_caja = next(r["total"] for r in inicio["resumen"] if "caja chica" in r["concepto"].casefold())
    anual0 = finanzas.leer_flujo_anual(wb0)
    marzo0 = anual0["salidas"]["meses"][2] or 0
    print(f"Hoy {hoy}; primer dia de la hoja {primera}; saldo de caja al empezar {inicio['saldo']:,.2f}\n")

    print("--- Caja chica del dia ('Flujo caja diario' columna G) ---")
    caso(1, "Caja chica de hoy = 3500 (vales)", finanzas.registrar_caja_chica, hoy, caja_hoy + 3500, "Vales de prueba",
         hoy=hoy, estado="OK")
    leido = finanzas.leer_flujo_diario(valores(), hoy=hoy)
    fila_hoy = dia_de(leido, hoy)
    revisar(fila_hoy["calculado"] and fila_hoy["caja_chica"] == caja_hoy + 3500
            and round(fila_hoy["egresos"] - (fila_hoy["materia_prima"] + fila_hoy["compras_externas"]), 2) == caja_hoy + 3500,
            f"el Excel sumo la caja chica a los egresos de hoy (fila {fila_hoy['fila']}: egresos {fila_hoy['egresos']:,.2f})")
    revisar(round(inicio["saldo"] - leido["saldo"], 2) == 3500,
            f"el saldo de caja ({leido['celda_saldo']}) bajo 3,500: {inicio['saldo']:,.2f} -> {leido['saldo']:,.2f}")
    saldo_ayer = dia_de(leido, ayer)["saldo"]          # ya calculado; la caja chica de hoy no lo mueve
    caso(2, "El mismo total otra vez: SIN_CAMBIO", finanzas.registrar_caja_chica, hoy.isoformat(), caja_hoy + 3500,
         "Vales de prueba", hoy=hoy, estado="SIN_CAMBIO")
    caso(3, "Total nuevo de hoy = 5000.50: REEMPLAZA al anterior", finanzas.registrar_caja_chica,
         hoy.strftime("%d/%m/%Y"), f"{caja_hoy + 5000.5:.2f}".replace(".", ","), "Vales y un flete", hoy=hoy, estado="OK")
    leido = finanzas.leer_flujo_diario(valores(), hoy=hoy)
    revisar(dia_de(leido, hoy)["caja_chica"] == caja_hoy + 5000.5 and round(inicio["saldo"] - leido["saldo"], 2) == 5000.5,
            "queda un solo total en el dia (5,000.50) y el saldo baja 5,000.50 respecto al inicio")
    caso(4, "Caja chica de ayer = 1200 (un dia que ya paso)", finanzas.registrar_caja_chica, ayer, caja_ayer + 1200,
         "Compra de bolsas", hoy=hoy, estado="OK")
    caso(5, "Caja chica del primer dia de la hoja = 800", finanzas.registrar_caja_chica, primera, caja_primera + 800,
         "Vale de arranque", hoy=hoy, estado="OK")
    leido = finanzas.leer_flujo_diario(valores(), hoy=hoy, dias=400)
    total_ahora = next(r["total"] for r in leido["resumen"] if "caja chica" in r["concepto"].casefold())
    revisar(round(total_ahora - total_caja, 2) == 7000.5 and round(inicio["saldo"] - leido["saldo"], 2) == 7000.5
            and round(saldo_ayer - dia_de(leido, ayer)["saldo"], 2) == 2000,
            "tres dias con caja chica: el total del periodo sube 7,000.50, el saldo de ayer baja 2,000 (800 + 1,200) "
            "y el de hoy 7,000.50")

    manana = hoy + timedelta(days=1)
    rechazos = (
        (6, "Fecha de mañana (futura)", (manana, 100, "x")),
        (7, "Fecha anterior al primer dia de la hoja", (primera - timedelta(days=1), 100, "x")),
        (8, "Fecha mal escrita", ("31-31-2026", 100, "x")),
        (9, "Monto negativo", (hoy, -1, "x")),
        (10, "Monto sobre el tope (500,001)", (hoy, 500001, "x")),
        (11, "Monto con letras", (hoy, "tres mil", "x")),
        (12, "Monto con separador de miles (1,500)", (hoy, "1,500", "x")),
        (13, "Motivo vacio", (hoy, 100, "   ")),
        (14, "Motivo que empieza con '='", (hoy, 100, "=SUM(A1)")),
        (15, "Motivo de mas de 120 letras", (hoy, 100, "x" * 121)),
    )
    antes = TEMPORAL.read_bytes()
    for numero, descripcion, argumentos in rechazos:
        caso(numero, f"{descripcion}: RECHAZAR", finanzas.registrar_caja_chica, *argumentos, hoy=hoy, debe_fallar=True)
    revisar(TEMPORAL.read_bytes() == antes, "los 10 rechazos no escribieron nada (el archivo quedo identico)")

    print("\n--- Salidas de efectivo del mes ('flujo caja anual' fila 19) ---")
    caso(16, "Salidas de efectivo de marzo = 25000", finanzas.registrar_salida_efectivo, "marzo", marzo0 + 25000,
         estado="OK")
    anual = finanzas.leer_flujo_anual(valores())
    otros_iguales = all(a == b for i, (a, b) in enumerate(zip(anual0["salidas"]["meses"], anual["salidas"]["meses"])) if i != 2)
    revisar(anual["salidas"]["meses"][2] == marzo0 + 25000 and otros_iguales,
            f"solo cambio la celda de marzo (fila {anual['salidas']['fila']}); los otros 11 meses quedaron igual")
    caso(17, "El mismo monto otra vez: SIN_CAMBIO", finanzas.registrar_salida_efectivo, 3, marzo0 + 25000,
         estado="SIN_CAMBIO")
    caso(18, "Salidas de efectivo de diciembre = 1000.75, con origen SaaS", finanzas.registrar_salida_efectivo, 12,
         "1000,75", origen="SaaS (administracion)", estado="OK")
    antes = TEMPORAL.read_bytes()
    for numero, descripcion, argumentos in ((19, "Mes 13", (13, 100)), (20, "Mes 'brumario'", ("brumario", 100)),
                                            (21, "Monto negativo", ("enero", -5)), (22, "Monto con letras", (1, "mucho")),
                                            (23, "Monto con separador de miles", (1, "25.000"))):
        caso(numero, f"{descripcion}: RECHAZAR", finanzas.registrar_salida_efectivo, *argumentos, debe_fallar=True)
    revisar(TEMPORAL.read_bytes() == antes, "los 5 rechazos no escribieron nada (el archivo quedo identico)")

    print("\n--- Eventos, formulas y regreso a los valores del inicio ---")
    wb = valores()
    nuevos = eventos(wb)[eventos_antes:]
    tipos = [e.get("TIPO_EVENTO") for e in nuevos]
    revisar(tipos == [finanzas.EVENTO_CAJA_CHICA] * 4 + [finanzas.EVENTO_SALIDAS] * 2,
            f"6 eventos nuevos, uno por cada cambio valido: {tipos}")
    revisar(all("->" in str(e.get("DETALLE")) for e in nuevos) and all("motivo:" in str(e.get("DETALLE")) for e in nuevos[:4]),
            "cada evento trae el valor anterior -> el nuevo; los de caja chica traen ademas el motivo")
    revisar([e.get("ORIGEN") for e in nuevos] == ["Robot Python"] * 5 + ["SaaS (administracion)"],
            "origen del evento: 'Robot Python' si nadie entrega otro; 'SaaS (administracion)' cuando lo entrega el SaaS")
    for evento in nuevos:
        print(f"         {evento.get('ID_EVENTO')} {evento.get('TIPO_EVENTO')}: {evento.get('DETALLE')}")

    caso(24, "Devolver la caja chica de hoy a su valor del inicio", finanzas.registrar_caja_chica, hoy, caja_hoy,
         "Regreso de la prueba", hoy=hoy, estado="OK")
    caso(25, "Devolver la de ayer", finanzas.registrar_caja_chica, ayer, caja_ayer, "Regreso de la prueba", hoy=hoy,
         estado="OK")
    caso(26, "Devolver la del primer dia", finanzas.registrar_caja_chica, primera, caja_primera, "Regreso de la prueba",
         hoy=hoy, estado="OK")
    leido = finanzas.leer_flujo_diario(valores(), hoy=hoy)
    revisar(round(leido["saldo"] - inicio["saldo"], 2) == 0,
            f"con la caja chica en sus valores del inicio, el saldo de caja vuelve a {inicio['saldo']:,.2f}")

    print("\n--- Supuestos financieros ('flujo caja proyecto' C28 y B31; 'Punto de Equilibrio' B11 y B22) ---")
    ind0 = finanzas.leer_indicadores(valores())
    eventos_antes = len(eventos(valores()))
    textos0 = {celda: load_workbook(TEMPORAL)[hoja][celda].value
               for hoja, celda in ((finanzas.HOJA_PROYECTO, "A29"), (finanzas.HOJA_EQUILIBRIO, "A23"))}
    print(f"Al empezar: tasa {ind0['tasa']:g} %, plazo {ind0['plazo']:g} meses, mezcla {ind0['mezcla'] * 100:g} %, utilidad "
          f"esperada {ind0['utilidad'] * 100:g} %; VAN {ind0['van']:,.2f}, TIR {ind0['tir'] * 100:.2f} %, punto de equilibrio "
          f"ajustado {ind0['equilibrio_ajustado']:,.0f} unidades")

    def iguales(a, b):
        return all(abs((a[k] or 0) - (b[k] or 0)) < 0.005 for k in ("van", "tir", "recuperacion", "equilibrio_ajustado",
                                                                  "cuota_anual", "tasa", "plazo", "mezcla", "utilidad"))

    tasa_nueva = round(ind0["tasa"] + 1.49, 2) if ind0["tasa"] <= 98 else round(ind0["tasa"] - 1.49, 2)
    caso(27, f"Tasa del prestamo {ind0['tasa']:g} % -> {tasa_nueva:g} %", finanzas.actualizar_supuesto, "tasa",
         f"{tasa_nueva:g}".replace(".", ","), estado="OK")
    r = caso.ultimo
    ind = r["indicadores"]
    tasa_descuento = finanzas.leer_tir_var(valores())["tasa"]
    revisar(r["celda"] == SUPUESTO_CELDA["tasa"] and ind["tasa"] == tasa_nueva and abs(tasa_descuento - tasa_nueva / 100) < 1e-9
            and (ind["cuota_anual"] > ind0["cuota_anual"]) == (tasa_nueva > ind0["tasa"]),
            f"la celda {r['celda']} guarda {ind['tasa']:g} (el porcentaje tal cual); la cuota anual pasa de "
            f"{ind0['cuota_anual']:,.2f} a {ind['cuota_anual']:,.2f} y 'TIR VAR' descuenta al {tasa_descuento * 100:g} %")
    caso(28, "La misma tasa otra vez: SIN_CAMBIO", finanzas.actualizar_supuesto, "tasa", tasa_nueva, estado="SIN_CAMBIO")
    plazo_nuevo = 36 if ind0["plazo"] != 36 else 48
    caso(29, f"Plazo {ind0['plazo']:g} -> {plazo_nuevo} meses: se guarda con AVISO", finanzas.actualizar_supuesto, "plazo",
         plazo_nuevo, estado="OK")
    r = caso.ultimo
    revisar(len(r["avisos"]) == 1 and "cinco años" in r["avisos"][0]
            and r["indicadores"]["cuota_mensual"] > ind["cuota_mensual"],
            f"plazo distinto de {finanzas.PLAZO_DEL_MODELO} meses: guardado con aviso; la cuota mensual sube a "
            f"{r['indicadores']['cuota_mensual']:,.2f}")
    caso(30, f"Devolver el plazo a {ind0['plazo']:g} meses", finanzas.actualizar_supuesto, "plazo", ind0["plazo"], estado="OK")
    revisar(bool(caso.ultimo["avisos"]) == (ind0["plazo"] != finanzas.PLAZO_DEL_MODELO),
            f"con {ind0['plazo']:g} meses " + ("hay aviso" if ind0["plazo"] != finanzas.PLAZO_DEL_MODELO else "no hay aviso"))
    caso(31, f"Devolver la tasa a {ind0['tasa']:g} %", finanzas.actualizar_supuesto, "tasa", ind0["tasa"], estado="OK")
    revisar(iguales(caso.ultimo["indicadores"], ind0),
            "con la tasa y el plazo del inicio, el VAN, la TIR, la recuperacion y el punto de equilibrio vuelven a su valor")

    mezcla_0, utilidad_0 = round(ind0["mezcla"] * 100, 2), round(ind0["utilidad"] * 100, 2)
    mezcla_nueva = mezcla_0 + 10 if mezcla_0 <= 90 else mezcla_0 - 10
    utilidad_nueva = utilidad_0 + 5 if utilidad_0 <= 95 else utilidad_0 - 5
    caso(32, f"Mezcla de ventas de Baguette {mezcla_0:g} % -> {mezcla_nueva:g} %", finanzas.actualizar_supuesto, "mezcla",
         f"{mezcla_nueva:g} %", origen="SaaS (administracion)", estado="OK")
    r = caso.ultimo
    wb = valores()
    equilibrio = finanzas.leer_punto_equilibrio(wb)
    por_celda = {f["celda"]: f["valor"] for f in equilibrio["filas"]}
    fila_mezcla = equilibrio["mezcla"]["fila"]
    proyecto = finanzas.leer_flujo_proyecto(wb)
    unidades = next(f for f in proyecto["filas"] if f["etiqueta"].casefold().startswith("cantidad"))["valores"]
    revisar(r["despues"] == round(mezcla_nueva / 100, 4) and abs(por_celda[f"B{fila_mezcla + 1}"] - (1 - mezcla_nueva / 100)) < 1e-9
            and unidades[0] + unidades[1] in (r["indicadores"]["equilibrio_ajustado"], r["indicadores"]["equilibrio_ajustado"] + 1,
                                              r["indicadores"]["equilibrio_ajustado"] - 1)
            and unidades[0] != unidades[1],
            f"la celda {r['celda']} guarda la fraccion {r['despues']}; CupCake queda en {por_celda[f'B{fila_mezcla + 1}'] * 100:g} % y "
            f"'flujo caja proyecto' lee las unidades nuevas ({unidades[0]:,.0f} Baguette y {unidades[1]:,.0f} CupCake)")
    caso(33, f"Utilidad esperada anual {utilidad_0:g} % -> {utilidad_nueva:g} %", finanzas.actualizar_supuesto, "utilidad",
         utilidad_nueva, estado="OK")
    r2 = caso.ultimo
    revisar(r2["despues"] == round(utilidad_nueva / 100, 4)
            and (r2["indicadores"]["equilibrio_ajustado"] > r["indicadores"]["equilibrio_ajustado"]) == (utilidad_nueva > utilidad_0),
            f"la celda {r2['celda']} guarda la fraccion {r2['despues']}; el punto de equilibrio ajustado pasa de "
            f"{r['indicadores']['equilibrio_ajustado']:,.0f} a {r2['indicadores']['equilibrio_ajustado']:,.0f} unidades")
    caso(34, f"Devolver la mezcla a {mezcla_0:g} %", finanzas.actualizar_supuesto, "mezcla", mezcla_0, estado="OK")
    caso(35, f"Devolver la utilidad esperada a {utilidad_0:g} %", finanzas.actualizar_supuesto, "utilidad", utilidad_0, estado="OK")
    revisar(iguales(caso.ultimo["indicadores"], ind0),
            f"con los cuatro supuestos del inicio: VAN {caso.ultimo['indicadores']['van']:,.2f}, TIR "
            f"{caso.ultimo['indicadores']['tir'] * 100:.2f} % y punto de equilibrio ajustado "
            f"{caso.ultimo['indicadores']['equilibrio_ajustado']:,.0f} unidades, iguales a los del inicio")
    antes = TEMPORAL.read_bytes()
    for numero, descripcion, argumentos in (
            (36, "Tasa 101 %", ("tasa", 101)), (37, "Tasa negativa", ("tasa", -1)), (38, "Tasa con letras", ("tasa", "doce")),
            (39, "Plazo 0", ("plazo", 0)), (40, "Plazo 12.5 (no es entero)", ("plazo", "12.5")),
            (41, f"Plazo {finanzas.TOPE_PLAZO + 1} (sobre el tope)", ("plazo", finanzas.TOPE_PLAZO + 1)),
            (42, "Mezcla 150 %", ("mezcla", 150)), (43, "Un supuesto que no existe ('iva')", ("iva", 13))):
        caso(numero, f"{descripcion}: RECHAZAR", finanzas.actualizar_supuesto, *argumentos, debe_fallar=True)
    revisar(TEMPORAL.read_bytes() == antes, "los 8 rechazos no escribieron nada (el archivo quedo identico)")
    nuevos = eventos(valores())[eventos_antes:]
    revisar([e.get("TIPO_EVENTO") for e in nuevos] == [finanzas.EVENTO_PRESTAMO] * 4 + [finanzas.EVENTO_EQUILIBRIO] * 4
            and all("->" in str(e.get("DETALLE")) for e in nuevos)
            and [e.get("ORIGEN") for e in nuevos] == ["Robot Python"] * 4 + ["SaaS (administracion)"] + ["Robot Python"] * 3,
            "8 eventos nuevos (4 CONFIG_PRESTAMO y 4 CONFIG_PUNTO_EQUILIBRIO), cada uno con el valor anterior -> el nuevo y su origen")
    for evento in nuevos[:1] + nuevos[4:5]:
        print(f"         {evento.get('ID_EVENTO')} {evento.get('TIPO_EVENTO')}: {evento.get('DETALLE')}")
    libro = load_workbook(TEMPORAL)
    revisar(all(libro[hoja][celda].value == textos0[celda]
                for hoja, celda in ((finanzas.HOJA_PROYECTO, "A29"), (finanzas.HOJA_EQUILIBRIO, "A23"))),
            f"los textos fijos no se tocaron: 'flujo caja proyecto'!A29 = {textos0['A29']!r} y 'Punto de Equilibrio'!A23 sigue igual")

    f_antes, f_despues = formulas(ORIGINAL), formulas(TEMPORAL)
    revisar(f_antes == f_despues, f"las {len(f_antes)} formulas del libro quedaron IDENTICAS al original")
    ws_anual = load_workbook(TEMPORAL)[finanzas.HOJA_ANUAL]
    fila_espacio = next(f["fila"] for f in finanzas.leer_flujo_anual(valores())["filas"]
                        if f["etiqueta"].casefold() == "gasto operacional") + 1
    revisar(all(ws_anual[f"{letra}{fila_espacio}"].value is None for letra in "A" + finanzas.COLUMNAS_MES),
            f"la fila {fila_espacio} de '{finanzas.HOJA_ANUAL}' (espacio de orden) sigue vacia: el robot no la toca")
    revisar(not buscar_errores(TEMPORAL.read_bytes()), "la copia temporal quedo con 0 errores de formula")

    print()
    fallos = resultados.count(False)
    if fallos:
        print(f"RESULTADO: {fallos} FALLO(S) de {len(resultados)} revisiones")
        return 1
    print(f"RESULTADO: TODO OK ({len(resultados)} de {len(resultados)} revisiones)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
