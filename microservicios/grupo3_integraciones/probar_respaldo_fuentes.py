"""Prueba del respaldo manual de Solera y Musi (Decisiones Pendientes #56, 3-oct-2026).

Trabaja SOLO sobre una carpeta temporal (datos_prueba/_prueba_respaldo_fuentes) que crea a
partir de datos_prueba/respaldo_fuentes/ y borra al terminar.
No usa internet: las respuestas de los sitios y de Gemini se SIMULAN (se imita la
verificacion contra robots que reciben GitHub Actions y Render).
No toca Google Drive, ni el Excel maestro, ni el cache, ni la carpeta de respaldos.

Uso (desde la carpeta ROBOT PANADERO):
    py -m microservicios.grupo3_integraciones.probar_respaldo_fuentes

Debe terminar en:  RESULTADO: TODO OK (0 lineas ERROR)
"""
import json
import os
import shutil
import sys
import time
from datetime import date

# Todo en modo local y sobre la carpeta temporal. Debe ir ANTES de importar el robot.
os.environ["ORIGEN_EXCEL"] = "local"
os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")
os.environ["RESPALDO_FUENTES_LOCAL_PATH"] = "datos_prueba/_prueba_respaldo_fuentes"
os.environ.pop("RESPALDO_DIAS_AVISO", None)

from openpyxl import Workbook  # noqa: E402

from conexion_drive import config  # noqa: E402

from . import competencia as comp  # noqa: E402
from . import precios_proveedores as pp  # noqa: E402
from . import respaldo_fuentes as rf  # noqa: E402

ORIGINAL = config.RAIZ_PROYECTO / "datos_prueba" / "respaldo_fuentes"
TEMPORAL = rf.carpeta_local()
CACHE_JSON = config.RAIZ_PROYECTO / "datos_prueba" / "cache_competencia.json"
VERIFICACION = b'<html><head><meta http-equiv="refresh" content="0;/.well-known/sgcaptcha/?r=x"></head></html>'
DIA = 24 * 3600
errores = 0


def revisar(condicion, texto):
    global errores
    if condicion:
        print(f"  OK     {texto}")
    else:
        errores += 1
        print(f"  ERROR  {texto}")


def poner_fecha(archivo, dias_atras):
    momento = time.time() - dias_atras * DIA
    os.utime(archivo, (momento, momento))


def reiniciar():
    """Deja al robot como al inicio de una corrida."""
    rf.olvidar()
    pp._solera_no_responde = None
    pp.respaldo_usado = None
    pp._primera_llamada = True
    comp.respaldos_usados.clear()


def precio_del_archivo(productos, sku):
    p = next(x for x in productos if str(x.get("sku")) == sku)
    return round(int(p["prices"]["price"]) / 10 ** int(p["prices"].get("currency_minor_unit", 0)))


def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:       # noqa: BLE001
        pass
    solera_original = sorted(ORIGINAL.glob("respaldo_solera*")) if ORIGINAL.is_dir() else []
    musi_original = sorted(ORIGINAL.glob("respaldo_musi*")) if ORIGINAL.is_dir() else []
    if not solera_original or not musi_original:
        sys.exit(f"Faltan las muestras respaldo_solera... y respaldo_musi... en {ORIGINAL}")
    tamanos = {a.name: a.stat().st_size for a in ORIGINAL.iterdir() if a.is_file()}
    if TEMPORAL.exists():
        shutil.rmtree(TEMPORAL)
    TEMPORAL.mkdir(parents=True)
    solera = TEMPORAL / "respaldo_solera.json"
    musi = TEMPORAL / "respaldo_musi.json"
    shutil.copyfile(solera_original[-1], solera)
    shutil.copyfile(musi_original[-1], musi)
    poner_fecha(solera, 0)
    poner_fecha(musi, 0)
    productos = json.loads(solera.read_text(encoding="utf-8-sig"))

    # Sin pausas ni internet: todo lo que saldria a la red se reemplaza por simulaciones.
    pp.PAUSA_MIN_S = pp.PAUSA_MAX_S = 0
    comp.PAUSA_WEB_S = comp.PAUSA_GEMINI_S = 0
    originales = (pp._get_json, comp._get_json, comp._pedir, comp._guardar_cache,
                  comp.leer_imagen_con_gemini)
    clave_antes = os.environ.get("GEMINI_API_KEY")
    pp.print = comp.print = lambda *a, **k: None      # el robot no escribe en pantalla durante la prueba
    try:
        print("=== 1) Leer la carpeta de respaldos (copia temporal) ===")
        revisar(rf.modo() == "local", f"modo local ({TEMPORAL.name})")
        r_solera, r_musi = rf.leer("solera"), rf.leer("musi")
        revisar(len(r_solera["datos"]) == len(pp.SOLERA),
                f"Solera: {r_solera['resumen']} (se esperan {len(pp.SOLERA)})")
        revisar(r_solera["dias"] == 0 and not r_solera["viejo"], "Solera: respaldo de hoy, no es viejo")
        revisar(r_musi["resumen"] == "1 pagina(s)", f"Musi: {r_musi['resumen']}")

        print("=== 2) Solera con verificacion contra robots -> precios desde el respaldo ===")
        llamadas = []

        def solera_bloqueada(url):
            llamadas.append(url)
            raise pp.FuenteNoResponde("el sitio pide verificacion contra robots (sgcaptcha)")
        pp._get_json = solera_bloqueada
        reiniciar()
        iguales = 0
        for codigo, (_, sku) in pp.SOLERA.items():
            precio, nota = pp.consultar_solera(codigo)
            iguales += precio == precio_del_archivo(productos, sku) and nota.startswith("RESPALDO")
        revisar(iguales == len(pp.SOLERA), f"{iguales} de {len(pp.SOLERA)} precios iguales al archivo, por SKU")
        revisar(len(llamadas) == 1, f"al sitio solo se le pregunto 1 vez (fueron {len(llamadas)}): no se le insiste")
        revisar(pp.respaldo_usado is not None and pp.respaldo_usado["nombre"] == solera.name,
                "queda anotado que archivo de respaldo se uso")

        print("=== 3) Escribir en una hoja de prueba (en memoria): limite del 50 % activo ===")
        wb = Workbook()
        ws = wb.active
        ws.title = pp.HOJA_MP
        for n, (codigo, (_, sku)) in enumerate(pp.SOLERA.items(), start=10):
            ws[f"F{n}"], ws[f"D{n}"] = codigo, precio_del_archivo(productos, sku)
        ws["D10"] = ws["D10"].value + 5            # un precio que cambio poco -> se escribe
        ws["D11"] = ws["D11"].value * 3            # un cambio mayor al 50 % -> se frena
        resultados = []
        for codigo, (_, sku) in pp.SOLERA.items():
            precio, nota = pp.consultar_solera(codigo)
            resultados.append({"hoja": pp.HOJA_MP, "codigo": codigo, "precio": precio, "nota": nota})
        resumen = pp.aplicar_precios(wb, resultados, limite=0.5)
        revisar(len(resumen["cambios"]) == 1 and len(resumen["frenados"]) == 1
                and resumen["iguales"] == len(pp.SOLERA) - 2 and not resumen["fallos"],
                f"1 cambio, 1 frenado por limite, {resumen['iguales']} sin cambio, 0 fallos")

        print("=== 4) Si Solera SI responde, el respaldo no se usa ===")
        reiniciar()
        codigo0, (slug0, sku0) = next(iter(pp.SOLERA.items()))
        pp._get_json = lambda url: [{"sku": sku0, "prices": {"price": "123400", "currency_minor_unit": 2}}]
        precio, nota = pp.consultar_solera(codigo0)
        revisar(precio == 1234 and nota.startswith("via") and pp.respaldo_usado is None,
                f"lectura directa: {precio} ({nota}); respaldo sin usar")
        pp._get_json = lambda url: []
        try:
            pp.consultar_solera(codigo0)
            revisar(False, "producto que ya no existe en el catalogo debe FALLAR")
        except pp.ConsultaFallida as e:
            revisar("no se encontro" in str(e) and pp.respaldo_usado is None,
                    "producto que ya no existe en el catalogo: FALLA y no se tapa con el respaldo")

        print("=== 5) Respaldo viejo, archivo equivocado, SKU faltante y carpeta ausente ===")
        pp._get_json = solera_bloqueada
        poner_fecha(solera, 15)
        reiniciar()
        r = rf.leer("solera")
        revisar(r["dias"] == 15 and r["viejo"] and "RESPALDO VIEJO" in rf.texto(r),
                f"respaldo de 15 dias: se usa y queda marcado ({rf.texto(r)[-45:]})")
        precio, _ = pp.consultar_solera(codigo0)
        revisar(precio == precio_del_archivo(productos, sku0), "respaldo viejo: el precio igual se entrega")

        equivocado = TEMPORAL / "respaldo_solera (1).json"
        equivocado.write_bytes(VERIFICACION)
        poner_fecha(equivocado, 0)
        reiniciar()
        r = rf.leer("solera")
        revisar(r["nombre"] == solera.name and len(r["descartados"]) == 1,
                "archivo mas nuevo que no es el catalogo: se descarta y se usa el anterior valido")
        equivocado.unlink()

        incompleto = TEMPORAL / "respaldo_solera_incompleto.json"
        incompleto.write_text(json.dumps([p for p in productos if str(p.get("sku")) != sku0]),
                              encoding="utf-8")
        poner_fecha(incompleto, 0)
        reiniciar()
        try:
            pp.consultar_solera(codigo0)
            revisar(False, "SKU que no viene en el respaldo debe FALLAR")
        except pp.ConsultaFallida as e:
            revisar(f"SKU {sku0}" in str(e), f"SKU que no viene en el respaldo: FALLA solo ese producto ({codigo0})")
        otro = [c for c in pp.SOLERA if c != codigo0][0]
        revisar(pp.consultar_solera(otro)[0] == precio_del_archivo(productos, pp.SOLERA[otro][1]),
                f"los demas productos del mismo respaldo siguen bien ({otro})")
        incompleto.unlink()
        poner_fecha(solera, 0)

        os.environ["RESPALDO_FUENTES_LOCAL_PATH"] = "datos_prueba/_carpeta_que_no_existe"
        reiniciar()
        try:
            pp.consultar_solera(codigo0)
            revisar(False, "sin carpeta de respaldos debe FALLAR")
        except pp.ConsultaFallida as e:
            revisar("sin respaldo" in str(e), "sin carpeta de respaldos: FALLA como antes (la celda no se toca)")
        os.environ["RESPALDO_FUENTES_LOCAL_PATH"] = "datos_prueba/_prueba_respaldo_fuentes"

        print("=== 6) Musi con verificacion contra robots -> pagina desde el respaldo ===")
        def musi_bloqueado(url):
            raise comp.ConsultaFallida("el sitio pide verificacion contra robots (sgcaptcha)")

        def red_bloqueada(url, datos=None, encabezados=None, espera=None):
            return VERIFICACION, "text/html"
        comp._get_json, comp._pedir = musi_bloqueado, red_bloqueada
        comp._guardar_cache = lambda *a, **k: None          # la prueba no escribe ningun cache
        reiniciar()
        datos_musi = comp.COMPETIDORES["Musi"]
        imagenes, notas = comp.consultar_paginas("Musi", datos_musi)
        visibles = [i for i in imagenes if i["visible"]]
        revisar(len(visibles) > 0 and any("RESPALDO manual" in n for n in notas),
                f"{len(imagenes)} imagen(es), {len(visibles)} visible(s), leidas del respaldo")

        hoy = date.today()
        if CACHE_JSON.exists():
            cache = json.loads(CACHE_JSON.read_text(encoding="utf-8"))
            en_cache = sum(comp._canonica(i["url"]) in cache for i in visibles)

            def gemini_prohibido(*a, **k):
                raise AssertionError("no debia llamarse a Gemini")
            comp.leer_imagen_con_gemini = gemini_prohibido
            reiniciar()
            if en_cache == len(visibles):
                resultado = comp.consultar_competidor("Musi", datos_musi, True, dict(cache), hoy)
                revisar(resultado["estado"] in ("SIN_PROMO", "VIGENTE"),
                        f"las {len(visibles)} imagenes salen del cache, 0 llamadas a Gemini -> {resultado['estado']}")
                resumen_txt, _, _ = comp.textos_para_excel(resultado, hoy)
                fecha_respaldo = resultado["respaldo"]["fecha"].strftime("%d-%m-%Y")
                revisar(fecha_respaldo in resumen_txt, f"el texto lleva la fecha del respaldo: {resumen_txt}")
            else:
                print(f"  (aviso) solo {en_cache} de {len(visibles)} imagenes estan en el cache JSON: "
                      "se omite la prueba de 0 llamadas")
            comp.leer_imagen_con_gemini = originales[4]

            print("=== 7) Imagen NUEVA (no esta en el cache): la abre Gemini por Google ===")
            cuerpos = []
            lectura = {"es_promocion": True, "producto": "Prueba", "precio_colones": 999,
                       "vigencia_desde": None, "vigencia_hasta": None, "texto_vigencia": None}

            def red_simulada(estado_google):
                def pedir(url, datos=None, encabezados=None, espera=None):
                    if "generativelanguage" not in url:
                        return VERIFICACION, "text/html"           # el sitio no entrega la imagen
                    cuerpos.append(json.loads(datos.decode("utf-8")))
                    respuesta = {"candidates": [{
                        "content": {"parts": [{"text": "```json\n" + json.dumps(lectura) + "\n```"}]},
                        "urlContextMetadata": {"urlMetadata": [{"urlRetrievalStatus": estado_google}]}}]}
                    return json.dumps(respuesta).encode("utf-8"), "application/json"
                return pedir
            os.environ["GEMINI_API_KEY"] = "clave-de-prueba"
            comp._modelo_elegido, comp._modelos_alternos, comp._cupo_agotado = "gemini-prueba", [], False
            sin_una = dict(cache)
            nueva = comp._canonica(visibles[-1]["url"])
            sin_una.pop(nueva, None)
            comp._pedir = red_simulada("URL_RETRIEVAL_STATUS_SUCCESS")
            reiniciar()
            resultado = comp.consultar_competidor("Musi", datos_musi, True, sin_una, hoy)
            revisar(resultado["estado"] != "FALLO" and sin_una.get(nueva, {}).get("leido") == lectura,
                    "el robot no pudo bajarla; Gemini la abrio por Google y la lectura quedo en el cache")
            revisar(len(cuerpos) == 1 and cuerpos[0].get("tools") == [{"url_context": {}}]
                    and visibles[-1]["url"] in cuerpos[0]["contents"][0]["parts"][0]["text"],
                    "a Gemini se le envio la direccion de la imagen con la herramienta URL context")

            sin_una.pop(nueva, None)
            comp._pedir = red_simulada("URL_RETRIEVAL_STATUS_ERROR")
            reiniciar()
            resultado = comp.consultar_competidor("Musi", datos_musi, True, sin_una, hoy)
            revisar(resultado["estado"] == "FALLO" and nueva not in sin_una,
                    "si Google tampoco puede abrirla: lectura incompleta, no se afirma nada (Regla #19)")
        else:
            print(f"  (aviso) no existe {CACHE_JSON.name}: se omiten las pruebas con el cache")

        print("=== 8) Detalles ===")
        revisar(comp._es_imagen(b"\xff\xd8\xff\xe0" + b"0" * 20, "text/html")
                and comp._es_imagen(b"\x89PNG\r\n" + b"0" * 20, "")
                and not comp._es_imagen(VERIFICACION, "text/html"),
                "se distingue una imagen real de una pagina de verificacion")
        os.environ["RESPALDO_FUENTES_LOCAL_PATH"] = "datos_prueba/_carpeta_que_no_existe"
        comp._get_json, comp._pedir = musi_bloqueado, red_bloqueada
        reiniciar()
        resultado = comp.consultar_competidor("Musi", datos_musi, True, {}, hoy)
        revisar(resultado["estado"] == "FALLO", "Musi sin carpeta de respaldos: FALLO como antes (celdas no se tocan)")
        os.environ["RESPALDO_FUENTES_LOCAL_PATH"] = "datos_prueba/_prueba_respaldo_fuentes"
    finally:
        (pp._get_json, comp._get_json, comp._pedir, comp._guardar_cache,
         comp.leer_imagen_con_gemini) = originales
        if clave_antes is None:
            os.environ.pop("GEMINI_API_KEY", None)
        else:
            os.environ["GEMINI_API_KEY"] = clave_antes
        del pp.print, comp.print
        shutil.rmtree(TEMPORAL, ignore_errors=True)

    print("=== 9) Nada del proyecto cambio ===")
    revisar(not TEMPORAL.exists(), "la carpeta temporal se borro")
    revisar({a.name: a.stat().st_size for a in ORIGINAL.iterdir() if a.is_file()} == tamanos,
            "la carpeta datos_prueba/respaldo_fuentes quedo igual")

    print()
    print("RESULTADO: TODO OK (0 lineas ERROR)" if errores == 0
          else f"RESULTADO: {errores} linea(s) ERROR - no instalar; avisar")


if __name__ == "__main__":
    main()
