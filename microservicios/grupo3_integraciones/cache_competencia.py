"""Cache de competencia en Excel: CACHE_COMPETENCIA.xlsx (Decisiones Pendientes #20).

Para que sirve (hoja Instrucciones del propio archivo):
  Guarda lo que Gemini leyo de cada imagen de promociones de la competencia. Cada imagen
  se lee UNA sola vez: si su URL ya esta en la tabla TablaCache, competencia.py usa esa
  lectura y no llama a Gemini. Ademas es un HISTORIAL ACUMULADO: una fila por lectura;
  si una imagen se relee, vale la lectura mas reciente. Nunca se borran filas.

Donde vive el cache (lo decide este modulo, competencia.py no cambia de uso):
  - Uso normal (Drive): el Excel de Google Drive cuyo ID esta en CACHE_COMPETENCIA_FILE_ID
    (.env en la PC; "Secrets" en GitHub Actions). Es el que usa la tarea diaria de las 7:00.
  - Con --local: la copia de prueba datos_prueba/CACHE_COMPETENCIA_prueba.xlsx.
  - Si no hay ninguno de los dos, competencia.py sigue con el archivo de antes
    (datos_prueba/cache_competencia.json): nada se rompe.

Patron del proyecto (Regla #1): descargar -> editar en memoria con openpyxl -> recalcular
con LibreOffice (hoja Control: filas usadas, imagenes distintas) -> subir.
Igual que el Excel maestro (Regla #12): NUNCA abrir este archivo con Google Sheets.

Decision del usuario (2-oct-2026): NO se construye nada para el limite de filas de la hoja
Control (B4 = 1,000,000). Con unas 9 imagenes y pocas nuevas por semana no se alcanza;
depurar el cache queda como tarea de la empresa en produccion (trabajo futuro).

Uso (desde la carpeta ROBOT PANADERO):
    py -m microservicios.grupo3_integraciones.cache_competencia --ver --local
    py -m microservicios.grupo3_integraciones.cache_competencia --ver
        (solo LEE: cuantas lecturas e imagenes hay; no escribe nada)
    py -m microservicios.grupo3_integraciones.cache_competencia --migrar-json --local
    py -m microservicios.grupo3_integraciones.cache_competencia --migrar-json
        (pasa al Excel las lecturas de datos_prueba/cache_competencia.json que aun no
         esten, con ORIGEN = "migrado JSON". Segura de repetir: la 2a corrida no agrega
         nada. El archivo JSON NO se borra.)
"""
import json
import logging
import os
import sys
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

# --local: igual que las migraciones. Debe ir ANTES de importar conexion_drive.config.
if "--local" in sys.argv:
    os.environ["ORIGEN_EXCEL"] = "local"
    os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

from conexion_drive import config  # noqa: E402  (carga el .env: CACHE_COMPETENCIA_FILE_ID)
from conexion_drive.tiempo import ahora_local  # noqa: E402

log = logging.getLogger("cache_competencia")

HOJA = "Cache"
TABLA = "TablaCache"
COLUMNAS = ["ID_LECTURA", "FECHA_LECTURA", "COMPETIDOR", "URL_IMAGEN", "FECHA_SUBIDA_IMAGEN",
            "ES_PROMOCION", "PRODUCTO", "PRECIO_CRC", "VIGENCIA_DESDE", "VIGENCIA_HASTA",
            "TEXTO_VIGENCIA", "ESTADO_AL_LEER", "MODELO_GEMINI", "ORIGEN"]
COPIA_LOCAL = "datos_prueba/CACHE_COMPETENCIA_prueba.xlsx"
JSON_VIEJO = config.RAIZ_PROYECTO / "datos_prueba" / "cache_competencia.json"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ORIGEN_GEMINI = "Gemini"
ORIGEN_JSON = "migrado JSON"


class ErrorCacheCompetencia(Exception):
    """El Excel del cache no se pudo leer o escribir, o no tiene la forma esperada."""


# ---------------------------------------------------------------- Donde vive el cache
def id_drive():
    return (os.getenv("CACHE_COMPETENCIA_FILE_ID") or "").strip()


def ruta_local():
    valor = (os.getenv("CACHE_COMPETENCIA_LOCAL_PATH") or COPIA_LOCAL).strip()
    ruta = Path(valor)
    return ruta if ruta.is_absolute() else config.RAIZ_PROYECTO / ruta


def modo():
    """'drive', 'local' o None (None = no hay Excel de cache: se usa el JSON de antes)."""
    if config.ORIGEN_EXCEL == "local":
        return "local" if ruta_local().exists() else None
    return "drive" if id_drive() else None


# ---------------------------------------------------------------- Conversiones
def _a_fecha(valor):
    """Texto 'AAAA-MM-DD', date o datetime -> date. Otra cosa -> None."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return date.fromisoformat(str(valor)[:10])
    except (TypeError, ValueError):
        return None


def _celda_fecha(valor):
    """Para escribir: una fecha valida va como fecha de Excel; si Gemini devolvio otro
    texto se guarda tal cual (asi no se pierde lo leido); vacio queda vacio."""
    if valor in (None, ""):
        return None
    fecha = _a_fecha(valor)
    return datetime(fecha.year, fecha.month, fecha.day) if fecha else str(valor)


def _texto_fecha(valor):
    """Para leer: fecha de Excel -> 'AAAA-MM-DD' (como la entrega Gemini); texto -> igual."""
    if valor in (None, ""):
        return None
    if isinstance(valor, (date, datetime)):
        return _a_fecha(valor).isoformat()
    return str(valor)


def _numero_entero(valor):
    """1350.0 -> 1350 (LibreOffice devuelve los enteros como decimal); lo demas, igual."""
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor


def fila_desde_lectura(id_lectura, momento, competidor, url, fecha_subida, leido, estado,
                       modelo, origen):
    """Convierte una lectura de Gemini en los 14 valores de una fila de TablaCache."""
    if isinstance(momento, date) and not isinstance(momento, datetime):
        momento = datetime(momento.year, momento.month, momento.day)
    return [
        id_lectura,
        momento,
        competidor,
        url,
        _celda_fecha(fecha_subida),
        "SI" if leido.get("es_promocion") else "NO",
        leido.get("producto"),
        leido.get("precio_colones"),
        _celda_fecha(leido.get("vigencia_desde")),
        _celda_fecha(leido.get("vigencia_hasta")),
        leido.get("texto_vigencia"),
        estado,
        modelo,
        origen,
    ]


def lectura_desde_fila(valores):
    """Lo contrario: de los 14 valores de una fila al dict que usa competencia.py."""
    fila = dict(zip(COLUMNAS, valores))
    leido = {
        "es_promocion": str(fila["ES_PROMOCION"] or "").strip().upper() == "SI",
        "producto": fila["PRODUCTO"],
        "precio_colones": _numero_entero(fila["PRECIO_CRC"]),
        "vigencia_desde": _texto_fecha(fila["VIGENCIA_DESDE"]),
        "vigencia_hasta": _texto_fecha(fila["VIGENCIA_HASTA"]),
        "texto_vigencia": fila["TEXTO_VIGENCIA"],
    }
    fecha = _a_fecha(fila["FECHA_LECTURA"])
    return {"leido": leido, "fecha_lectura": fecha.isoformat() if fecha else None}


# ---------------------------------------------------------------- El Excel del cache
class CacheExcel:
    """El Excel CACHE_COMPETENCIA abierto en memoria. Se abre una vez por corrida."""

    def __init__(self):
        self.modo = modo()
        if not self.modo:
            raise ErrorCacheCompetencia(
                "No hay Excel de cache: falta CACHE_COMPETENCIA_FILE_ID en el .env (Drive) "
                f"o la copia local {ruta_local()} (con --local).")
        self.destino = ("COPIA LOCAL " + str(ruta_local())) if self.modo == "local" \
            else "GOOGLE DRIVE (CACHE_COMPETENCIA.xlsx)"
        self.wb = load_workbook(BytesIO(self._descargar()))
        if HOJA not in self.wb.sheetnames or TABLA not in self.wb[HOJA].tables:
            raise ErrorCacheCompetencia(
                f"El Excel del cache no tiene la hoja '{HOJA}' con la tabla '{TABLA}'.")
        self.ws = self.wb[HOJA]
        encabezados = [self.ws.cell(1, c).value for c in range(1, len(COLUMNAS) + 1)]
        if encabezados != COLUMNAS:
            raise ErrorCacheCompetencia(
                f"Las columnas de {TABLA} no son las esperadas. Encontre: {encabezados}")
        self.pendientes = 0

    # --- leer
    def filas(self):
        """Lista de (numero_de_fila, [14 valores]) de las filas que tienen URL."""
        col_url = COLUMNAS.index("URL_IMAGEN") + 1
        resultado = []
        for n in range(2, self.ws.max_row + 1):
            if self.ws.cell(n, col_url).value not in (None, ""):
                resultado.append((n, [self.ws.cell(n, c).value
                                      for c in range(1, len(COLUMNAS) + 1)]))
        return resultado

    def como_dict(self):
        """{url: {"leido": {...}, "fecha_lectura": "AAAA-MM-DD"}} igual que el JSON de antes.
        Si una imagen tiene varias lecturas, vale la mas reciente (la de mas abajo)."""
        cache = {}
        for _, valores in self.filas():
            cache[valores[COLUMNAS.index("URL_IMAGEN")]] = lectura_desde_fila(valores)
        return cache

    def resumen(self):
        filas = self.filas()
        urls = {v[COLUMNAS.index("URL_IMAGEN")] for _, v in filas}
        return {"lecturas": len(filas), "imagenes": len(urls)}

    # --- escribir
    def _siguiente_id(self):
        numeros = [0]
        for _, valores in self.filas():
            texto = str(valores[0] or "")
            if texto.startswith("LEC-") and texto[4:].isdigit():
                numeros.append(int(texto[4:]))
        return f"LEC-{max(numeros) + 1:06d}"

    def agregar(self, competidor, url, fecha_subida, leido, estado, modelo,
                origen=ORIGEN_GEMINI, momento=None):
        """Agrega UNA fila al final de TablaCache (en memoria; falta guardar())."""
        filas = self.filas()
        fila = (filas[-1][0] + 1) if filas else 2
        valores = fila_desde_lectura(self._siguiente_id(), momento or ahora_local(), competidor,
                                     url, fecha_subida, leido, estado, modelo, origen)
        for columna, valor in enumerate(valores, start=1):
            celda = self.ws.cell(fila, columna)
            if fila > 2:                      # mismo formato que la primera fila de datos
                modelo_celda = self.ws.cell(2, columna)
                celda.number_format = modelo_celda.number_format
                celda.font = copy(modelo_celda.font)
                celda.alignment = copy(modelo_celda.alignment)
                celda.border = copy(modelo_celda.border)
                celda.fill = copy(modelo_celda.fill)
            if isinstance(valor, str) and valor.startswith("="):
                valor = "'" + valor           # nunca escribir texto que Excel tome por formula
            celda.value = valor
        tabla = self.ws.tables[TABLA]
        tabla.ref = f"A1:N{fila}"
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
        self.pendientes += 1
        return valores[0]

    def guardar(self):
        """Recalcula (hoja Control) y guarda: copia local o Google Drive."""
        if not self.pendientes:
            return False
        memoria = BytesIO()
        self.wb.save(memoria)
        contenido = memoria.getvalue()
        if config.RECALCULAR:
            from conexion_drive.recalculo import buscar_errores, recalcular_bytes
            contenido = recalcular_bytes(contenido)
            errores = buscar_errores(contenido, maximo=20)
            if errores:
                raise ErrorCacheCompetencia(
                    f"El cache quedo con errores de formula; NO se guardo. Celdas: {errores[:10]}")
        self._subir(contenido)
        log.info("Cache de competencia: %s fila(s) nueva(s) guardada(s) en %s.",
                 self.pendientes, self.destino)
        self.pendientes = 0
        return True

    # --- Drive / copia local
    def _descargar(self):
        if self.modo == "local":
            return ruta_local().read_bytes()
        from conexion_drive.drive_excel import _explicar_error_google, _servicio_drive
        from googleapiclient.http import MediaIoBaseDownload
        try:
            datos = _servicio_drive().files().get(
                fileId=id_drive(), fields="id,name,mimeType", supportsAllDrives=True).execute()
            if datos.get("mimeType") != MIME_XLSX:
                raise ErrorCacheCompetencia(
                    f"El cache en Drive no es un .xlsx (es {datos.get('mimeType')}). "
                    "No debe convertirse a Google Sheets.")
            solicitud = _servicio_drive().files().get_media(fileId=id_drive(),
                                                            supportsAllDrives=True)
            memoria = BytesIO()
            descargador = MediaIoBaseDownload(memoria, solicitud)
            terminado = False
            while not terminado:
                _, terminado = descargador.next_chunk()
            return memoria.getvalue()
        except ErrorCacheCompetencia:
            raise
        except Exception as error:
            raise ErrorCacheCompetencia(
                "No pude descargar CACHE_COMPETENCIA.xlsx: "
                + _explicar_error_google(error).replace("EXCEL_MAESTRO_FILE_ID",
                                                        "CACHE_COMPETENCIA_FILE_ID")) from error

    def _subir(self, contenido):
        if self.modo == "local":
            ruta = ruta_local()
            temporal = ruta.with_suffix(".tmp")
            temporal.write_bytes(contenido)
            temporal.replace(ruta)
            return
        from conexion_drive.drive_excel import _explicar_error_google, _servicio_drive
        from googleapiclient.http import MediaIoBaseUpload
        try:
            media = MediaIoBaseUpload(BytesIO(contenido), mimetype=MIME_XLSX, resumable=True)
            _servicio_drive().files().update(
                fileId=id_drive(), media_body=media, fields="id,version,modifiedTime",
                supportsAllDrives=True).execute()
        except Exception as error:
            raise ErrorCacheCompetencia(
                "No pude subir CACHE_COMPETENCIA.xlsx: "
                + _explicar_error_google(error).replace("EXCEL_MAESTRO_FILE_ID",
                                                        "CACHE_COMPETENCIA_FILE_ID")) from error


# ---------------------------------------------------------------- Migrar el JSON viejo
def migrar_json(cache_excel, ruta_json=None):
    """Pasa al Excel las lecturas del JSON que aun no esten. Devuelve (nuevas, ya_estaban)."""
    from .competencia import COMPETIDORES, evaluar, fecha_subida   # import tardio (sin circulo)

    ruta_json = Path(ruta_json or JSON_VIEJO)
    if not ruta_json.exists():
        raise ErrorCacheCompetencia(f"No existe el archivo {ruta_json}")
    viejo = json.loads(ruta_json.read_text(encoding="utf-8"))
    ya = set(cache_excel.como_dict())
    nuevas = ya_estaban = 0
    for url, dato in viejo.items():
        if url in ya:
            ya_estaban += 1
            continue
        leido = dato.get("leido") or {}
        dia = _a_fecha(dato.get("fecha_lectura")) or ahora_local().date()
        subida = fecha_subida(url)
        competidor = next((nombre for nombre, d in COMPETIDORES.items()
                           if url.startswith(d["base"])), "desconocido")
        estado, _ = evaluar(leido, subida, dia)
        cache_excel.agregar(competidor, url, subida, leido, estado, "no registrado",
                            origen=ORIGEN_JSON, momento=dia)
        nuevas += 1
    return nuevas, ya_estaban


def main():
    logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")
    args = sys.argv[1:]
    if "--ver" not in args and "--migrar-json" not in args:
        print(__doc__)
        return
    cache = CacheExcel()
    print(f"=== Cache de competencia: {cache.destino} ===")
    antes = cache.resumen()
    print(f"Antes: {antes['lecturas']} lectura(s) de {antes['imagenes']} imagen(es).")
    if "--migrar-json" in args:
        nuevas, ya_estaban = migrar_json(cache)
        if nuevas:
            cache.guardar()
            despues = cache.resumen()
            print(f"Migradas {nuevas} lectura(s) del JSON (ORIGEN '{ORIGEN_JSON}'); "
                  f"{ya_estaban} ya estaban.")
            print(f"Despues: {despues['lecturas']} lectura(s) de {despues['imagenes']} imagen(es).")
        else:
            print(f"Ya estaba migrado: las {ya_estaban} lectura(s) del JSON ya estan en el "
                  "Excel. No se escribio nada.")
    else:
        for _, v in cache.filas():
            fila = dict(zip(COLUMNAS, v))
            print(f"  {fila['ID_LECTURA']} | {fila['COMPETIDOR']:<8} | {fila['ES_PROMOCION']} | "
                  f"{str(fila['PRODUCTO'] or '-')[:45]:<45} | {fila['ESTADO_AL_LEER']} | "
                  f"{fila['ORIGEN']}")
        print("Modo --ver: no se escribio nada.")
    print("LISTO.")


if __name__ == "__main__":
    main()
