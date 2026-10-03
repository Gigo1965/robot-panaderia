"""Prueba del cache de competencia en Excel (Decisiones Pendientes #20).

Trabaja SOLO sobre una copia temporal (datos_prueba/_prueba_cache_excel.xlsx) que crea a
partir de datos_prueba/CACHE_COMPETENCIA_prueba.xlsx y borra al terminar.
No toca Google Drive, ni el Excel maestro, ni la copia de prueba, ni el JSON. No llama a
Gemini ni a internet.

Uso (desde la carpeta ROBOT PANADERO):
    py -m microservicios.grupo3_integraciones.probar_cache_excel

Debe terminar en:  RESULTADO: TODO OK (0 lineas ERROR)
"""
import json
import os
import shutil
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

# Todo en modo local y sobre la copia temporal. Debe ir ANTES de importar el robot.
os.environ["ORIGEN_EXCEL"] = "local"
os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")
os.environ["CACHE_COMPETENCIA_LOCAL_PATH"] = "datos_prueba/_prueba_cache_excel.xlsx"

from openpyxl import load_workbook  # noqa: E402

from conexion_drive import config  # noqa: E402

from . import cache_competencia as cc  # noqa: E402
from . import competencia  # noqa: E402

ORIGINAL = config.RAIZ_PROYECTO / "datos_prueba" / "CACHE_COMPETENCIA_prueba.xlsx"
TEMPORAL = cc.ruta_local()
errores = 0


def revisar(condicion, texto):
    global errores
    if condicion:
        print(f"  OK     {texto}")
    else:
        errores += 1
        print(f"  ERROR  {texto}")


def main():
    if not ORIGINAL.exists():
        sys.exit(f"Falta la copia de prueba del cache: {ORIGINAL}")
    if not cc.JSON_VIEJO.exists():
        sys.exit(f"Falta el archivo {cc.JSON_VIEJO}")
    shutil.copyfile(ORIGINAL, TEMPORAL)
    tamano_original = ORIGINAL.stat().st_size
    try:
        viejo = json.loads(cc.JSON_VIEJO.read_text(encoding="utf-8"))

        print("=== 1) Abrir la copia temporal del cache ===")
        cache = cc.CacheExcel()
        revisar(cache.modo == "local", f"modo local ({cache.destino})")
        antes = cache.resumen()
        print(f"         tenia {antes['lecturas']} lectura(s) de {antes['imagenes']} imagen(es)")

        print("=== 2) Migrar el JSON (ORIGEN 'migrado JSON') ===")
        nuevas, ya = cc.migrar_json(cache)
        revisar(nuevas + ya == len(viejo), f"{nuevas} nuevas + {ya} que ya estaban = {len(viejo)} del JSON")
        cache.guardar()
        cache = cc.CacheExcel()                       # volver a abrir desde el disco
        despues = cache.resumen()
        revisar(despues["lecturas"] == antes["lecturas"] + nuevas,
                f"lecturas en el Excel: {antes['lecturas']} -> {despues['lecturas']}")
        leidas = cache.como_dict()
        iguales = decisiones = 0
        for url, dato in viejo.items():
            if leidas.get(url, {}).get("leido") == dato["leido"]:
                iguales += 1
            for hoy in (date(2025, 6, 15), date(2026, 10, 2)):
                if (competencia.evaluar(leidas[url]["leido"], competencia.fecha_subida(url), hoy)
                        == competencia.evaluar(dato["leido"], competencia.fecha_subida(url), hoy)):
                    decisiones += 1
        revisar(iguales == len(viejo), f"{iguales} de {len(viejo)} lecturas identicas a las del JSON")
        revisar(decisiones == 2 * len(viejo),
                f"misma decision (vigente / vencida / precio) en {decisiones} de {2 * len(viejo)} casos")
        ids = [v[0] for _, v in cache.filas()]
        revisar(len(set(ids)) == len(ids) and all(str(i).startswith("LEC-") for i in ids),
                f"ID_LECTURA unicos y con forma LEC-000000 (ultimo: {ids[-1]})")

        print("=== 3) Segunda migracion: no debe agregar nada ===")
        nuevas2, ya2 = cc.migrar_json(cache)
        revisar(nuevas2 == 0 and ya2 == len(viejo), f"0 nuevas, {ya2} ya estaban")
        revisar(cache.guardar() is False, "sin cambios: no se guardo de nuevo")

        print("=== 4) Lectura nueva de Gemini (simulada) y relectura de la misma imagen ===")
        url = "https://tiendasmusi.com/wp-content/uploads/2026/10/PRUEBA-ROBOT.jpg"
        leido = {"es_promocion": True, "producto": "Prueba del robot", "precio_colones": 1234,
                 "vigencia_desde": "2026-10-01", "vigencia_hasta": "2026-10-31",
                 "texto_vigencia": "=texto que empieza con igual"}
        hoy = date(2026, 10, 2)
        competencia._cache_excel = cache
        competencia._guardar_cache({}, nueva=("Musi", url, competencia.fecha_subida(url), leido, hoy))
        leido2 = dict(leido, precio_colones=999, texto_vigencia=None)
        competencia._guardar_cache({}, nueva=("Musi", url, competencia.fecha_subida(url), leido2, hoy))
        cache = cc.CacheExcel()
        final = cache.resumen()
        revisar(final["lecturas"] == despues["lecturas"] + 2 and final["imagenes"] == despues["imagenes"] + 1,
                f"2 filas nuevas y 1 imagen nueva ({final['lecturas']} lecturas, {final['imagenes']} imagenes)")
        revisar(cache.como_dict()[url]["leido"]["precio_colones"] == 999,
                "si una imagen se relee, vale la lectura mas reciente (999)")
        ultima = dict(zip(cc.COLUMNAS, cache.filas()[-1][1]))
        revisar(str(ultima["ESTADO_AL_LEER"]).startswith("VIGENTE") and ultima["ORIGEN"] == "Gemini",
                f"ESTADO_AL_LEER '{ultima['ESTADO_AL_LEER']}' y ORIGEN '{ultima['ORIGEN']}'")
        penultima = dict(zip(cc.COLUMNAS, cache.filas()[-2][1]))
        revisar(not str(penultima["TEXTO_VIGENCIA"]).startswith("="),
                "un texto que empieza con '=' no se escribe como formula")

        print("=== 5) Forma del Excel despues de guardar ===")
        contenido = TEMPORAL.read_bytes()
        wb = load_workbook(BytesIO(contenido))
        valores = load_workbook(BytesIO(contenido), data_only=True)
        tabla = wb["Cache"].tables["TablaCache"]
        revisar(tabla.ref == f"A1:N{final['lecturas'] + 1}", f"TablaCache = {tabla.ref}")
        revisar(wb.sheetnames == ["Cache", "Control", "Instrucciones"], "3 hojas: Cache, Control, Instrucciones")
        revisar(str(wb["Control"]["B5"].value).startswith("=COUNTA("), "Control B5 sigue siendo formula")
        if config.RECALCULAR:
            revisar(valores["Control"]["B5"].value == final["lecturas"],
                    f"Control 'Filas usadas' = {valores['Control']['B5'].value}")
            revisar(valores["Control"]["B7"].value == final["imagenes"],
                    f"Control 'Imagenes distintas' = {valores['Control']['B7'].value}")
        revisar(ORIGINAL.stat().st_size == tamano_original, "la copia de prueba del cache NO se toco")
    finally:
        competencia._cache_excel = None
        for archivo in (TEMPORAL, TEMPORAL.with_suffix(".tmp")):
            if archivo.exists():
                archivo.unlink()

    print(f"\nRESULTADO: {'TODO OK' if not errores else 'HAY ERRORES'} ({errores} lineas ERROR)")


if __name__ == "__main__":
    main()
