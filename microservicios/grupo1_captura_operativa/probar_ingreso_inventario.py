"""Prueba de Ingreso de Inventario SOBRE LA COPIA LOCAL del Excel (no toca Google Drive).

Usa la misma copia que las otras pruebas:  datos_prueba/maestro_prueba.xlsx

Uso (desde la carpeta ROBOT PANADERO):
    py -m microservicios.grupo1_captura_operativa.probar_ingreso_inventario

#32 (V15, 27-sep-2026): pasos 6 a 8 revisan el costo congelado. Al final debe
decir "RESULTADO: TODO OK (0 lineas ERROR)".
V15 (28-sep-2026): paso 9 = el cafe con codigo Ca01 (caja de la MP) se rechaza en
compras y el autocompletado solo ofrece Cf01. La leche vigilada usa el nombre
actual de la TCM (antes 'Leche Suli...', que ya no existe y no se mostraba).
"""
import os

# Forzar modo local ANTES de importar el resto (asi nunca toca Drive)
os.environ["ORIGEN_EXCEL"] = "local"
os.environ.setdefault("EXCEL_LOCAL_PATH", "datos_prueba/maestro_prueba.xlsx")

import logging                                                   # noqa: E402
import time as reloj                                             # noqa: E402

from openpyxl import load_workbook                               # noqa: E402

from conexion_drive import leer_excel                            # noqa: E402
from .ingreso_inventario import (  # noqa: E402
    CostoNoDisponible, ProductoNoReconocido, UsarCodigoDeVenta, _costo_congelado, buscar_opciones, leer_costos_tcm,
    leer_productos_inventario, registrar_ingreso, registrar_ingresos,
)

logging.basicConfig(level=logging.WARNING)
VIGILADOS = ("Baguette", "Leche Semidescremada Dos Pinos -1 L")   # V15: nombre actual de la TCM (fila 10)


def existencias():
    ws = leer_excel()["Control de Inventario"]
    return {ws.cell(f, 1).value: ws.cell(f, 6).value for f in range(5, ws.max_row + 1)
            if ws.cell(f, 1).value in VIGILADOS}


def paso(titulo):
    print(f"\n--- {titulo} ---")


ERRORES = []


def revisar(condicion, texto):
    """Imprime OK o ERROR y cuenta los errores para el resumen final."""
    print(f"  {'OK   ' if condicion else 'ERROR'} {texto}")
    if not condicion:
        ERRORES.append(texto)


def celdas_ingreso(fila):
    """Lee E (formula o numero), el valor de E, F y L de una fila, directo de la copia local."""
    ruta = os.environ["EXCEL_LOCAL_PATH"]
    formulas = load_workbook(ruta)["Ingreso de Inventario"]
    valores = load_workbook(ruta, data_only=True)["Ingreso de Inventario"]
    return formulas[f"E{fila}"].value, valores[f"E{fila}"].value, valores[f"F{fila}"].value, valores[f"L{fila}"].value


if __name__ == "__main__":
    marca = int(reloj.time()) % 100000
    antes = existencias()
    print("Existencia ANTES (Control de Inventario):", antes)

    paso("0) Autocompletado del SaaS: lo que ve la persona mientras escribe")
    productos = leer_productos_inventario(leer_excel())
    for escrito in ("har", "lech", "bag", "mejorante 5", "cafe"):
        opciones = ", ".join(f"{nombre[:32]} ({codigo})" for codigo, nombre in buscar_opciones(productos, escrito, 3))
        print(f"  escribe '{escrito}'  ->  {opciones}")

    costos_antes = leer_costos_tcm(leer_excel())

    paso("1) Factura de compra: 10 Harina (materia prima) + 12 Leche (reventa)")
    compra = registrar_ingresos(
        [{"producto": "Ha01", "cantidad": 10}, {"producto": "L01", "cantidad": 12}],
        factura=f"F-PRUEBA-{marca}", proveedor="Proveedor A")
    print(" ", compra)

    paso("2) Lote de produccion propia: 500 Baguette")
    print(" ", registrar_ingreso("P01", 500, factura=f"L-PRUEBA-{marca}"))

    paso("3) La misma factura otra vez (debe decir DUPLICADO)")
    print(" ", registrar_ingresos(
        [{"producto": "Ha01", "cantidad": 10}, {"producto": "L01", "cantidad": 12}],
        factura=f"F-PRUEBA-{marca}", proveedor="Proveedor A"))

    for numero, escrito in ((4, "MEJORANTE STAR PAN PLUS 5 KILOS"), (5, "CAFE NARANJO PURO 250 GR X 10 UNIDADES")):
        paso(f"{numero}) Nombre mal escrito (caso real): debe rechazar y sugerir")
        try:
            registrar_ingreso(escrito, 2, factura=f"F-PRUEBA-{marca}-{numero}", proveedor="Proveedor B")
        except ProductoNoReconocido as error:
            print(f"  Rechazado correctamente: {error}")

    paso("6) #32: el costo de la compra queda como NUMERO igual al de la TCM")
    for linea in compra.get("registradas", []):
        e_texto, e_valor, f_valor, _ = celdas_ingreso(linea["fila"])
        esperado = costos_antes.get(linea["producto"])
        revisar(isinstance(e_texto, (int, float)),
                f"fila {linea['fila']} E es numero (no formula): {str(e_texto)[:40]}")
        revisar(e_valor == esperado, f"fila {linea['fila']} E = {e_valor} (TCM: {esperado})")
        revisar(f_valor == linea["cantidad"] * esperado,
                f"fila {linea['fila']} F = {f_valor} (= {linea['cantidad']} x {esperado})")

    paso("7) #32 + Regla #32: compra de cafe (Cf01) = costo por paquete y 'Compras Externas Netas'")
    cafe = registrar_ingreso("Cf01", 10, factura=f"F-PRUEBA-{marca}-CAFE", proveedor="Proveedor B")
    fila_cafe = cafe["registradas"][0]["fila"]
    e_texto, e_valor, _, l_valor = celdas_ingreso(fila_cafe)
    nombre_cafe = cafe["registradas"][0]["producto"]
    revisar(isinstance(e_texto, (int, float)) and e_valor == costos_antes.get(nombre_cafe),
            f"fila {fila_cafe} E = {e_valor} (TCM: {costos_antes.get(nombre_cafe)})")
    revisar(l_valor == "Compras Externas Netas", f"fila {fila_cafe} L = {l_valor!r}")

    paso("8) #32: producto sin costo valido en la TCM -> se rechaza sin escribir")
    for costo in (0, None, -5, "abc"):
        try:
            _costo_congelado("Producto de prueba", {"Producto de prueba": costo})
            revisar(False, f"costo {costo!r} debio rechazarse")
        except CostoNoDisponible:
            revisar(True, f"costo {costo!r} rechazado correctamente")

    paso("9) V15: cafe con codigo Ca01 (caja de la MP) -> se rechaza y sugiere Cf01")
    opciones_cafe = [codigo for codigo, _ in buscar_opciones(productos, "cafe", 8)]
    revisar("Ca01" not in opciones_cafe and "Cf01" in opciones_cafe,
            f"autocompletado 'cafe' solo ofrece Cf01: {opciones_cafe}")
    filas_antes = leer_excel()["Ingreso de Inventario"].max_row
    for escrito in ("Ca01", "ca01"):
        try:
            registrar_ingreso(escrito, 1, factura=f"F-PRUEBA-{marca}-CA01", proveedor="Proveedor B")
            revisar(False, f"'{escrito}' debio rechazarse")
        except UsarCodigoDeVenta as error:
            revisar(error.codigo_venta == "Cf01" and error.factor == 10, f"'{escrito}' rechazado: {error}")
    revisar(leer_excel()["Ingreso de Inventario"].max_row == filas_antes,
            f"no se escribio ninguna fila (max_row sigue en {filas_antes})")

    despues = existencias()
    print("\nExistencia DESPUES (Control de Inventario):", despues)
    print("  Baguette debe subir 500 y la Leche 12.")
    print("Abre datos_prueba/maestro_prueba.xlsx y revisa el final de 'Ingreso de Inventario'.")
    if ERRORES:
        print(f"\nRESULTADO: HAY {len(ERRORES)} lineas ERROR:")
        for texto in ERRORES:
            print("  -", texto)
    else:
        print("\nRESULTADO: TODO OK (0 lineas ERROR)")
